"""Лист «Сверка со списком»: организации из списков — в одной строке отмеченные столбцы списка и платежи
из реестров (только если выбраны и списки, и реестры)."""
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from analysis_reconcile import reconcile
from report_fast import title_rows
from report_styles import BORDER, CENTER, NUM, fill_color
from utils import format_money

HEAD_ROW = 4
TOP = 10                # в подсказке — крупнейшие связки
LIST_COLOR, REG_COLOR = "DBEAFE", "DCFCE7"     # заголовки: столбцы из списка / из реестров
GRAY = "Серый заголовок — общий столбец (из списка).\n"
BLUE = "Синий заголовок — столбец из СПИСКА (вы отметили его галочкой).\n"
GREEN = "Зелёный заголовок — посчитано по РЕЕСТРАМ платежей.\n"
HINTS = {"ИНН": GRAY + "ИНН организации из списка — по нему ищутся её платежи в реестрах.",
         "Название": GRAY + "Название организации, как в списке.",
         "Где в списке": GRAY + "В каких таблицах (листах) списка есть эта организация.",
         "Заплатил всего": GREEN + "Сумма платежей, где эта организация — плательщик (по всем реестрам). "
                                   "Наведите на сумму — кому платила.",
         "Платежей": GREEN + "Сколько платежей она сделала.",
         "Получил": GREEN + "Сумма платежей, где эта организация — получатель. Наведите на сумму — от кого."}


def _cell(ws, v, **style):
    c = WriteOnlyCell(ws, value=v)
    for k, val in style.items():
        setattr(c, k, val)
    return c


def who_hint(entries, title: str, i: int = None) -> str | None:
    """Подсказка «кому платил / от кого получал»: крупные сверху; i — одна часть реестров, None — все."""
    pick = [(n, inn, s[i] if i is not None else sum(s), c[i] if i is not None else sum(c)) for n, inn, s, c in entries]
    pick = sorted((x for x in pick if x[3]), key=lambda x: -x[2])
    if not pick:
        return None
    lines = [title] + [f"• {n or '—'} ({inn}): {format_money(s)} — {c} плат." for n, inn, s, c in pick[:TOP]]
    if len(pick) > TOP:
        lines.append(f"…и ещё {len(pick) - TOP} на {format_money(sum(x[2] for x in pick[TOP:]))}")
    return "\n".join(lines)


def _money(ws, v, hint=None):
    """Сумма с подсказкой при наведении (кому платил / от кого получал)."""
    if not v:
        return None
    c = _cell(ws, v, number_format=NUM)
    if hint:
        c.comment = Comment(hint, "Распределение по ГИБ", width=480, height=24 + 17 * (hint.count("\n") + 1))
    return c


def write_reconcile_sheet(wb, data) -> dict | None:
    """→ {"orgs": сколько организаций в списке} или None, если сверять нечего."""
    reg, heads, rows = reconcile(data)
    if not rows:
        return None
    ws = wb.create_sheet("Сверка со списком")
    ws.sheet_properties.tabColor = "16A34A"
    ws.sheet_view.showGridLines = False
    own = ["№", "ИНН", "Название", "Где в списке"]
    from_list = [f"{p.label}: {h}" for p, h in heads]
    from_reg = [f"Заплатил: {p.label}" for p in reg] + ["Заплатил всего", "Платежей", "Получил"]
    for i, w in enumerate([6, 13, 46, 30] + [20] * len(from_list) + [18] * len(reg) + [18, 10, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"D{HEAD_ROW + 1}"
    title_rows(ws, f"СВЕРКА СО СПИСКОМ — {len(rows)}", "16A34A",
               "В одной строке: синие столбцы — из списка (отмеченные столбцы), зелёные — из реестров платежей. "
               "Наведите на заголовок — откуда столбец; на сумму платежей — кому платил.", len(own) + len(from_list) + len(from_reg))
    hints = dict(HINTS, **{f"Заплатил: {p.label}": GREEN + f"Реестр «{p.source}»: сколько заплатила за «{p.label}». "
                           "Наведите на сумму — кому платила." for p in reg})
    hints.update({t: BLUE + f"Файл «{p.source}», {p.label}, столбец «{h}»." for t, (p, h) in zip(from_list, heads)})
    cells = []
    for t, color in [(t, "E2E8F0") for t in own] + [(t, LIST_COLOR) for t in from_list] + [(t, REG_COLOR) for t in from_reg]:
        c = _cell(ws, t, font=Font(bold=True), fill=fill_color(color), border=BORDER, alignment=CENTER)
        if t in hints:
            c.comment = Comment(hints[t], "Распределение по ГИБ", width=340, height=110)
        cells.append(c)
    ws.append(cells)
    for k, r in enumerate(rows, 1):
        per_part = [_money(ws, v, who_hint(r.to, f"Кому платил за «{p.label}»:", i))
                    for i, (v, p) in enumerate(zip(r.paid, reg))]
        ws.append([k, r.inn, r.name, r.where] + [_money(ws, v) for v in r.values] + per_part
                  + [_money(ws, round(sum(r.paid), 2), who_hint(r.to, "Кому платил (всего):")), r.payments or None,
                     _money(ws, r.got, who_hint(r.frm, "От кого получал:"))])
    ws.auto_filter.ref = f"A{HEAD_ROW}:{get_column_letter(len(cells))}{HEAD_ROW + len(rows)}"
    return {"orgs": len(rows)}
