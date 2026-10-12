"""Чтение Excel для предпросмотра (openpyxl) и запись в Бучет через сам Excel (COM).

Запись идёт через установленный Microsoft Excel, а не openpyxl: так сохраняются
картинки, форматирование, а формулы (Ҳамагӣ, Ҷамъ) растягиваются самим Excel.
"""
import os
import re
import shutil
import tempfile
from datetime import datetime

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

XL_TO_RIGHT = -4161
XL_DOWN = -4121
XL_FORMAT_FROM_LEFT_OR_ABOVE = 0
XL_CALC_MANUAL = -4135
XL_OPENXML = 51
XL_XLSM = 52
XL_GREEN = 220 + 252 * 256 + 231 * 65536   # светло-зелёный (#DCFCE7), Excel хранит как BGR

CELL_RE = re.compile(r"^\s*([A-Za-z]{1,3})\s*(\d+)\s*$")
GIB_HEADER_RE = re.compile(r"^\s*(\d{4})(?!\d)")


# ---------------------------------------------------------------- адреса ячеек

def parse_cell(addr: str):
    """'B5' -> (5, 2). None, если адрес некорректный."""
    m = CELL_RE.match(addr or "")
    if not m:
        return None
    return int(m.group(2)), column_index_from_string(m.group(1).upper())


def cell_name(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def gib_code(value) -> str | None:
    """Заголовок столбца -> код ГИБ (2111, '2111', '2111 Музди маош')."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    m = GIB_HEADER_RE.match(str(value))
    return m.group(1) if m else None


# ---------------------------------------------------------------- чтение

class Book:
    """Значения всех листов книги для показа в интерфейсе."""

    def __init__(self, path: str):
        self.path = path
        read_path = _xls_to_xlsx(path) if path.lower().endswith(".xls") else path
        wb = openpyxl.load_workbook(read_path, read_only=True, data_only=True)
        self.sheets: dict[str, list[list]] = {}
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            width = max((len(r) for r in rows), default=0)
            # обрезаем хвостовые пустые строки
            while rows and all(v is None for v in rows[-1]):
                rows.pop()
            self.sheets[ws.title] = [r + [None] * (width - len(r)) for r in rows]
        wb.close()

    @staticmethod
    def value(data, row: int, col: int):
        if 1 <= row <= len(data) and 1 <= col <= len(data[row - 1]):
            return data[row - 1][col - 1]
        return None


def _xls_to_xlsx(path: str) -> str:
    """Старый .xls читаем через Excel: сохраняем временную копию в .xlsx."""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    out = os.path.join(tempfile.gettempdir(), f"gib_preview_{os.getpid()}.xlsx")
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(os.path.abspath(path), ReadOnly=True)
        wb.SaveAs(out, FileFormat=XL_OPENXML)
        wb.Close(False)
    finally:
        xl.Quit()
    return out


# ---------------------------------------------------------------- запись

class WriteError(Exception):
    pass


def fill_in_place(path, sheet_name, purpose_col, header_row, items, progress=lambda p, t: None):
    """Режим «Заполнить ГИБ в файле»: суммы ставятся в те же строки, где стоят назначения.

    items: список (строка, {код_ГИБ: сумма}).
    header_row=None — шапки нет (данные с первой строки): вставляем над ними пустую строку под шапку.
    """
    import pythoncom
    import pywintypes
    import win32com.client

    if not items:
        raise WriteError("Нет строк для заполнения.")
    pythoncom.CoInitialize()
    progress(2, "Подключаюсь к Excel…")
    xl, wb, own = _open_workbook(win32com.client, pywintypes, path)
    backup = None
    old_calc = old_screen = None
    try:
        if wb.ReadOnly:
            raise WriteError("Файл открыт только для чтения (возможно, его открыл другой пользователь).")
        progress(5, "Делаю копию файла…")
        backup = _make_backup(path, wb, own)
        try:
            ws = wb.Worksheets(sheet_name)
        except pywintypes.com_error:
            raise WriteError(f"Лист «{sheet_name}» не найден.")

        old_screen = xl.ScreenUpdating
        xl.ScreenUpdating = False
        try:
            old_calc = xl.Calculation
            xl.Calculation = XL_CALC_MANUAL
        except pywintypes.com_error:
            pass

        header_inserted = False
        if header_row is None:
            first = min(r for r, _ in items)
            ws.Rows(first).Insert(Shift=XL_DOWN)
            header_row = first
            items = [(r + 1, g) for r, g in items]
            header_inserted = True

        progress(15, "Проверяю столбцы ГИБ…")
        needed = sorted({code for _, gibs in items for code in gibs})
        created = _ensure_gib_columns(ws, header_row, needed, purpose_col)

        progress(50, "Ставлю суммы…")
        gib_cols = _gib_columns(ws, header_row)
        r1 = min(r for r, _ in items)
        r2 = max(r for r, _ in items)
        replaced = 0
        for code, col in gib_cols.items():
            mine = {r: g[code] for r, g in items if code in g}
            if not mine:
                continue
            cur = _col_values(ws, col, r1, r2, "Formula")
            out = list(cur)
            for r, amount in mine.items():
                old = cur[r - r1]
                if old not in ("", None) and _num_or_none(old) != amount:
                    replaced += 1
                out[r - r1] = amount
            ws.Range(ws.Cells(r1, col), ws.Cells(r2, col)).Formula = [(v,) for v in out]
            for r in mine:
                ws.Cells(r, col).Interior.Color = XL_GREEN      # видно, что вписала программа
        for code in created:
            ws.Cells(header_row, gib_cols[code]).Interior.Color = XL_GREEN

        if old_calc is not None:
            xl.Calculation = old_calc
            old_calc = None
        xl.Calculate()
        progress(90, "Сверяю итоги…")
        check = _verify_rows(ws, items, header_row, r2)
        progress(95, "Сохраняю…")
        wb.Save()
        progress(100, "Готово")
        return {
            "filled": len(items),
            "created": created,
            "replaced": replaced,
            "header_inserted": header_inserted,
            "backup": backup,
            "left_open": not own,
            "check": check,
        }
    finally:
        try:
            if old_calc is not None:
                xl.Calculation = old_calc
            if old_screen is not None:
                xl.ScreenUpdating = old_screen
        except Exception:
            pass
        if own:
            try:
                wb.Close(False)
            finally:
                xl.Quit()


def embed_automation(src, dst, template, sheet_name, header_row, purpose_col, progress=lambda p, t: None):
    """«Встроить автоматику в Excel»: копия файла в формате .xlsm с макросом распределения по ГИБ.

    Листы файла копируются в готовый шаблон с макросом (assets/gib_template.xlsm) — так не нужна
    настройка «Доверять доступ к объектной модели VBA». Лист с назначениями помечается скрытыми
    именами GIB_ON / GIB_HDR / GIB_COL: макрос работает только на нём. Исходный файл не меняется.
    После встраивания макрос сразу заполняет ГИБ для назначений, которые уже есть в файле.
    """
    import pythoncom
    import pywintypes
    import win32com.client

    if os.path.normcase(os.path.abspath(src)) == os.path.normcase(os.path.abspath(dst)):
        raise WriteError("Новый файл должен отличаться от исходного.")
    if not os.path.exists(template):
        raise WriteError(f"Не найден шаблон с макросом:\n{template}")
    pythoncom.CoInitialize()
    progress(5, "Запускаю Excel…")
    try:
        xl = win32com.client.DispatchEx("Excel.Application")
    except pywintypes.com_error:
        raise WriteError("Не удалось запустить Microsoft Excel.")
    xl.Visible = False
    xl.DisplayAlerts = False
    xl.EnableEvents = False
    xl.ScreenUpdating = False
    book = source = None
    try:
        book = xl.Workbooks.Open(os.path.abspath(template), ReadOnly=True)
        source = xl.Workbooks.Open(os.path.abspath(src), ReadOnly=True, UpdateLinks=0)
        progress(20, "Копирую листы…")
        source.Sheets.Copy(Before=book.Sheets(1))
        book.Worksheets("_gib_tmp").Delete()
        ws = book.Worksheets(sheet_name)

        # верхняя строка — под кнопку «Пересчитать ГИБ» и надпись «Макросы выключены»
        if xl.WorksheetFunction.CountA(ws.Rows(1)) > 0 or ws.Range("A1").MergeCells:
            ws.Rows(1).Insert(Shift=XL_DOWN)
            header_row += 1
        if ws.Rows(1).RowHeight < 30:
            ws.Rows(1).RowHeight = 30

        for name, value in (("GIB_ON", 1), ("GIB_HDR", header_row), ("GIB_COL", purpose_col)):
            ws.Names.Add(Name=name, RefersTo=f"={value}", Visible=False)

        progress(35, "Добавляю кнопку…")
        _add_controls(ws)

        source.Close(False)
        source = None
        progress(50, "Сохраняю…")
        if os.path.exists(dst):
            os.remove(dst)
        book.SaveAs(os.path.abspath(dst), FileFormat=XL_XLSM)
        # формулы между листами не должны ссылаться на исходный файл
        for link in book.LinkSources(1) or ():
            if os.path.normcase(link) == os.path.normcase(os.path.abspath(src)):
                book.ChangeLink(link, os.path.abspath(dst), 1)

        progress(65, "Заполняю ГИБ…")
        xl.EnableEvents = True
        book.Activate()
        ws.Activate()
        xl.Run("GIB_RecalcSheet")
        report = xl.Run("GIB_LastReport")
        if str(report).startswith("ERROR"):
            raise WriteError(f"Макрос сообщил об ошибке: {report[7:]}")
        progress(90, "Сохраняю…")
        book.Save()
        progress(100, "Готово")
        return {"path": os.path.abspath(dst), "report": report}
    finally:
        try:
            if source is not None:
                source.Close(False)
            if book is not None:
                book.Close(False)
        finally:
            xl.Quit()


def _add_controls(ws):
    """Кнопка «Пересчитать ГИБ» и красная надпись «Макросы выключены» (её прячет сам макрос)."""
    btn = ws.Buttons().Add(3, 4, 130, 22)
    btn.Name = "btnGibRecalc"
    btn.Caption = "Пересчитать ГИБ"
    btn.OnAction = "GIB_RecalcSheet"
    btn.PrintObject = False

    warn = ws.Shapes.AddTextbox(1, 2, 1, 640, 28)
    warn.Name = "gibMacroWarning"
    warn.Fill.ForeColor.RGB = 38 * 256 * 256 + 38 * 256 + 220          # #DC2626 (BGR)
    warn.Line.Visible = False
    tf = warn.TextFrame2
    tf.TextRange.Text = ("⚠ Макросы выключены — нажмите «Включить содержимое» на жёлтой панели вверху. "
                         "Без этого ГИБ не заполняются.")
    tf.TextRange.Font.Fill.ForeColor.RGB = 0xFFFFFF
    tf.TextRange.Font.Bold = True
    tf.TextRange.Font.Size = 10
    tf.VerticalAnchor = 3                                              # по центру
    try:
        warn.DrawingObject.PrintObject = False
    except Exception:
        pass


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def _num_or_none(v):
    try:
        return round(float(str(v).replace(",", ".")), 2)
    except ValueError:
        return None


def _verify_rows(ws, items, header_row, last_row):
    """Сверка после записи: читаем ячейки обратно из Excel.

    По каждому ГИБ: сколько должны были поставить, сколько реально стоит в этих строках,
    и сходится ли итог под таблицей («Ҷамъ») с суммой столбца.
    """
    used_last, _ = _used_bounds(ws)
    expected = {}
    for _, gibs in items:
        for code, amount in gibs.items():
            expected[code] = round(expected.get(code, 0.0) + amount, 2)

    result = []
    for code, col in sorted(_gib_columns(ws, header_row).items()):
        values = _col_values(ws, col, header_row + 1, last_row)
        written = round(sum(_num(values[r - header_row - 1]) for r, g in items if code in g), 2)
        column_sum = round(sum(_num(v) for v in values), 2)
        total = None
        if used_last > last_row:
            formulas = _col_values(ws, col, last_row + 1, used_last, "Formula")
            for i, f in enumerate(formulas):
                if isinstance(f, str) and f.startswith("="):
                    total = round(_num(ws.Cells(last_row + 1 + i, col).Value), 2)
                    break
        exp = expected.get(code, 0.0)
        if not exp and not column_sum:
            continue
        result.append({
            "code": code,
            "expected": exp,
            "written": written,
            "column_sum": column_sum,
            "total": total,
            "written_ok": abs(written - exp) < 0.005,
            "total_ok": total is None or abs(total - column_sum) < 0.005,
        })
    return {
        "gibs": result,
        "written_sum": round(sum(r["written"] for r in result), 2),
        "ok": all(r["written_ok"] and r["total_ok"] for r in result),
    }


def _open_workbook(win32, pywintypes, path):
    """Если книга уже открыта у пользователя в Excel — пишем прямо в неё."""
    full = os.path.normcase(os.path.abspath(path))
    try:
        running = win32.GetActiveObject("Excel.Application")
        for wb in running.Workbooks:
            if os.path.normcase(wb.FullName) == full:
                return running, wb, False
    except pywintypes.com_error:
        pass
    try:
        xl = win32.DispatchEx("Excel.Application")
    except pywintypes.com_error:
        raise WriteError("Не удалось запустить Microsoft Excel. Для записи нужен установленный Excel.")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(os.path.abspath(path), UpdateLinks=0)
    except pywintypes.com_error as e:
        xl.Quit()
        raise WriteError(f"Excel не смог открыть файл: {e}")
    return xl, wb, True


def _make_backup(path, wb, own):
    """Копия файла до изменений в папку «_резервные копии» рядом с ним.

    Если файл открыт у пользователя в Excel — сохраняем копию из Excel, чтобы в неё
    попали и несохранённые правки. Без копии файл не меняем.
    """
    folder = os.path.join(os.path.dirname(os.path.abspath(path)), "_резервные копии")
    name, ext = os.path.splitext(os.path.basename(path))
    dst = os.path.join(folder, f"{name} {datetime.now():%Y-%m-%d %H-%M-%S}{ext}")
    try:
        os.makedirs(folder, exist_ok=True)
        if own:
            shutil.copy2(path, dst)
        else:
            wb.SaveCopyAs(dst)
    except Exception as e:
        raise WriteError(f"Не удалось сделать копию файла, поэтому файл не изменён.\n{e}")
    return dst


def _used_bounds(ws):
    ur = ws.UsedRange
    return ur.Row + ur.Rows.Count - 1, ur.Column + ur.Columns.Count - 1


def _row_values(ws, row, last_col):
    v = ws.Range(ws.Cells(row, 1), ws.Cells(row, last_col)).Value
    return list(v[0]) if isinstance(v, tuple) else [v]


def _col_values(ws, col, r1, r2, prop="Value"):
    v = getattr(ws.Range(ws.Cells(r1, col), ws.Cells(r2, col)), prop)
    return [x[0] for x in v] if isinstance(v, tuple) else [v]


def _set_column(ws, row, col, values):
    ws.Range(ws.Cells(row, col), ws.Cells(row + len(values) - 1, col)).Value = [(v,) for v in values]


def _gib_columns(ws, header_row):
    _, last_col = _used_bounds(ws)
    cols = {}
    for i, v in enumerate(_row_values(ws, header_row, last_col), start=1):
        code = gib_code(v)
        if code and code not in cols:
            cols[code] = i
    return cols


# ----- столбцы

def _ensure_gib_columns(ws, header_row, needed, pcol):
    """Создаёт недостающие столбцы ГИБ по возрастанию кода. Возвращает список созданных."""
    created = []
    if needed and ws.Cells(header_row, pcol).Value in (None, ""):
        ws.Cells(header_row, pcol).Value = "Номгӯи корхонаҳо"     # пустой Бучет — подписываем столбец
    for code in needed:
        cols = _gib_columns(ws, header_row)
        if code in cols:
            continue
        created.append(code)
        if not cols:
            # столбцов ГИБ ещё нет — ставим в конец шапки, правее назначений
            _, last_col = _used_bounds(ws)
            vals = _row_values(ws, header_row, last_col)
            last = max((i for i, v in enumerate(vals, 1) if v not in (None, "")), default=0)
            ws.Cells(header_row, max(last, pcol) + 1).Value = int(code)
            continue
        last_gib_col = max(cols.values())
        if (max(cols, key=int) < code and _header(ws, header_row, last_gib_col + 1) in (None, "")
                and not _has_formulas(ws, last_gib_col)):
            # таблица без итогов (например, новый Бучет) — просто дописываем справа
            ws.Cells(header_row, last_gib_col + 1).Value = int(code)
            continue

        # Новый столбец вставляем строго внутрь диапазона (SUM(G:Z), итоги, объединённый заголовок),
        # иначе Excel не растянет формулы. Порядок столбцов — по возрастанию кода.
        ordered = sorted(cols.items(), key=lambda kv: kv[1])
        first_col, last_gib_col = ordered[0][1], ordered[-1][1]
        bigger = [c for k, c in ordered if k > code]
        swap = None
        if bigger and bigger[0] != first_col:
            pos = bigger[0]                                    # между соседями
        elif bigger:
            pos, swap = first_col + 1, first_col               # самый маленький: после первого + обмен
        elif _header(ws, header_row, last_gib_col + 1) not in (None, ""):
            pos = last_gib_col + 1                             # самый большой, справа есть «Бар маблағ»
        elif len(ordered) > 1:
            pos, swap = last_gib_col, last_gib_col + 1         # самый большой: перед последним + обмен
        else:
            pos = last_gib_col + 1

        ws.Columns(pos).Insert(Shift=XL_TO_RIGHT, CopyOrigin=XL_FORMAT_FROM_LEFT_OR_ABOVE)
        neighbour = pos - 1 if pos - 1 >= first_col else pos + 1
        _copy_column_formulas(ws, neighbour, pos, header_row)
        ws.Cells(header_row, pos).Value = int(code)
        if swap is not None:
            _swap_constants(ws, pos, swap)
    return created


def _header(ws, header_row, col):
    return ws.Cells(header_row, col).Value


def _has_formulas(ws, col):
    last_row, _ = _used_bounds(ws)
    return any(isinstance(f, str) and f.startswith("=") for f in _col_values(ws, col, 1, last_row, "Formula"))


def _copy_column_formulas(ws, src, dst, header_row):
    """Формулы соседнего столбца ГИБ (например, итог «Ҷамъ») копируем в новый столбец."""
    last_row, _ = _used_bounds(ws)
    formulas = _col_values(ws, src, 1, last_row, "FormulaR1C1")
    for r, f in enumerate(formulas, start=1):
        if r != header_row and isinstance(f, str) and f.startswith("="):
            ws.Cells(r, dst).FormulaR1C1 = f


def _swap_constants(ws, a, b):
    """Меняет местами значения двух столбцов (кроме формул — они у каждого считают свой столбец)."""
    last_row, _ = _used_bounds(ws)
    fa = _col_values(ws, a, 1, last_row, "Formula")
    fb = _col_values(ws, b, 1, last_row, "Formula")
    va = _col_values(ws, a, 1, last_row)
    vb = _col_values(ws, b, 1, last_row)
    for r in range(1, last_row + 1):
        is_fa = isinstance(fa[r - 1], str) and fa[r - 1].startswith("=")
        is_fb = isinstance(fb[r - 1], str) and fb[r - 1].startswith("=")
        if is_fa or is_fb or va[r - 1] == vb[r - 1]:
            continue
        ws.Cells(r, a).Value = vb[r - 1]
        ws.Cells(r, b).Value = va[r - 1]
