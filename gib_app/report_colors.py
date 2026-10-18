"""Цвет ошибок — общий для листов Excel и экрана результата."""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ERR_COLOR = "FED7AA"


def cf_fill(color: str):
    """Заливка для условного форматирования: Excel берёт цвет фона (bgColor), а не цвет узора."""
    return PatternFill("solid", start_color=color, end_color=color)


def add_error_color(ws, first_row: int, last_row: int, error_col: int):
    """Ячейка «Ошибки» — оранжевая, если не пустая."""
    if last_row < first_row:
        return
    col = get_column_letter(error_col)
    ws.conditional_formatting.add(f"{col}{first_row}:{col}{last_row}", FormulaRule(
        formula=[f'LEN(${col}{first_row})>0'], fill=cf_fill(ERR_COLOR), font=Font(bold=True)))
