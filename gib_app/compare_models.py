"""Таблица с ИНН на листе (результат compare_finder.find_tables)."""
from dataclasses import dataclass, field
from openpyxl.utils import get_column_letter


@dataclass
class SideList:
    key_col: int
    name_col: int | None
    cols: list                                   # столбцы списка (для подписи)
    value_cols: list                             # числовые столбцы
    rows: dict = field(default_factory=dict)     # ключ сравнения -> строка
    keys: dict = field(default_factory=dict)     # ключ сравнения -> ИНН как в файле
    all_rows: dict = field(default_factory=dict) # ключ сравнения -> все строки (для дубликатов)
    year: str | None = None
    block: list = field(default_factory=list)    # столбцы самого списка (№, название, ИНН…)
    side: str = ""                               # «левая часть» / «правая часть» листа
    header: int = 0                              # строка шапки
    last_row: int = 0                            # последняя заполненная строка таблицы
    in_text: bool = False                        # ИНН внутри текста: «ООО А (ИНН 123456789)»
    fixed_money: bool = False                    # столбцы сумм выбрал пользователь — суммируются все
    by_name: bool = False                        # ИНН нет — ключ — название (без кавычек и форм)
    separate: bool = False                       # столбцы сумм не складывать: каждый — отдельной частью
    repeats: bool = False                        # организация в нескольких строках — норма (операции, выписка)
    hidden_rows: set = field(default_factory=set)  # скрытые строки листа (номера строк Excel)
    total_col: int | None = None                 # колонка «Итого» в каждой строке (сверка с суммой колонок)
    shared: set = field(default_factory=set)     # ключи из строк, где было несколько организаций
    row_names: dict = field(default_factory=dict)  # (ключ, строка) → название (если ячейка разбита на части)

    @property
    def span(self) -> str:
        """«A–C» — где список виден в Excel."""
        return f"{get_column_letter(min(self.block))}–{get_column_letter(max(self.block))}"
