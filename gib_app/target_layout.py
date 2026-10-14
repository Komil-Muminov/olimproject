"""Таблица в файле, куда переносится выписка (через Excel COM): столбцы, свободные строки, итоги."""
import re
from collections import Counter
from dataclasses import dataclass

from openpyxl.utils import get_column_letter

from constants import XL_DOWN
from excel_columns import col_values, insert_column, row_values, sum_row_below, used_bounds
from excel_reconcile import SIMPLE_SUM_RE
from gib_columns import target_cols
from statement import doc_text
from utils import find_col, norm_text, to_number

BALANCE_AMOUNT_RE = re.compile(r"-\(?(-?[\d.]+)\)?,2\)$")


@dataclass
class Target:
    header_row: int
    purpose: int
    first_row: int
    last_filled: int
    sum_row: int | None
    num: int | None = None
    doc: int | None = None
    date: int | None = None
    rash: int | None = None


def locate(ws, header_row: int, purpose_col: int, user_cols: dict | None = None) -> Target:
    """Таблица на листе. user_cols — столбцы, указанные пользователем (с 0): num, doc, date."""
    _, last_col = used_bounds(ws)
    head = row_values(ws, header_row, last_col)
    found = target_cols(head)
    found.update({k: v for k, v in (user_cols or {}).items() if k != "purpose"})
    one = lambda k: found[k] + 1 if found.get(k) is not None else None
    t = Target(header_row, purpose_col, header_row + 1, header_row, sum_row_below(ws, header_row, purpose_col))
    t.num, t.doc, t.date = one("num"), one("doc"), one("date")
    t.rash = find_col([norm_text(v) for v in head], ("расход",), purpose_col) or ensure_reconciliation(ws, header_row, head)
    end = (t.sum_row - 1) if t.sum_row else max(used_bounds(ws)[0], header_row + 1)
    texts = col_values(ws, purpose_col, header_row + 1, end)
    nums = col_values(ws, t.num, header_row + 1, end) if t.num else [None] * len(texts)
    # первая строка данных: с номером р/т или с назначением (строка под шапкой может быть пустой)
    for i, (txt, n) in enumerate(zip(texts[:5], nums[:5])):
        if to_number(n) is not None or str(txt or "").strip():
            t.first_row = header_row + 1 + i
            break
    filled = [header_row + 1 + i for i, v in enumerate(texts) if str(v or "").strip()]
    t.last_filled = max(filled) if filled else t.first_row - 1
    return t


def ensure_reconciliation(ws, header_row: int, head: list) -> int:
    """Нет «Расходы» — добавить «Ҳамагӣ» и «Расходы» справа от шапки (туда пойдут суммы из выписки,
    на них держится сверка). → номер столбца «Расходы»."""
    pos = max((c for c, v in enumerate(head, 1) if v not in (None, "")), default=0) + 1
    insert_column(ws, header_row, pos, "Ҳамагӣ")
    insert_column(ws, header_row, pos + 1, "Расходы")
    return pos + 1


def _amount(formula, value):
    """Сумма из выписки в «Расходы»: число или =ROUND(Ҳамагӣ-сумма,2)."""
    m = BALANCE_AMOUNT_RE.search(str(formula or ""))
    return round(float(m.group(1)), 2) if m else (round(to_number(value), 2) if to_number(value) is not None else None)


def payment_key(purpose, amount, doc=None):
    return norm_text(purpose), amount, doc_text(doc)


def split_new(payments, seen: Counter, use_doc: bool) -> tuple:
    """Платежи, которых ещё нет в файле, и сколько пропущено. Одинаковые платежи
    считаются поштучно: два одинаковых в выписке и один в файле — перенесётся один."""
    seen, new = Counter(seen), []
    for p in payments:
        k = payment_key(p.purpose, p.amount, p.doc if use_doc else None)
        if seen[k] > 0:
            seen[k] -= 1
        else:
            new.append(p)
    return new, len(payments) - len(new)


def existing_keys(ws, t: Target, use_doc: bool) -> Counter:
    """Что уже перенесено раньше: назначение + сумма (+ № документа)."""
    if t.last_filled < t.first_row:
        return Counter()
    r1, r2 = t.first_row, t.last_filled
    texts = col_values(ws, t.purpose, r1, r2)
    forms = col_values(ws, t.rash, r1, r2, "Formula") if t.rash else [None] * len(texts)
    vals = col_values(ws, t.rash, r1, r2) if t.rash else [None] * len(texts)
    docs = col_values(ws, t.doc, r1, r2) if use_doc else [None] * len(texts)
    return Counter(payment_key(p, _amount(f, v), d) for p, f, v, d in zip(texts, forms, vals, docs)
                   if str(p or "").strip())


def make_room(ws, t: Target, count: int) -> int:
    """Свободные строки сразу после последней заполненной; не хватает до «Чамъ» — вставляются.
    Возвращает первую строку для записи."""
    start = t.last_filled + 1
    if not t.sum_row:
        return start
    need = start + count - t.sum_row
    if need > 0:
        ws.Range(ws.Rows(t.sum_row), ws.Rows(t.sum_row + need - 1)).Insert(Shift=XL_DOWN, CopyOrigin=0)
        t.sum_row += need
        extend_sum_row(ws, t)
    return start


NEW_HEADERS = ("р/т", "Номгӯи корхонаҳо", "№ ҳуҷҷат", "Сана", "Ҳамагӣ", "Расходы")
NEW_WIDTHS = (6, 60, 12, 12, 14, 14)


def create_table(ws) -> tuple:
    """Пустой лист → таблица как в Бучете: шапка (строка 1), одна пустая строка, «Чамъ».
    Возвращает (строка шапки, столбец назначений)."""
    n = len(NEW_HEADERS)
    ws.Range(ws.Cells(1, 1), ws.Cells(1, n)).Value = [NEW_HEADERS]
    ws.Cells(3, 1).Value = "Чамъ"
    for c in (5, 6):
        L = get_column_letter(c)
        ws.Cells(3, c).Formula = f"=SUM({L}2:{L}2)"
    for c, w in enumerate(NEW_WIDTHS, 1):
        ws.Columns(c).ColumnWidth = w
    ws.Columns(2).WrapText = True
    ws.Range(ws.Cells(1, 1), ws.Cells(1, n)).Font.Bold = True
    ws.Range(ws.Cells(3, 1), ws.Cells(3, n)).Font.Bold = True
    ws.Range(ws.Cells(1, 1), ws.Cells(3, n)).Borders.LineStyle = 1
    return 1, 2


def extend_sum_row(ws, t: Target):
    """Итоги «Чамъ» вида =SUM(G6:G9) растягиваются на всю таблицу (с новыми строками)."""
    _, last_col = used_bounds(ws)
    forms = ws.Range(ws.Cells(t.sum_row, 1), ws.Cells(t.sum_row, last_col)).Formula
    for c, f in enumerate(forms[0] if isinstance(forms, tuple) else (forms,), 1):
        m = SIMPLE_SUM_RE.match(str(f or ""))
        if m and m.group(1).upper() == m.group(2).upper():
            L = get_column_letter(c)
            ws.Cells(t.sum_row, c).Formula = f"=SUM({L}{t.first_row}:{L}{t.sum_row - 1})"
