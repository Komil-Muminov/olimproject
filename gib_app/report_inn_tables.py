"""«Всё по ИНН»: строки обычных таблиц (не реестров), где встречается ИНН, — на любом листе, в любом
столбце. Каждое значение строки подписано: файл, лист, строка и заголовок столбца."""
import os
from datetime import date, datetime

from openpyxl.utils import get_column_letter

from analysis_list_map import AS_LIST, split_cell
from compare_finder import find_tables, key_in_text, key_of, match_key
from file_format import real_path

DATE_FMT = "%d.%m.%Y"


def _match(v, target: str) -> bool:
    k = key_of(v) or key_in_text(v)
    return bool(k) and match_key(k) == target


def _text(v) -> str:
    return " ".join(str(v).split()) if v not in (None, "") else ""


def _show(v):
    """Значение для отчёта: дата — текстом, целое число — без «.0»."""
    if isinstance(v, (datetime, date)):
        return v.strftime(DATE_FMT)
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _plain_header(data, r: int) -> int:
    """Строка шапки над строкой r, если таблица не узнана: ближайшая выше, где хотя бы два текста."""
    for h in range(r - 1, 0, -1):
        if sum(isinstance(v, str) and len(v.strip()) > 1 for v in data[h - 1]) >= 2:
            return h
    return 0


def _row_hit(data, r: int, c: int, tables) -> tuple:
    """→ (название, [(столбец, значение)]) — строка r таблицы, в которой ИНН найден в столбце c."""
    side = next((s for s in tables if c == s.key_col or c in s.cols), None)
    header = side.header if side else _plain_header(data, r)
    cols = sorted(set(side.cols) | {side.key_col}) if side else range(1, len(data[r - 1]) + 1)
    row = data[r - 1]
    name = _text(row[side.name_col - 1]) if side and side.name_col else next(
        (_text(v) for v in row if isinstance(v, str) and len(v.strip()) > 3 and not key_of(v)), "")
    cells = []
    for col in cols:
        v = row[col - 1] if col - 1 < len(row) else None
        if v in (None, ""):
            continue
        head = _text(data[header - 1][col - 1]) if header and col - 1 < len(data[header - 1]) else ""
        if not head and header > 1 and col - 1 < len(data[header - 2]):     # шапка в две строки
            head = _text(data[header - 2][col - 1])
        cells.append((head or "(без названия)", col, _show(v)))
    return name, cells


def table_hits(path: str, inn: str) -> list:
    """→ [(лист, строка в Excel, название, [(заголовок, буква столбца, значение)])]."""
    from python_calamine import CalamineWorkbook

    target, out = match_key(inn), []
    wb = CalamineWorkbook.from_path(real_path(path))
    for name in wb.sheet_names:
        sheet = wb.get_sheet_by_name(name)
        row0, col0 = sheet.start or (0, 0)
        data = [[None if v == "" else v for v in r] for r in sheet.to_python()]
        found = [(r, c) for r, row in enumerate(data, 1) for c, v in enumerate(row, 1) if _match(v, target)]
        if not found:
            continue
        tables = (find_tables(data) or (0, []))[1]
        for r in dict.fromkeys(r for r, _ in found):            # строка — один раз, даже если ИНН в ней дважды
            c = next(c for rr, c in found if rr == r)
            title, cells = _row_hit(data, r, c, tables)
            out.append((name, row0 + r, title, [(h, get_column_letter(col0 + col), v) for h, col, v in cells]))
    return out


def chosen_hits(path: str, inn: str, sheet_maps: list) -> list:
    """Листы-списки, настроенные пользователем: ИНН — только в выбранном столбце, строка — целиком, заголовки —
    из выбранной строки заголовков. → как table_hits."""
    from python_calamine import CalamineWorkbook

    target, out = match_key(inn), []
    wb = CalamineWorkbook.from_path(real_path(path))
    for m in sheet_maps:
        c, name_c = m["cols"].get("inn"), m["cols"].get("name")
        if c is None or m.get("transpose"):
            continue
        sheet = wb.get_sheet_by_name(m["sheet"])
        rows = list(sheet.iter_rows())          # строки — с первой строки листа, столбцы — с первого заполненного
        heads = rows[m["head"]] if m["head"] < len(rows) else []
        for r in range(m["head"] + 1, len(rows)):
            row = rows[r]
            if c >= len(row) or not any(_match(v, target) for v in split_cell(row[c])):
                continue
            cells = [((_text(heads[i]) if i < len(heads) else "") or "(без названия)",
                      get_column_letter(m["col0"] + i + 1), _show(v)) for i, v in enumerate(row) if v not in (None, "")]
            out.append((m["sheet"], r + 1, _text(row[name_c]) if name_c is not None and name_c < len(row) else "", cells))
    return out


def tables_report_rows(paths: list, inn: str, skip: set, progress=lambda p, t: None, mappings=None) -> tuple:
    """Обычные таблицы из paths (кроме реестров skip). Листы, настроенные пользователем (mappings), — только
    они и только выбранный столбец ИНН. → (строки листа «Из таблиц», строки «Где нашлось»)."""
    rows, where, mappings = [], [], mappings or {}
    for k, path in enumerate(paths):
        chosen = mappings.get(path)
        lists = [m for m in chosen if m["kind"] == AS_LIST] if isinstance(chosen, list) else None
        if (lists is None and path in skip) or lists == []:
            continue
        fname = os.path.basename(path)
        progress(60 + int(25 * k / len(paths)), f"Ищу в {fname}…")
        try:
            hits = chosen_hits(path, inn, lists) if lists else table_hits(path, inn)
        except Exception:
            continue
        for sheet, r, title, cells in hits:
            where.append([fname, sheet, r, "таблица", title, len(cells), None])
            rows += [[len(rows) + i, fname, sheet, r, f"{letter}: {head}", v]
                     for i, (head, letter, v) in enumerate(cells, 1)]
    return rows, where
