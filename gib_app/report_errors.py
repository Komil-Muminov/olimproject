"""Листы ошибок: «Ошибки в реестрах» и «Ошибки в таблицах» (только если есть). Ошибки по группам —
у каждой группы цветной заголовок (наведите — подсказка), пояснение и одинаковый порядок столбцов."""
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from report_fast import DATE, title_rows
from report_styles import BORDER, CENTER, NUM, fill_color

PAY_HEAD = ["Файл", "Строка", "№ док.", "Дата", "Сумма"]
# группа: (заголовок, цвет, что это и как проверить, шапка)
GROUPS = {
    "dup": ("ДВОЙНОЙ ПЛАТЁЖ", "B91C1C",
            "Один и тот же платёж записан дважды: совпадают №, дата, плательщик, получатель, сумма и основание. "
            "Проверьте, не провели ли платёж два раза. «Повтор строки» — где этот платёж уже был.",
            PAY_HEAD + ["Повтор строки", "Плательщик", "ИНН плательщика", "Получатель", "ИНН получателя",
                        "Основание"]),
    "bad_pay": ("НЕПРАВИЛЬНЫЙ ИНН", "EA580C",
                "ИНН должен быть из 9 цифр. «ИНН в названии» — если в названии организации вписан ИНН, скорее "
                "всего правильный именно он.",
                PAY_HEAD + ["Чей ИНН", "Организация", "Что стоит", "Что не так", "ИНН в названии"]),
    "odd": ("СТРАННАЯ ДАТА", "64748B",
            "Дата платёжки больше чем на месяц раньше или позже периода файла (например, 2004 или 2029 год). "
            "Возможна опечатка в дате.",
            PAY_HEAD + ["Период файла", "Плательщик", "ИНН плательщика", "Получатель", "ИНН получателя"]),
    "dup_row": ("ПОВТОР ИНН В ТАБЛИЦЕ", "7C3AED",
                "Один ИНН записан в одной таблице несколько раз — возможно, организация задвоена.",
                ["Файл", "Таблица", "Строки", "ИНН", "Организация", "Сколько раз"]),
    "bad_list": ("НЕПРАВИЛЬНЫЙ ИНН", "EA580C",
                 "ИНН в таблице выглядит неправильно. «Как должно быть» — если программа может подсказать.",
                 ["Файл", "Таблица", "Строка", "Что стоит", "Организация", "Что не так", "Как должно быть"]),
}
# лист: (группы, ширины столбцов с «№», столбцы сумм, столбцы дат — с 1, без «№»)
SHEETS = {
    "Ошибки в реестрах": (("dup", "bad_pay", "odd"), [6, 13, 10, 16, 12, 16, 16, 34, 16, 34, 16, 50], (5,), (4,)),
    "Ошибки в таблицах": (("dup_row", "bad_list"), [6, 16, 40, 16, 16, 44, 36, 16], (), ()),
}
BIGGEST_FIRST = ("dup", "bad_pay")      # крупные суммы сверху


def _cell(ws, v, **style):
    c = WriteOnlyCell(ws, value=v)
    for k, val in style.items():
        setattr(c, k, val)
    return c


def write_errors_sheets(wb, details: dict) -> list:
    """→ названия созданных листов (если ошибок нет — один лист «Ошибки» с «не найдено»)."""
    made = []
    for name, (groups, widths, money, dates) in SHEETS.items():
        if any(details.get(g) for g in groups):
            _sheet(wb, name, [(g, details[g]) for g in groups if details.get(g)], widths, money, dates)
            made.append(name)
    if not made:
        ws = wb.create_sheet("Ошибки")
        title_rows(ws, "ОШИБКИ", "16A34A", "Ошибок не найдено ✓", 4)
        made.append("Ошибки")
    return made


def _sheet(wb, name, groups, widths, money, dates):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = "C2410C"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    title_rows(ws, f"{name.upper()} — {sum(len(r) for _, r in groups)}", "C2410C",
               "Ошибки по группам. В каждой группе: что это, как проверить и строки, где найдено.", len(widths))
    for g, rows in groups:
        title, color, hint, head = GROUPS[g]
        if g in BIGGEST_FIRST:
            rows = sorted(rows, key=lambda r: -(r[4] or 0))
        band = fill_color(color)
        t = _cell(ws, f"{title} — {len(rows)}", font=Font(bold=True, size=13, color="FFFFFF"), fill=band,
                  alignment=Alignment(vertical="center", indent=1))
        t.comment = Comment(hint, "Распределение по ГИБ", width=320, height=90)
        ws.append([t] + [_cell(ws, None, fill=band) for _ in head])
        ws.append([_cell(ws, hint, font=Font(italic=True, color="475569"))])
        ws.append([_cell(ws, h, font=Font(bold=True), fill=fill_color("E2E8F0"), border=BORDER, alignment=CENTER)
                   for h in ["№"] + head])
        for k, r in enumerate(rows, 1):
            ws.append([k] + [_cell(ws, v, number_format=NUM) if i in money and v is not None
                             else _cell(ws, v, number_format=DATE) if i in dates and v is not None
                             else v for i, v in enumerate(r, 1)])
        ws.append([])
