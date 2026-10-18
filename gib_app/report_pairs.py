"""Лист «Кто кому платит»: все пары плательщик → получатель из реестров — суммы по частям."""
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from report_fast import title_rows
from report_styles import BORDER, CENTER, NUM, fill_color

HEAD_ROW = 4
HINTS = {"Плательщик": "Кто платил (самое частое написание названия).",
         "Получатель": "Кому платили.",
         "Всего": "Сумма всех платежей этого плательщика этому получателю.",
         "Платежей": "Сколько было платежей."}


def _cell(ws, v, **style):
    c = WriteOnlyCell(ws, value=v)
    for k, val in style.items():
        setattr(c, k, val)
    return c


def write_pairs_sheet(wb, data):
    """Только для реестров (у обычных таблиц нет пар). Крупные связки сверху."""
    if not data.pairs:
        return None
    parts = data.parts
    ws = wb.create_sheet("Кто кому платит")
    ws.sheet_properties.tabColor = "2563EB"
    ws.sheet_view.showGridLines = False
    head = ["№", "Плательщик", "ИНН плательщика", "Получатель", "ИНН получателя"] + [p.label for p in parts] + \
        ["Всего", "Платежей"]
    widths = [7, 40, 13, 40, 13] + [17] * len(parts) + [18, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"C{HEAD_ROW + 1}"
    title_rows(ws, f"КТО КОМУ ПЛАТИТ — {len(data.pairs):,}".replace(",", " "), "2563EB",
               "Каждая строка — плательщик → получатель: сколько он ему заплатил в каждой части.", len(head))
    hints = dict(HINTS)
    hints.update({p.label: f"Сумма платежей этой связки за «{p.label}»." for p in parts})
    row = []
    for t in head:
        c = _cell(ws, t, font=Font(bold=True), fill=fill_color("E2E8F0"), border=BORDER, alignment=CENTER)
        if hints.get(t):
            c.comment = Comment(hints[t], "Распределение по ГИБ", width=320, height=100)
        row.append(c)
    ws.append(row)
    rows = sorted(data.pairs, key=lambda x: -sum(x.sums))
    for k, x in enumerate(rows, 1):
        ws.append([k, x.payer, x.payer_inn, x.rec, x.rec_inn]
                  + [_cell(ws, x.sums[i], number_format=NUM) if x.counts[i] else None for i in range(len(parts))]
                  + [_cell(ws, round(sum(x.sums), 2), number_format=NUM), sum(x.counts)])
    last = HEAD_ROW + len(rows)
    ws.auto_filter.ref = f"A{HEAD_ROW}:{get_column_letter(len(head))}{last}"
    return ws
