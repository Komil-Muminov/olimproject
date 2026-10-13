"""Форматирование текста, предупреждений и сводок для пользовательского интерфейса."""
import os
import sys
from openpyxl.utils import get_column_letter

from utils import format_money


def rows_list(rows: list, limit=8) -> str:
    shown = ", ".join(str(r) for r in rows[:limit])
    return ("строка " if len(rows) == 1 else "строки ") + shown + (
        f" и ещё {len(rows) - limit}" if len(rows) > limit else "")


def template_path() -> str:
    base = getattr(sys, "_MEIPASS", os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
    return os.path.join(base, "assets", "gib_template.xlsm")


def format_warnings(plan, after=False) -> str:
    w = []
    if plan.no_gib:
        w.append(f"⚠ Нет ГИБ в назначении — {rows_list(plan.no_gib)}. "
                 + ("В Excel эти ячейки красные — заполните вручную." if after
                    else "Программа это пропустит и выделит красным."))
    if plan.mismatch:
        w.append(f"⚠ Сумма ГИБ не равна сумме платежа — {rows_list(plan.mismatch)}. Проверьте.")
    if plan.long_codes:
        shown = ", ".join(f"строка {r} ({', '.join(f'{c} → {c[:4]}' for c in codes)})"
                          for r, codes in plan.long_codes[:6])
        more = f" и ещё {len(plan.long_codes) - 6}" if len(plan.long_codes) > 6 else ""
        w.append(f"⚠ В ГИБ больше 4 цифр — взяты первые 4: {shown}{more}. Проверьте"
                 + (" (в Excel эти строки жёлтые)." if after else "."))
    return "<br>".join(w)


def format_plan_description(plan) -> tuple[str, bool]:
    lay = plan.layout
    if not lay and plan.header_only:
        h = plan.header_only
        return (
            f"Назначений с ГИБ в файле пока нет.<br><br>"
            f"Можно <b>встроить автоматику</b>: после этого ГИБ будут проставляться сами, когда вы "
            f"вставите назначения в столбец {get_column_letter(h.purpose_col)} "
            f"(лист «{plan.sheet}», заголовок в строке {h.header_row}).", False
        )
    if not lay:
        return "В этом файле не нашлось назначений с «ГИБ XXXX-сумма».", False
    if not plan.items:
        return (f"Все суммы ГИБ уже проставлены ✓<br><br>"
                f"<span style='color:#64748B'>Лист «{plan.sheet}», "
                f"назначений с ГИБ: {len(plan.done)}</span>", False)

    parts = [f"Найдено назначений с ГИБ: <b>{len(plan.items)}</b> на сумму <b>{format_money(plan.total)}</b>",
             f"ГИБ: <b>{', '.join(plan.codes)}</b>"]
    if plan.new_codes:
        parts.append(f"Будут созданы столбцы: <b>{', '.join(plan.new_codes)}</b>")
    else:
        parts.append("Все нужные столбцы ГИБ уже есть")
    if plan.done:
        parts.append(f"<span style='color:#64748B'>Уже заполнено раньше: {len(plan.done)} — их не трогаем</span>")
    parts.append("Сверка с выпиской — в столбце «Расходы» сразу за «Ҳамагӣ»: 0 — сумма совпала")
    parts.append(f"<span style='color:#64748B'>Лист «{plan.sheet}», назначения в столбце "
                 f"{get_column_letter(lay.purpose_col)}, строки {lay.first_row}–{lay.last_row}</span>")
    return "<br>".join(parts), True


def format_done_summary(res: dict, plan) -> tuple[list, list]:
    check = res["check"]
    lines = [f"Заполнено строк: <b>{res['filled']}</b> на сумму <b>{format_money(check['written_sum'])}</b>"]
    if res["created"]:
        lines.append(f"Созданы столбцы ГИБ: <b>{', '.join(res['created'])}</b>")
    if res.get("header_inserted"):
        lines.append("Над таблицей добавлена строка с номерами ГИБ")
    lines.append("Всё, что вписала программа, подсвечено <span style='background:#DCFCE7'>&nbsp;зелёным&nbsp;</span>.")
    if res.get("left_open"):
        lines.append("Файл был открыт в Excel — изменения уже в нём и сохранены.")

    recon = check.get("recon")
    if recon and recon["empty"]:
        lines.append("📋 Сверка: вставьте суммы из выписки в столбец «Расходы» и снова нажмите «Заполнить ГИБ» — "
                     "в ячейках появится баланс (0 — сумма совпала).")
    elif recon and not recon["bad_rows"]:
        lines.append("✓ Сверка с выпиской: во всех строках «Расходы» баланс 0 — суммы совпали.")

    warn = [format_warnings(plan, after=True)] if format_warnings(plan) else []
    if recon and recon["bad_rows"] and not recon["empty"]:
        warn.append(f"⚠ Сверка с выпиской: не сходится — {rows_list(recon['bad_rows'])}, "
                    f"общий баланс {format_money(recon['total_diff'])}. В столбце «Расходы» такие ячейки красные.")
    bad = [g for g in check["gibs"] if not (g["written_ok"] and g["total_ok"])]
    for g in bad:
        warn.append(f"⚠ ГИБ {g['code']}: итог «{format_money(g['total'] or 0)}» не равен сумме столбца {format_money(g['column_sum'])}")
    return lines, warn


def format_embed_summary(res: dict, plan) -> tuple[list, str]:
    lines = [f"Создан файл: <b>{os.path.basename(res['path'])}</b>"]
    if plan.items:
        lines.append(f"Сразу заполнено назначений: <b>{len(plan.items)}</b>"
                     + (f", созданы столбцы: <b>{', '.join(plan.new_codes)}</b>" if plan.new_codes else ""))
    lines.append("Сверка с выпиской — столбец «Расходы» сразу за «Ҳамагӣ».")
    lines += ["", "<b>Как работать дальше:</b>",
              "1. Откройте этот файл в Excel и нажмите «Включить содержимое» на жёлтой панели.",
              "2. Вставляйте назначения в столбец назначений — ГИБ проставятся сами.",
              "   Суммы из выписки вставляйте в «Расходы» — там появится баланс (0 — сумма совпала).",
              "3. Кнопки вверху листа: «Пересчитать ГИБ», «Добавить строки»,",
              "   «Зафиксировать распределение» и «Снять фиксацию»."]
    warn = "Исходный файл не изменён." + (
        "<br>" + format_warnings(plan, after=True).replace("Пропущено, заполните вручную.",
                                                           "В Excel эти строки подсвечены красным.")
        if plan.no_gib else "")
    return lines, warn
