"""«Сводная таблица»: отдельный файл — лист «Данные» (одна строка = организация × часть) и настоящая
сводная таблица Excel со срезами на листе «Сводная» (строит excel_pivot.build_pivot)."""
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter

from analysis_model import FLAGS, status
from excel_pivot import build_pivot
from report_styles import NUM

HEAD = ["ИНН", "Название", "Кто", "Часть", "Сумма", "Записей", "Что случилось", "Ошибки"]
WIDTHS = [13, 46, 13, 26, 18, 10, 18, 30]
NO_ERRORS = "нет ошибок"


def pivot_rows(data) -> list:
    """Одна строка на организацию в каждой части, где она есть. Части пронумерованы («1. Янв–Апр 2024»),
    чтобы в сводной шли по порядку времени."""
    labels = [f"{i}. {p.label}" for i, p in enumerate(data.parts, 1)]
    out = []
    for e in data.entities:
        st = status(e, data.parts, data.compare) or "—"
        errors = ", ".join(f for f in FLAGS if f in e.flags) or NO_ERRORS
        for i, p in enumerate(data.parts):
            if e.counts[i]:
                out.append([e.inn, e.name, e.role, labels[i], e.sums[i] if p.has_sums else None, e.counts[i],
                            st, errors])
    return out


def write_pivot_report(path, data, progress=lambda p, t: None) -> dict:
    progress(10, "Готовлю данные для сводной…")
    rows = pivot_rows(data)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Данные")
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.append(HEAD)
    for r in rows:
        if r[4] is not None:
            c = WriteOnlyCell(ws, value=r[4])
            c.number_format = NUM
            r = r[:4] + [c] + r[5:]
        ws.append(r)
    progress(35, "Сохраняю данные…")
    wb.save(path)
    with_sums = any(p.has_sums for p in data.parts)
    build_pivot(path, len(rows) + 1, len(HEAD), with_sums, len(data.roles) > 1, progress)
    progress(100, "Готово")
    return {"path": path, "rows": len(rows), "orgs": len(data.entities), "parts": len(data.parts)}
