"""Перед глубоким анализом: что программа думает о каждом файле — листы (первые строки) и на каждом листе
колонка ИНН и название, найденные по содержимому. Решает пользователь — догадка подставляется в окно
«Листы и ИНН», не нашлось — пользователь вписывает сам или берёт лист как есть."""
from openpyxl.utils import get_column_letter

from analysis_list_map import AS_LIST
from analysis_lists import money_cols
from compare_finder import find_tables
from excel_sheet_info import sheet_scan, workbook_sheets
from file_format import real_path

REPEAT_SHARE = 0.2          # столько ИНН встречаются в нескольких строках — похоже на операции (выписку)
PREVIEW_ROWS = 300          # столько первых строк листа показывается в окне
FULL_READ_LIMIT = 200_000   # таблицы больше — без поиска ИНН по содержимому (долго)


def _sheet_rows(sheet) -> list:
    rows = []
    for i, values in enumerate(sheet.iter_rows()):
        rows.append(["" if v is None else v for v in values])
        if i + 1 >= PREVIEW_ROWS:
            break
    return rows


def _list_guess(wb, name: str) -> dict | None:
    """Список: таблица с ИНН по содержимому (ИНН, название, суммы — как нашла бы программа)."""
    sheet = wb.get_sheet_by_name(name)
    data = [[None if v == "" else v for v in r] for r in sheet.to_python()]     # с первой заполненной строки
    found = find_tables(data)
    if not found:
        return None
    s = found[1][0]
    return {"sheet": name, "kind": AS_LIST, "head": (sheet.start or (0, 0))[0] + s.header - 1,
            "cols": {"inn": s.key_col - 1, "name": s.name_col - 1 if s.name_col else None},
            "sums": [c - 1 for c in money_cols(data, s)],
            "repeats": sum(len(r) > 1 for r in s.all_rows.values()) > REPEAT_SHARE * len(s.all_rows)}


def sheet_inn(wb, name: str, col0: int) -> dict:
    """Колонка ИНН и название на листе — буквами («C», «B»), как их вписал бы пользователь. Не нашлось — None."""
    g = _list_guess(wb, name)
    if not g:
        return {"inn": None, "name": None, "repeats": False}
    name_col = g["cols"]["name"]
    return {"inn": get_column_letter(col0 + g["cols"]["inn"] + 1),
            "name": get_column_letter(col0 + name_col + 1) if name_col is not None else None, "repeats": g["repeats"]}


def propose(path: str) -> dict:
    """→ {"path", "sheets": {лист: {"rows", "row0", "col0", "height", "inn", "name"}}} (пустые листы пропущены)."""
    from python_calamine import CalamineWorkbook

    wb, info = CalamineWorkbook.from_path(real_path(path)), workbook_sheets(path)
    sheets = {}
    for name in wb.sheet_names:
        sheet = wb.get_sheet_by_name(name)
        rows = _sheet_rows(sheet)
        if any(v != "" for r in rows for v in r):
            # iter_rows: строки — с первой строки листа (row0 = 0), столбцы — с первого заполненного (col0)
            formulas = sheet_scan(path, name, formulas=True)["formulas"] if sheet.height <= FULL_READ_LIMIT else {}
            sheets[name] = {"rows": rows, "row0": 0, "col0": (sheet.start or (0, 0))[1], "height": sheet.height,
                            "hidden": info.get(name, {}).get("hidden", False),
                            "formulas": {c - (sheet.start or (0, 0))[1]: n for c, n in formulas.items()}}
    for name, sh in sheets.items():             # колонка ИНН на каждом листе — для окна «Листы и ИНН»
        sh.update(sheet_inn(wb, name, sh["col0"]) if sh["height"] <= FULL_READ_LIMIT else {"inn": None, "name": None})
    return {"path": path, "sheets": sheets}


def propose_files(paths: list, progress=lambda p, t: None) -> list:
    out = []
    for k, path in enumerate(paths):
        progress(int(100 * k / len(paths)), f"Смотрю файл {k + 1} из {len(paths)}…")
        out.append(propose(path))
    return out
