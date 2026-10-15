"""Что не видно в значениях ячеек: скрытые листы и строки, формулы без сохранённого результата. Читается прямо
из .xlsx / .xlsm (это zip с xml); для старого .xls таких сведений нет — пусто."""
import re
import zipfile

from file_format import real_path

SHEET_RE = re.compile(r'<sheet\b[^>]*>')
ROW_RE = re.compile(r'<row\b[^>]*>')
CELL_RE = re.compile(r'<c\b[^>]*?\br="([A-Z]+)\d+"[^>]*?(?:/>|>(.*?)</c>)', re.S)
VALUE_RE = re.compile(r"<v>[^<]")
FORMULA_SCAN_LIMIT = 20_000_000     # большие листы (реестры) — без поиска формул (долго)


def _attr(tag: str, name: str) -> str | None:
    m = re.search(rf'\b{name}="([^"]*)"', tag)
    return m.group(1) if m else None


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n - 1


def workbook_sheets(path: str) -> dict:
    """{лист: {"hidden": скрыт ли, "xml": путь к листу внутри файла}}; не .xlsx — {}."""
    path = real_path(path)
    if not zipfile.is_zipfile(path):
        return {}
    with zipfile.ZipFile(path) as z:
        if "xl/workbook.xml" not in z.namelist():         # не книга Excel нового формата (например, .xls)
            return {}
        book = z.read("xl/workbook.xml").decode("utf-8", "ignore")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8", "ignore")
    targets = {_attr(t, "Id"): _attr(t, "Target") for t in re.findall(r"<Relationship\b[^>]*>", rels)}
    out = {}
    for tag in SHEET_RE.findall(book):
        target = (targets.get(_attr(tag, "r:id")) or "").lstrip("/")
        name = _attr(tag, "name").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
        out[name] = {"hidden": _attr(tag, "state") in ("hidden", "veryHidden"),
                     "xml": target if target.startswith("xl/") else "xl/" + target}
    return out


def sheet_scan(path: str, sheet: str, formulas: bool = False) -> dict:
    """{"hidden_rows": {номера строк Excel}, "formulas": {столбец (с 0): сколько формул без значения}}."""
    info = workbook_sheets(path).get(sheet)
    if not info:
        return {"hidden_rows": set(), "formulas": {}}
    with zipfile.ZipFile(real_path(path)) as z:
        size = z.getinfo(info["xml"]).file_size
        xml = z.read(info["xml"]).decode("utf-8", "ignore")
    hidden = {int(_attr(t, "r")) for t in ROW_RE.findall(xml) if _attr(t, "hidden") in ("1", "true") and _attr(t, "r")}
    empty = {}
    if formulas and size <= FORMULA_SCAN_LIMIT:
        for col, body in CELL_RE.findall(xml):
            if body and "<f" in body and not VALUE_RE.search(body):      # формула, а результата нет
                empty[_col_index(col)] = empty.get(_col_index(col), 0) + 1
    return {"hidden_rows": hidden, "formulas": empty}


HIDDEN_LIST_LIMIT = 2000    # скрытых строк больше — в отчёте только сводка (иначе это весь файл, например фильтр)


def hidden_note(path: str, sheet: str, hidden_rows: set) -> dict | None:
    """Пометка для «Как считали» / экрана итога: скрытый лист и скрытые строки. → None, если скрытого нет."""
    info = workbook_sheets(path).get(sheet, {})
    if not info.get("hidden") and not hidden_rows:
        return None
    return {"file": path, "sheet": sheet, "hidden_sheet": bool(info.get("hidden")), "hidden_rows": len(hidden_rows),
            "listed": len(hidden_rows) <= HIDDEN_LIST_LIMIT}


def cell_value(path: str, sheet: str, col: int, row: int):
    """Значение ячейки (столбец с 0, строка Excel с 1) — для сверки с «Итого» в файле."""
    from python_calamine import CalamineWorkbook

    sh = CalamineWorkbook.from_path(real_path(path)).get_sheet_by_name(sheet)
    c = col - (sh.start or (0, 0))[1]
    for i, values in enumerate(sh.iter_rows()):     # строки — с первой строки листа
        if i == row - 1:
            return values[c] if 0 <= c < len(values) and values[c] != "" else None
    return None


def total_check(path: str, m: dict, program: float) -> dict | None:
    """Сверка: итог в ячейке файла (её вписал пользователь) и сумма, которую насчитала программа."""
    if not m.get("total_cell"):
        return None
    from openpyxl.utils import get_column_letter
    from utils import to_number

    col, row = m["total_cell"]
    raw = cell_value(path, m["sheet"], col, row)
    in_file = to_number(raw)
    return {"file": path, "sheet": m["sheet"], "cell": f"{get_column_letter(col + 1)}{row}", "raw": raw,
            "in_file": in_file, "program": round(program, 2),
            "diff": None if in_file is None else round(program - in_file, 2)}


WHITE_RE = re.compile(r'<color\b[^>]*\b(?:rgb="(?:[0-9A-F]{2})?FFFFFF"|theme="0")[^>]*/>', re.I)   # «AARRGGBB»
CELL_TAG_RE = re.compile(r"<c\b([^>]*?)(?<!/)>(.*?)</c>", re.S)        # пустые <c …/> пропускаются


def _items(xml: str, block: str, item: str) -> list:
    """Элементы <item> внутри <block> (fonts → font, fills → fill, cellXfs → xf)."""
    m = re.search(rf"<{block}\b.*?</{block}>", xml, re.S)
    return re.findall(rf"<{item}\b[^>]*/>|<{item}\b.*?</{item}>", m.group(0), re.S) if m else []


def _white_styles(z) -> set:
    """Номера стилей ячеек, где шрифт белый, а заливки нет (или она белая) — такой текст в Excel не видно."""
    styles = z.read("xl/styles.xml").decode("utf-8", "ignore")
    white_font = {i for i, f in enumerate(_items(styles, "fonts", "font")) if WHITE_RE.search(f)}
    clear_fill = {i for i, f in enumerate(_items(styles, "fills", "fill"))
                  if 'patternType="solid"' not in f or re.search(r'<fgColor\b[^>]*rgb="(?:[0-9A-F]{2})?FFFFFF"', f, re.I)}
    return {i for i, xf in enumerate(_items(styles, "cellXfs", "xf"))
            if int(_attr(xf, "fontId") or 0) in white_font and int(_attr(xf, "fillId") or 0) in clear_fill}


def white_cells(path: str, sheet: str) -> list:
    """[(ячейка, текст)] — непустые ячейки белым шрифтом (большие листы не проверяются — долго)."""
    info = workbook_sheets(path).get(sheet)
    if not info:
        return []
    with zipfile.ZipFile(real_path(path)) as z:
        names = z.namelist()
        if "xl/styles.xml" not in names or z.getinfo(info["xml"]).file_size > FORMULA_SCAN_LIMIT:
            return []
        white = _white_styles(z)
        if not white:
            return []
        cells = [(attrs, body) for attrs, body in CELL_TAG_RE.findall(z.read(info["xml"]).decode("utf-8", "ignore"))
                 if _attr(attrs, "s") and int(_attr(attrs, "s")) in white]
        shared = re.findall(r"<si>(.*?)</si>", z.read("xl/sharedStrings.xml").decode("utf-8", "ignore"), re.S) \
            if cells and "xl/sharedStrings.xml" in names else []
    out = []
    for attrs, body in cells:
        v = re.search(r"<v>([^<]*)</v>", body)
        text = v.group(1) if v else "".join(re.findall(r"<t[^>]*>([^<]*)</t>", body))
        if _attr(attrs, "t") == "s" and text.isdigit() and int(text) < len(shared):
            text = "".join(re.findall(r"<t[^>]*>([^<]*)</t>", shared[int(text)]))
        if text.strip():
            out.append((_attr(attrs, "r"), text))
    return out
