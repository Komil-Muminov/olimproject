"""Генерация листов категорий для отчёта сравнения списков."""
from openpyxl.styles import Alignment, Border, Font
from openpyxl.utils import get_column_letter

from report_styles import (
    BORDER, CENTER, COLORS, MED, NUM, THIN,
    apply_diff_colors, fill_color, styled_cell
)


def create_category_sheet(wb, sheet_name, title, hint, key, items, metrics, y1, y2,
                          name_title="Организация", note_title="Примечание", note=lambda it: ""):
    """Лист категории: цветной заголовок, пояснение, таблица, «Итого», фильтр.
    items: (ИНН, название, [(было, стало)] по показателям). → (лист, строка «Итого» или None)."""
    color, light = COLORS[key]
    ws = wb.create_sheet(sheet_name)
    ws.sheet_properties.tabColor = color
    ws.sheet_view.showGridLines = False
    first_val = 4
    n_cols = 3 + 3 * len(metrics) + 1
    note_col = n_cols
    for i, w in enumerate([6, 13, 50] + [16, 16, 15] * len(metrics) + [50], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    styled_cell(ws, 1, 1, f"{title} — {len(items)}", font=Font(bold=True, size=14, color="FFFFFF"),
                fill=fill_color(color), alignment=Alignment(vertical="center", indent=1))
    for c in range(2, n_cols + 1):
        ws.cell(1, c).fill = fill_color(color)
    ws.row_dimensions[1].height = 30
    styled_cell(ws, 2, 1, hint, font=Font(italic=True, color="475569"))

    h1, h2 = 4, 5
    for c, t in ((1, "№"), (2, "ИНН / РМА"), (3, name_title), (note_col, note_title)):
        ws.merge_cells(start_row=h1, start_column=c, end_row=h2, end_column=c)
        styled_cell(ws, h1, c, t)
    for i, m in enumerate(metrics):
        c = first_val + 3 * i
        ws.merge_cells(start_row=h1, start_column=c, end_row=h1, end_column=c + 2)
        styled_cell(ws, h1, c, m)
        for j, t in enumerate((y1, y2, "Разница")):
            styled_cell(ws, h2, c + j, t)
    for r in (h1, h2):
        for c in range(1, n_cols + 1):
            styled_cell(ws, r, c, font=Font(bold=True, color="0F172A"), fill=fill_color("E2E8F0"),
                        border=BORDER, alignment=CENTER)
    ws.row_dimensions[h1].height = 34
    ws.freeze_panes = ws.cell(h2 + 1, 4)

    row = h2 + 1
    if not items:
        styled_cell(ws, row, 3, "нет", font=Font(italic=True, color="94A3B8"))
        return ws, None
    for i, (inn, name, vals) in enumerate(items, 1):
        styled_cell(ws, row, 1, i, alignment=Alignment(horizontal="center"))
        styled_cell(ws, row, 2, inn)
        styled_cell(ws, row, 3, name)
        for m, (v1, v2) in enumerate(vals):
            c = first_val + 3 * m
            styled_cell(ws, row, c, v1, number_format=NUM)
            styled_cell(ws, row, c + 1, v2, number_format=NUM)
            l1, l2 = get_column_letter(c), get_column_letter(c + 1)
            styled_cell(ws, row, c + 2, f"={l2}{row}-{l1}{row}", number_format=NUM)
        styled_cell(ws, row, note_col, note((inn, name, vals)) or None, font=Font(color="475569"))
        for c in range(1, n_cols + 1):
            ws.cell(row, c).border = BORDER
        row += 1
    last = row - 1
    ws.auto_filter.ref = f"A{h2}:{get_column_letter(n_cols)}{last}"
    for i in range(len(metrics)):
        apply_diff_colors(ws, first_val + 3 * i + 2, h2 + 1, last + 1)
    if not metrics:
        return ws, None
    styled_cell(ws, row, 3, "Итого", font=Font(bold=True))
    for c in range(first_val, first_val + 3 * len(metrics)):
        let = get_column_letter(c)
        styled_cell(ws, row, c, f"=SUM({let}{h2 + 1}:{let}{last})", number_format=NUM, font=Font(bold=True))
    for c in range(1, n_cols + 1):
        styled_cell(ws, row, c, fill=fill_color(light), border=Border(top=MED, bottom=THIN, left=THIN, right=THIN))
    ws.print_title_rows = f"{h1}:{h2}"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws, row
