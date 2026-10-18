"""Лист «Одна орг. во всех файлах»: строки всех таблиц как есть (все колонки), рядом — по ИНН. Таблица — лист файла,
настроенный пользователем (несколько листов одного файла — разные таблицы). Все организации; «Есть в файлах» —
в каких таблицах она встречается. Каждая колонка подписана: «файл: колонка» (у файла несколько листов — «файл · лист: колонка»)."""
import os
from datetime import date, datetime

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from analysis_list_map import NAME_KEY, RAW, org_key, split_cell
from compare_finder import key_in_text, key_of, match_key
from file_format import real_path
from report_fast import _cell, title_rows
from report_styles import BORDER, CENTER, fill_color

def _text(v) -> str:
    return " ".join(str(v).split()) if v not in (None, "") else ""


def _show(v):
    """Значение для отчёта: дата — текстом, целое число — без «.0»."""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


MAX_ROWS = 500_000          # больше — не помещается в память и в разумный Excel: показываются первые
BASE = ["№", "ИНН", "Есть в файлах", "Файлов", "Файл", "Лист", "Строка"]
STRIPE = "F1F5F9"          # серая полоса — каждая вторая организация
SHEET = "Одна орг. во всех файлах"


def _sources(paths: list, mappings: dict) -> list:
    """Листы с ИНН (листы «без ИНН» — отдельно, как есть)."""
    return [(p, m) for p in paths for m in (mappings.get(p) if isinstance(mappings.get(p), list) else [])
            if m["kind"] != RAW]


def _rows(path: str, m: dict):
    """(строка Excel, значения) — строки данных листа по настройке пользователя."""
    from python_calamine import CalamineWorkbook

    sheet = CalamineWorkbook.from_path(real_path(path)).get_sheet_by_name(m["sheet"])
    for i, values in enumerate(sheet.iter_rows()):          # строки — с первой строки листа
        if i <= m["head"]:
            continue
        if m.get("last") and i + 1 > m["last"]:
            return
        yield i + 1, values


def _keys(values, m: dict) -> list:
    by_name = bool(m.get("by_name"))
    c = m["cols"].get("name") if by_name else (m["cols"].get("inn") if m["kind"] == "list" else m["cols"].get("payer_inn"))
    out = []
    for v in split_cell(values[c] if c is not None and c < len(values) else None):
        if by_name:
            out += [NAME_KEY + org_key(v)] if org_key(v) else []
        elif (k := key_of(v) or key_in_text(v)):
            out.append(match_key(k) or k)
    return out


def _collect(sources: list) -> tuple:
    """Два прохода: в каких таблицах какой ИНН, затем все строки.
    → (строки, шапки источников, обрезано ли, {ИНН: номера таблиц})."""
    per_table = [{k for _, values in _rows(path, m) for k in _keys(values, m)} for path, m in sources]
    where = {}
    for src, keys in enumerate(per_table):
        for k in keys:
            where.setdefault(k, []).append(src)
    rows, heads = [], []
    for src, (path, m) in enumerate(sources):
        heads.append(_head_row(path, m))
        for r, values in _rows(path, m):
            for k in _keys(values, m):
                if len(rows) >= MAX_ROWS:
                    return rows, heads + [[]] * (len(sources) - len(heads)), True, where
                rows.append((k, src, r, [_show(v) if v != "" else None for v in values]))
    return rows, heads, False, where


def _table_name(path: str, m: dict, sources: list) -> str:
    """Подпись таблицы: имя файла; у файла несколько листов — ещё и лист."""
    name = os.path.basename(path)
    return f"{name} · {m['sheet']}" if sum(p == path for p, _ in sources) > 1 else name


def _head_row(path: str, m: dict) -> list:
    head = m.get("labels_row", m["head"])         # строка с названиями колонок (над итогами, если они есть)
    if head < 0:
        return []
    from python_calamine import CalamineWorkbook

    sheet = CalamineWorkbook.from_path(real_path(path)).get_sheet_by_name(m["sheet"])
    for i, values in enumerate(sheet.iter_rows()):
        if i == head:
            return [_text(v) for v in values]
    return []


def _stripes(ws, n_cols: int, n_rows: int):
    """Организации через одну — серым (по чётности «№»): видно, где кончается одна и начинается другая."""
    if n_rows:
        ws.conditional_formatting.add(f"A5:{get_column_letter(n_cols)}{4 + n_rows}",
                                      FormulaRule(formula=["ISEVEN($A5)"], fill=PatternFill("solid", start_color=STRIPE, end_color=STRIPE)))


def write_all_by_inn_sheet(wb, paths: list, mappings: dict) -> dict | None:
    """Если есть хоть один лист с ИНН, указанный пользователем.
    → {"orgs", "multi", "rows", "cut"} или None; multi — сколько организаций есть в 2+ таблицах."""
    sources = _sources(paths, mappings or {})
    if not sources:
        return None
    rows, heads, cut, where = _collect(sources)
    widths = [len(heads[s]) for s in range(len(sources))]
    for _, src, _, values in rows:
        widths[src] = max(widths[src], len(values))
    names = [_table_name(p, m, sources) for p, m in sources]
    ws = wb.create_sheet(SHEET)
    ws.sheet_properties.tabColor = "7C3AED"
    ws.freeze_panes = "C5"
    labels = BASE + [f"{names[s]}: {get_column_letter(m['col0'] + i + 1)}"
                     + (f" {heads[s][i]}" if i < len(heads[s]) and heads[s][i] else "")
                     for s, (p, m) in enumerate(sources) for i in range(widths[s])]
    orgs = sorted({k for k, *_ in rows})
    number = {k: i for i, k in enumerate(orgs, 1)}
    note = f" Показаны первые {MAX_ROWS:,} строк — сузьте строки в поле ИНН.".replace(",", " ") if cut else ""
    title_rows(ws, f"ОДНА ОРГАНИЗАЦИЯ ВО ВСЕХ ФАЙЛАХ — {len(orgs):,} организаций".replace(",", " "),
               "7C3AED", "Здесь собраны все строки по каждой организации из всех файлов. Чтобы посмотреть одну "
               "организацию, поставьте фильтр в колонке «ИНН». «Есть в файлах» — где организация встречается; "
               "«Файлов» = 1 — она есть только в одном файле. Над каждой колонкой написано, из какого она файла."
               + note, len(labels))
    ws.append([_cell(ws, t, font=Font(bold=True), fill=fill_color("E2E8F0"), border=BORDER, alignment=CENTER) for t in labels])
    _stripes(ws, len(labels), len(rows))
    offsets = [len(BASE) + sum(widths[:s]) for s in range(len(sources))]
    for k, src, r, values in sorted(rows, key=lambda x: (x[0], x[1], x[2])):
        path, m = sources[src]
        line = [None] * len(labels)
        line[:len(BASE)] = [number[k], k[1:] if k.startswith(NAME_KEY) else (k.zfill(9) if k.isdigit() and len(k) < 9 else k),
                            ", ".join(names[i] for i in where[k]), len(where[k]), os.path.basename(path), m["sheet"], r]
        line[offsets[src]:offsets[src] + len(values)] = values
        ws.append(line)
    ws.auto_filter.ref = f"A4:{get_column_letter(len(labels))}{4 + len(rows)}"
    return {"orgs": len(orgs), "multi": sum(len(where[k]) > 1 for k in orgs), "rows": len(rows), "cut": cut}
