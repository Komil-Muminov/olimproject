"""Колонки, вписанные от руки в карточке файла → настройка чтения листа (mapping).
Строки берутся из поля ИНН: «C5-C300» — с 5 по 300; «C» — с первой строки, где в колонке C стоит ИНН."""
from openpyxl.utils import get_column_letter

from column_spec import SpecError, one_column, parse
from inn_keys import key_in_text, key_of
from utils import is_total_label, to_number


def _first_data_row(rows: list, col: int) -> int:
    """Номер (с 0) первой строки, где в колонке ИНН стоит ИНН."""
    for i, r in enumerate(rows):
        v = r[col] if col < len(r) else ""
        if key_of(v) or key_in_text(v):
            return i
    letter = get_column_letter(col + 1)
    raise SpecError(f"В колонке {letter} в первых {len(rows)} строках нет ИНН — укажите, с какой строки начинать, "
                    f"например {letter}5")


def labels_row(rows: list, head: int) -> int:
    """Строка с названиями колонок (для подписей): ближайшая над данными, где хотя бы два текста и это не
    «Итого / Ҳамагӣ». Нет такой — строка прямо над данными."""
    for j in range(head, max(head - 6, -1), -1):
        texts = [v for v in rows[j] if isinstance(v, str) and v.strip() and to_number(v) is None]
        if len(texts) >= 2 and not is_total_label(texts[0]):           # строка итогов — «Итого» первым словом
            return j
    return head


def to_mapping(sheet_name: str, sheet: dict, typed: dict) -> dict:
    """typed — что вписано в поля ({"inn": "C5-C300", "sums": "G, H"}). Ошибка ввода → SpecError с понятным текстом.
    → {"sheet", "head", "last", "col0", "cols": {"inn"}, "sums", "labels_row"}; столбцы — с 0 от col0."""
    col0 = sheet["col0"]

    def rel(c: int, title: str) -> int:
        if c < col0:
            raise SpecError(f"{title}: колонка {get_column_letter(c + 1)} на этом листе пустая")
        return c - col0

    inn = one_column(typed.get("inn"), "ИНН")
    if not inn:
        raise SpecError("Укажите колонку ИНН (например, C или C5-C300)")
    sums = [rel(c, "Сумма") for c, _, _ in parse(typed.get("sums"))]
    col, first, last = inn
    head = first - 2 if first else _first_data_row(sheet["rows"], rel(col, "ИНН")) - 1
    return {"sheet": sheet_name, "head": head, "last": last, "col0": col0, "cols": {"inn": rel(col, "ИНН")},
            "sums": sums, "labels_row": labels_row(sheet["rows"], head) if head >= 0 else head}
