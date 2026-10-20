"""Excel глубокого анализа: лист «Анализ» (цвет строки — что случилось, столбец «Ошибки», легенда и
подсказки при наведении) и листы с ошибками подробно."""
from collections import Counter
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from analysis_model import FLAG_HINTS, FLAGS, change, status, status_hints
from report_colors import COLORS, ERR_COLOR, add_row_colors, cf_fill
from report_errors import write_errors_sheets
from report_months import write_months_sheet
from report_all_by_inn import write_all_by_inn_sheet
from report_reconcile import write_reconcile_sheet
from report_settings import write_settings_sheet
from report_threshold import write_threshold_sheet
from report_pairs import write_pairs_sheet
from report_fast import title_rows
from report_styles import BORDER, CENTER, NUM, fill_color

HEAD_ROW = 5
CHANGE_HEAD = ["Изменение", "Изменение, %"]
AUTHOR = "Распределение по ГИБ"


def _errors_cell(ws, e):
    """«Ошибки»: что найдено; если «Итого» в строках не сходится — при наведении где и на сколько (первые 15)."""
    text = ", ".join(f for f in FLAGS if f in e.flags)
    if not e.hover:
        return text
    more = f"\n…и ещё {len(e.hover) - 15}" if len(e.hover) > 15 else ""
    return _cell(ws, text, "«Итого» не сходится:\n" + "\n".join(f"• {h}" for h in e.hover[:15]) + more)


def _name_cell(ws, e):
    """Название; если у ИНН несколько названий — при наведении все варианты (первые 20)."""
    if len(e.variants) < 2:
        return e.name
    more = f"\n…и ещё {len(e.variants) - 20}" if len(e.variants) > 20 else ""
    return _cell(ws, e.name, "Как записано в файлах:\n" + "\n".join(f"• {v}" for v in e.variants[:20]) + more)


def _cell(ws, v, hint=None, **style):
    c = WriteOnlyCell(ws, value=v)
    for k, val in style.items():
        setattr(c, k, val)
    if hint:
        c.comment = Comment(hint, AUTHOR, width=320, height=110)
    return c


def write_analysis_report(path, data, progress=lambda p, t: None, mappings=None, threshold=None) -> dict:
    parts, with_role = data.parts, len(data.roles) > 1
    pair = data.compare
    rows = sorted(((e, status(e, parts, pair) if pair else "") for e in data.entities), key=lambda r: -r[0].total)
    progress(15, "Пишу лист «Анализ»…")
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Анализ")
    ws.sheet_view.showGridLines = False
    head = ["№", "ИНН", "Название"] + (["Кто"] if with_role else []) + [p.label for p in parts] + \
        ["Всего"] + (CHANGE_HEAD if pair else []) + ["Что случилось", "Ошибки", "Названий"]
    n = len(head)
    first_sum = 4 + with_role
    st_col = first_sum + len(parts) + 1 + (2 if pair else 0)
    err_col = st_col + 1
    widths = [7, 13, 46] + [12] * with_role + [17] * len(parts) + [18] + ([18, 12] if pair else []) + [18, 30, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"D{HEAD_ROW + 1}"
    what = f" · «Что случилось» — «{parts[pair[0]].label}» → «{parts[pair[1]].label}»" if pair else ""
    title_rows(ws, "ГЛУБОКИЙ АНАЛИЗ", "0F172A", f"Создано {datetime.now():%d.%m.%Y %H:%M} · части: "
               + " → ".join(f"«{p.label}»" for p in parts) + what, n)
    hints_st = status_hints(parts, pair) if pair else {}
    _legend(ws, hints_st)
    hints = _head_hints(parts, with_role, hints_st, pair)
    ws.append([_cell(ws, t, hints.get(t), font=Font(bold=True), fill=fill_color("E2E8F0"), border=BORDER,
                     alignment=CENTER) for t in head])
    for k, (e, st) in enumerate(rows, 1):
        sums = [(_cell(ws, e.sums[i], number_format=NUM) if p.has_sums else "есть") if e.counts[i] else None
                for i, p in enumerate(parts)]
        ws.append([k, e.inn, _name_cell(ws, e)] + ([e.role] if with_role else []) + sums
                  + [_cell(ws, e.total, number_format=NUM)] + (_change_cells(ws, e, parts, pair) if pair else [])
                  + [st, _errors_cell(ws, e),
                     e.names if e.names > 1 else None])
        if k % 20000 == 0:
            progress(15 + int(60 * k / len(rows)), "Пишу лист «Анализ»…")
    last = HEAD_ROW + len(rows)
    if rows:
        ws.auto_filter.ref = f"A{HEAD_ROW}:{get_column_letter(n)}{last}"
        add_row_colors(ws, HEAD_ROW + 1, last, n, st_col, err_col)
        names_col = get_column_letter(n)        # один ИНН — разные названия: название и «Названий» оранжевые
        for rng in (f"C{HEAD_ROW + 1}:C{last}", f"{names_col}{HEAD_ROW + 1}:{names_col}{last}"):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f"${names_col}{HEAD_ROW + 1}>1"],
                                                           fill=cf_fill(ERR_COLOR), font=Font(bold=True)))
    progress(68, "Пишу лист «Сверка со списком»…")
    recon = write_reconcile_sheet(wb, data)
    progress(69, "Пишу лист «Одна орг. во всех файлах»…")
    all_inn = write_all_by_inn_sheet(wb, data.paths, mappings)
    big = write_threshold_sheet(wb, data.paths, mappings, threshold)
    progress(70, "Пишу лист «Кто кому платит»…")
    write_pairs_sheet(wb, data)
    write_months_sheet(wb, data)
    progress(80, "Пишу листы с ошибками…")
    write_errors_sheets(wb, data.details)
    write_settings_sheet(wb, data, mappings, threshold)
    progress(95, "Сохраняю…")
    wb.save(path)
    return {"path": path, "rows": len(rows), "statuses": Counter(st for _, st in rows if st),
            "flags": Counter(f for e, _ in rows for f in e.flags), "reconcile": recon, "all_by_inn": all_inn,
            "threshold": big}


def _change_cells(ws, e, parts, pair) -> list:
    delta, pct = change(e, parts, pair)
    return [_cell(ws, delta, number_format=NUM) if delta is not None else None,
            _cell(ws, pct, number_format="0.0%") if pct is not None else None]


def _legend(ws, hints_st: dict):
    """Строка 4: цвета «что случилось» и цвет ошибок — с подсказками при наведении."""
    cells = [None, None, _cell(ws, "Цвет строки (наведите — подсказка) →", font=Font(bold=True),
                               alignment=Alignment(horizontal="right", vertical="center"))]
    if hints_st:
        cells += [_cell(ws, st, hints_st[st], fill=fill_color(c), border=BORDER, alignment=CENTER)
                  for st, c in COLORS.items()]
    else:
        cells += [_cell(ws, "сравнения нет (одна часть или реестр + список)", font=Font(italic=True, color="64748B"))]
    cells += [_cell(ws, "Ошибки", "\n".join(f"• {f}: {FLAG_HINTS[f]}" for f in FLAGS) +
                    "\nПодробно — на листах ошибок (у каждого вида — свой лист).", fill=fill_color(ERR_COLOR), border=BORDER,
                    alignment=CENTER, font=Font(bold=True))]
    ws.append(cells)


def _head_hints(parts, with_role, hints_st: dict, pair) -> dict:
    hints = {"ИНН": "ИНН организации (как записан в файле).",
             "Название": "Самое частое написание названия у этого ИНН.",
             "Кто": "плательщик — кто платил; получатель — кому платили; организация — из обычной таблицы с ИНН.",
             "Всего": "Сумма по всем частям.",
             "Что случилось": "\n".join(hints_st.values()) or "Сравнения нет: одна часть или реестр вместе со списком (для них — лист «Сверка со списком»).",
             "Ошибки": "\n".join(f"• {f}: {h}" for f, h in FLAG_HINTS.items()),
             "Названий": "Сколько разных написаний названия у этого ИНН (пусто — одно)."}
    if pair:
        a, b = f"«{parts[pair[0]].label}»", f"«{parts[pair[1]].label}»"
        hints["Изменение"] = f"Сумма в {b} минус сумма в {a}. Пусто — организации нет в одной из них."
        hints["Изменение, %"] = (f"На сколько процентов изменилась сумма: {a} → {b}. "
                                 "Фильтром можно отобрать, например, больше 50%.")
    hints.update({p.label: f"Сумма за часть «{p.label}» (файл {p.source})." if p.has_sums else
                  f"В таблице «{p.label}» (файл {p.source}) нет сумм — «есть» значит, что организация там есть."
                  for p in parts})
    return hints
