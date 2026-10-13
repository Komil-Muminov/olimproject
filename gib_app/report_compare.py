"""Сборка полного Excel-отчета о сравнении двух списков («Муқоиса»)."""
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font
from openpyxl.utils import get_column_letter

from compare import sorted_changes
from report_compare_category import create_category_sheet
from report_styles import (
    BORDER, CENTER, COLORS, MED, NUM, THIN, VCENTER,
    apply_diff_colors, fill_color, styled_cell
)


def write_compare_report(path, src, sheet, diff, a, b):
    years = bool(a.year and b.year)
    y1, y2 = (a.year, b.year) if years else (f"Левая ({a.span})", f"Правая ({b.span})")
    in1 = f"в {a.year}" if years else f"в левой части ({a.span})"
    in2 = f"в {b.year}" if years else f"в правой части ({b.span})"
    metrics = [p[0] for p in diff.pairs]

    cats = [
        ("new", "Новые" if years else "Только справа",
         "🆕 НОВЫЕ" if years else f"🆕 ТОЛЬКО В ПРАВОЙ ЧАСТИ ({b.span})",
         f"Организации, которые есть {in2}, но которых не было {in1}." if years
         else f"Организации, которые есть {in2}, но которых нет {in1}."),
        ("gone", "Пропавшие" if years else "Только слева",
         "❌ ПРОПАВШИЕ" if years else f"❌ ТОЛЬКО В ЛЕВОЙ ЧАСТИ ({a.span})",
         f"Организации, которые были {in1}, но которых нет {in2}." if years
         else f"Организации, которые есть {in1}, но которых нет {in2}."),
    ]
    if metrics:
        cats += [("changed", "Изменения", "🔄 ИЗМЕНИЛАСЬ СУММА",
                  f"Организации, которые есть и {in1}, и {in2}, но сумма у них другая. Крупные разницы — сверху."),
                 ("same", "Без изменений", "✓ БЕЗ ИЗМЕНЕНИЙ",
                  f"Организации, которые есть и {in1}, и {in2}, и суммы совпадают.")]
    else:
        cats += [("both", "Есть в обоих", "✓ ЕСТЬ В ОБОИХ СПИСКАХ",
                  f"Организации, которые есть и {in1}, и {in2}. Суммы не сравнивались — {in2} их нет.")]
    cats += [("dup", "Дубликаты", "👥 ДУБЛИКАТЫ",
              "Один и тот же ИНН записан в списке несколько раз — проверьте, не задвоена ли организация."),
             ("names", "Разные названия", "✏ РАЗНЫЕ НАЗВАНИЯ",
              "Один ИНН, а название записано по-разному. Организации всё равно сопоставлены по ИНН.")]

    def side_vals(vals, side):
        return [(v, None) if side == 1 else (None, v) for v in vals]

    by_size = lambda o: -sum(abs(v or 0) for pair in o[2] for v in pair)
    dup_notes = {(k, n): f"{'левая' if i == 1 else 'правая'} часть · строки {', '.join(map(str, rr))} "
                         f"({len(rr)} раза)" for i, k, n, rr in diff.duplicates}
    names2 = {k: n2 for k, _, n2 in diff.org_names}
    data = {
        "new": (sorted([(k, n, side_vals(v, 2)) for k, n, v in diff.org_only2], key=by_size), {}),
        "gone": (sorted([(k, n, side_vals(v, 1)) for k, n, v in diff.org_only1], key=by_size), {}),
        "changed": (sorted_changes(diff), {}),
        "same": (diff.org_same, {}),
        "both": ([(k, n, []) for k, n, _ in diff.org_same + diff.org_changed], {}),
        "dup": ([(k, n, [(None, None)] * len(metrics)) for _, k, n, _ in diff.duplicates],
                dict(note_title="Где повторяется", note=lambda it: dup_notes.get((it[0], it[1]), ""))),
        "names": (sorted([(k, n1, [(None, None)] * len(metrics)) for k, n1, _ in diff.org_names], key=lambda o: o[1]),
                  dict(name_title=f"Как записано {in1}", note_title=f"Как записано {in2}",
                       note=lambda it: names2.get(it[0], ""))),
    }

    wb = Workbook()
    summary = wb.active
    summary.title = "Итого"
    subtotals = {}
    for key, sheet_name, title, hint in cats:
        items, extra = data[key]
        tbl_metrics = metrics if key in ("new", "gone", "changed", "same") else []
        _, sub = create_category_sheet(wb, sheet_name, title, hint, key, items, tbl_metrics, y1, y2, **extra)
        subtotals[key] = (sheet_name, sub, len(items))

    _write_summary_sheet(summary, src, sheet, a, b, cats, subtotals, metrics, y1, y2, diff)
    wb.save(path)
    return path


def _write_summary_sheet(ws, src, sheet, a, b, cats, subtotals, metrics, y1, y2, diff):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "0F172A"
    first_val = 3
    n_cols = 2 + 3 * len(metrics) + 1
    note_col = n_cols
    for i, w in enumerate([44, 14] + [16, 16, 15] * len(metrics) + [70], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    styled_cell(ws, 1, 1, "Сравнение двух списков по ИНН / РМА", font=Font(bold=True, size=16, color="0F172A"))
    ws.row_dimensions[1].height = 28
    styled_cell(ws, 2, 1, f"Файл: {os.path.basename(src)}   ·   лист «{sheet}»   ·   создано {datetime.now():%d.%m.%Y %H:%M}",
                font=Font(color="64748B"))
    styled_cell(ws, 3, 1, f"Сравниваются: {a.where()}   ↔   {b.where()}", font=Font(bold=True, color="334155"))
    styled_cell(ws, 4, 1, "Каждая категория — на своём листе (вкладки внизу). Название категории в таблице — ссылка на её лист.",
                font=Font(italic=True, color="64748B"))

    row = 6
    head = ["Категория", "Организаций"]
    for r in (row, row + 1):
        for c in range(1, n_cols + 1):
            styled_cell(ws, r, c, font=Font(bold=True, color="0F172A"), fill=fill_color("E2E8F0"), border=BORDER, alignment=CENTER)
    for c, t in enumerate(head, 1):
        ws.merge_cells(start_row=row, start_column=c, end_row=row + 1, end_column=c)
        styled_cell(ws, row, c, t)
    ws.merge_cells(start_row=row, start_column=note_col, end_row=row + 1, end_column=note_col)
    styled_cell(ws, row, note_col, "Что это значит")
    for i, m in enumerate(metrics):
        c = first_val + 3 * i
        ws.merge_cells(start_row=row, start_column=c, end_row=row, end_column=c + 2)
        styled_cell(ws, row, c, m)
        for j, t in enumerate((y1, y2, "Разница")):
            styled_cell(ws, row + 1, c + j, t)
    ws.row_dimensions[row].height = 34
    row += 2

    top = row
    total_keys = [k for k, *_ in cats if k not in ("dup", "names")]
    for key, sheet_name, title, hint in cats:
        if key == "names":
            continue
        color, light = COLORS[key]
        name, sub, count = subtotals[key]
        link = styled_cell(ws, row, 1, title, font=Font(bold=True, size=12, color=color, underline="single"),
                           alignment=Alignment(vertical="center", indent=1))
        link.hyperlink = f"#'{name}'!A1"
        styled_cell(ws, row, 2, count, font=Font(bold=True, size=12), alignment=Alignment(horizontal="center", vertical="center"))
        if sub and key in total_keys:
            for c in range(len(metrics) * 3):
                styled_cell(ws, row, first_val + c, f"='{name}'!{get_column_letter(4 + c)}{sub}", number_format=NUM,
                            font=Font(bold=True), alignment=VCENTER)
        styled_cell(ws, row, note_col, hint, font=Font(color="475569"), alignment=VCENTER)
        for c in range(1, n_cols + 1):
            styled_cell(ws, row, c, fill=fill_color(light), border=BORDER)
        ws.row_dimensions[row].height = 26
        row += 1

    styled_cell(ws, row, 1, "ВСЕГО ПО СПИСКАМ", font=Font(bold=True, size=13, color="FFFFFF"),
                alignment=Alignment(vertical="center", indent=1))
    styled_cell(ws, row, 2, f"{len(a.rows)} → {len(b.rows)}", font=Font(bold=True, size=12, color="FFFFFF"),
                alignment=Alignment(horizontal="center", vertical="center"))
    for c in range(len(metrics) * 3):
        let = get_column_letter(first_val + c)
        rows = [str(top + i) for i, (k, *_) in enumerate(c_ for c_ in cats if c_[0] != "names") if k in total_keys]
        styled_cell(ws, row, first_val + c, "=" + "+".join(f"{let}{r}" for r in rows), number_format=NUM,
                    font=Font(bold=True, size=12, color="FFFFFF"), alignment=VCENTER)
    styled_cell(ws, row, note_col, f"{a.where()}  →  {b.where()}", font=Font(bold=True, color="E2E8F0"), alignment=VCENTER)
    for c in range(1, n_cols + 1):
        styled_cell(ws, row, c, fill=fill_color("334155"), border=Border(top=MED, bottom=MED, left=THIN, right=THIN))
    ws.row_dimensions[row].height = 30
    if metrics:
        row += 1
        styled_cell(ws, row, 1, "Проверка — суммы самих списков:", font=Font(italic=True, color="64748B"),
                    alignment=Alignment(indent=1))
        for i, (_, s1, s2) in enumerate(diff.totals):
            c = first_val + 3 * i
            styled_cell(ws, row, c, s1, number_format=NUM, font=Font(italic=True, color="64748B"))
            styled_cell(ws, row, c + 1, s2, number_format=NUM, font=Font(italic=True, color="64748B"))
    for i in range(len(metrics)):
        apply_diff_colors(ws, first_val + 3 * i + 2, top, row)

    row += 2
    name, _, count = subtotals["names"]
    link = styled_cell(ws, row, 1, f"✏ РАЗНЫЕ НАЗВАНИЯ — {count}", font=Font(bold=True, color=COLORS["names"][0], underline="single"))
    link.hyperlink = f"#'{name}'!A1"
    styled_cell(ws, row, 2, None)
    styled_cell(ws, row, 3, "один ИНН, а название записано по-разному — на суммы не влияет", font=Font(color="475569"))
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
