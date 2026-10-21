"""Строки файлов как есть по условию (ИНН, порог суммы): сбор из всех взятых файлов и лист Excel —
сверху где нашлось и сумма по колонкам «Сумма», ниже строки, каждая колонка подписана «файл: колонка»."""
import os

from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from report_fast import _cell, title_rows
from report_source import data_rows, head_row, show, sources as table_sources, table_name
from report_styles import BORDER, CENTER, NUM, fill_color
from utils import format_money, to_num_safe

BASE = ["Файл", "Лист", "Строка", "Сумма строки"]
MAX_ROWS = 500_000          # больше — не помещается в разумный Excel: показываются первые


def row_sum(values, m: dict):
    """Сумма строки по колонкам «Сумма» из карточки; колонок нет — None."""
    money = m.get("sums") or []
    return round(sum(to_num_safe(values[c]) for c in money if c < len(values)), 2) if money else None


def in_range(total, threshold) -> bool:
    """Порог (от, до): пусто — без ограничения; у строки без суммы порог не проверить — не подходит."""
    if not threshold:
        return True
    lo, hi = threshold
    return total is not None and (lo is None or total >= lo) and (hi is None or total <= hi)


def threshold_text(threshold) -> str:
    lo, hi = threshold
    if lo is not None and hi is not None:
        return f"= {format_money(lo)}" if lo == hi else f"{format_money(lo)} – {format_money(hi)}"
    return f"от {format_money(lo)}" if lo is not None else f"до {format_money(hi)}"


def collect(paths: list, mappings: dict, keep, progress=lambda p, t: None) -> dict:
    """keep(значения, настройка) → брать ли строку: False — нет, True — да, [метки] — да, с метками (найденные ИНН).
    → {"sources", "heads", "hits", "per", "by_tag": {(метка, таблица): [строк, сумма]}, "cut"}."""
    sources, hits, heads, per, by_tag, cut = table_sources(paths, mappings or {}), [], [], [], {}, False
    for src, (p, m) in enumerate(sources):
        progress(int(90 * src / max(len(sources), 1)), f"Ищу в {os.path.basename(p)}…")
        heads.append(head_row(p, m))
        count, total = 0, 0.0
        for r, values in data_rows(p, m):
            got = keep(values, m)
            if not got:
                continue
            tags = got if isinstance(got, list) else []
            count, s = count + 1, row_sum(values, m)
            total += s or 0.0
            for t in tags:
                stat = by_tag.setdefault((t, src), [0, 0.0])
                stat[0], stat[1] = stat[0] + 1, stat[1] + (s or 0.0)
            if len(hits) < MAX_ROWS:
                hits.append((src, r, s, [show(v) if v != "" else None for v in values], tags))
            else:
                cut = True
        per.append((table_name(p, m, sources), count, round(total, 2) if m.get("sums") else None))
    return {"sources": sources, "heads": heads, "hits": hits, "per": per, "by_tag": by_tag, "cut": cut}


def _tag_rows(found: dict, tags: list) -> list:
    """Сводка по меткам: (ИНН, таблица, строк, сумма) — в порядке ввода; не найденный — одной строкой."""
    out = []
    for t in tags:
        stats = sorted((src, st) for (tag, src), st in found["by_tag"].items() if tag == t)
        out += [(t, found["per"][src][0], st[0], round(st[1], 2) if found["sources"][src][1].get("sums") else None)
                for src, st in stats] or [(t, "не найден ни в одном файле", 0, None)]
    return out


def write_rows_sheet(wb, name: str, title: str, hint: str, found: dict, tags: list | None = None):
    """Лист: «Где нашлось» (таблица, строк, сумма), по меткам (если tags — несколько ИНН) и строки как есть."""
    sources, heads, hits, per = found["sources"], found["heads"], found["hits"], found["per"]
    base = BASE[:3] + ["ИНН"] + BASE[3:] if tags else BASE
    widths = [len(h) for h in heads]
    for src, _, _, values, _ in hits:
        widths[src] = max(widths[src], len(values))
    labels = base + [f"{per[s][0]}: {get_column_letter(m['col0'] + i + 1)}" + (f" {heads[s][i]}" if i < len(heads[s]) and heads[s][i] else "")
                     for s, (_, m) in enumerate(sources) for i in range(widths[s])]
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = "7C3AED"
    note = f" Показаны первые {MAX_ROWS:,} строк.".replace(",", " ") if found["cut"] else ""
    title_rows(ws, title, "7C3AED", hint + note, len(labels))
    head = lambda texts: ws.append([_cell(ws, t, font=Font(bold=True), fill=fill_color("E2E8F0"), border=BORDER,
                                          alignment=CENTER) for t in texts])
    money = lambda v: (lambda c: (setattr(c, "number_format", NUM), c)[1])(WriteOnlyCell(ws, value=v))
    head(["Где нашлось", "Строк", "Сумма (колонки «Сумма»)"])
    for table, count, total in per:
        ws.append([table, count, money(total) if total is not None else "сумма не указана"])
    ws.append([])
    if tags:
        head(["ИНН", "Где нашлось", "Строк", "Сумма (колонки «Сумма»)"])
        for t, table, count, total in _tag_rows(found, tags):
            ws.append([t, table, count or None, money(total) if total is not None else None])
        ws.append([])
    head(labels)
    offsets = [len(base) + sum(widths[:s]) for s in range(len(sources))]
    for src, r, s, values, row_tags in hits:
        p, m = sources[src]
        line = [None] * len(labels)
        line[:len(base)] = [os.path.basename(p), m["sheet"], r] + ([", ".join(row_tags)] if tags else []) \
            + [money(s) if s is not None else None]
        line[offsets[src]:offsets[src] + len(values)] = values
        ws.append(line)
    return ws
