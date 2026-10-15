"""Обычные файлы со списками организаций (таблицы с ИНН на любых листах) → части и организации
глубокого анализа. Ошибки — те же правила, что в проверке ИНН (inn_check)."""
import os

from openpyxl.utils import get_column_letter

from analysis_list_map import FIND, mapped_table, side_guess, transpose
from analysis_model import BAD_INN, BY_NAME, DUP_ROW, HIDDEN, ORG, SHARED, Entity, Part
from analysis_registry import inn_problem
from excel_sheet_info import hidden_note, sheet_scan
from compare_finder import find_tables, name_at
from analysis_checks import distinct_names
from inn_check import check_inn
from registry_stats import INN_DIGITS, inn_ok
from utils import (cell_val, has_invisible, is_total_label, norm_text, not_a_number, number_as_text, show_invisible,
                   to_num_safe)

# деньги в двух столбцах (ушло / пришло) — считаются оба
TWO_SIDED = ("дебет", "кредит", "debit", "credit", "приход", "расход")


def _named(data, s, words) -> list:
    """Столбцы с числами, чей заголовок начинается с одного из words."""
    return [c for c in s.value_cols if norm_text(cell_val(data, s.header, c) or "").startswith(words)]


def money_cols(data, s) -> list:
    """Столбцы сумм таблицы: Дебет и Кредит (Приход/Расход) — оба; иначе столбец «Сумма», иначе первый с числами."""
    if s.fixed_money:
        return s.value_cols
    both = _named(data, s, TWO_SIDED)
    return both if len(both) > 1 else (_named(data, s, FIND["amount"])[:1] or s.value_cols[:1])


def money_unclear(sheets: dict) -> dict | None:
    """Несколько столбцов с числами и ни один не «Сумма» (и не Дебет/Кредит) — какой считать деньгами,
    скажет пользователь. → догадка для окна выбора столбцов или None."""
    for sheet, data in sheets.items():
        found = find_tables(data)
        for s in (found[1] if found else []):
            if len(s.value_cols) > 1 and len(_named(data, s, TWO_SIDED)) < 2 and not _named(data, s, FIND["amount"]):
                return side_guess(data, sheet, s)
    return None


def payments_like(sheets: dict) -> bool:
    """Два столбца ИНН в одной таблице (плательщик и получатель) — похоже на реестр платежей."""
    for data in sheets.values():
        found = find_tables(data)
        if found and len({tuple(s.block) for s in found[1]}) < len(found[1]):
            return True
    return False


class ListCollector:
    """Собирает таблицы с ИНН из нескольких файлов: одна организация (ИНН) — одна строка."""

    def __init__(self):
        self.parts = []
        self.orgs = {}          # ключ ИНН (без нулей впереди) → {"inn", "names", "sums", "counts", "dup"}
        self.details = {"dup_row": [], "bad_list": [], "hidden_list": [], "bad_num_list": [], "text_num_list": [],
                        "spaces_list": []}      # группы листов ошибок в таблицах
        self.bad = set()        # ИНН (как в файле) с ошибкой
        self.notes = []         # скрытые листы и строки (сводка)
        self.last_total = 0.0   # сумма последней прочитанной таблицы (каждая строка один раз) — для сверки с «Итого»

    def add_file(self, path: str, sheets: dict, mapping: dict = None) -> int:
        """→ сколько таблиц с ИНН нашлось в файле. mapping — столбцы, указанные пользователем."""
        fname, before = os.path.basename(path), len(self.parts)
        if mapping:                     # столбцы указал пользователь — читается только этот лист
            data = sheets[mapping["sheet"]]
            data = transpose(data) if mapping.get("transpose") else data       # организации по столбцам
            s = mapped_table(data, mapping)
            if not mapping.get("transpose"):    # у перевёрнутой строки листа — это столбцы таблицы
                s.hidden_rows = sheet_scan(path, mapping["sheet"])["hidden_rows"]
            note = hidden_note(path, mapping["sheet"], s.hidden_rows)
            self.notes += [note] if note else []
            if note and not note["listed"]:
                s.hidden_rows = set()           # скрыта почти вся таблица — только сводка
            rows = {r for rr in s.all_rows.values() for r in rr}
            self.last_total = sum(to_num_safe(cell_val(data, r, c)) for r in rows for c in money_cols(data, s))
            if s.all_rows:
                self._add_side(fname, mapping["sheet"], data, s, 1)
                self._mapped_errors(fname, mapping["sheet"], data, s)
            return len(self.parts) - before
        for sheet, data in sheets.items():
            found = find_tables(data)
            if found:
                for s in found[1]:
                    self._add_side(fname, sheet, data, s, len(found[1]))
        if len(self.parts) > before:
            self._add_errors(fname, sheets)
        return len(self.parts) - before

    def _add_side(self, fname, sheet, data, s, on_sheet):
        """Таблица → часть; столбцы «показывать отдельно» (План, Факт) → по части на каждый столбец."""
        where = f"«{sheet}» {s.span}" if on_sheet > 1 else f"«{sheet}»"
        metric = money_cols(data, s)
        head = lambda c: " ".join(str(cell_val(data, s.header, c) or "").split()) or f"столбец {get_column_letter(c)}"
        for group in ([[c] for c in metric] if s.separate and metric else [metric]):
            label = f"{where}: {head(group[0])}" if s.separate and group else where
            self._add_part(fname, data, s, f"{s.year} · {label}" if s.year else label, group, [head(c) for c in group])

    def _add_part(self, fname, data, s, label, metric, heads):
        i = len(self.parts)
        self.parts.append(Part(label, fname, "list", (int(s.year) if s.year else 9999, 0, i), has_sums=bool(metric),
                               heads=heads))
        for k, rows in s.all_rows.items():
            o = self.orgs.setdefault(k, {"inn": s.keys[k], "names": [], "sums": {}, "counts": {}, "cols": {},
                                         "dup": False, "notes": set()})
            o["names"] += [n for n in (s.row_names.get((k, r)) or name_at(data, r, s) for r in rows) if n]
            o["notes"] |= {SHARED} if k in s.shared else set()
            o["notes"] |= {BY_NAME} if s.by_name else set()
            o["notes"] |= {HIDDEN} if s.hidden_rows.intersection(rows) else set()
            o["cols"][i] = [round(sum(to_num_safe(cell_val(data, r, c)) for r in rows), 2) for c in metric]
            o["sums"][i] = round(sum(o["cols"][i]), 2)
            o["counts"][i] = len(rows)
            o["dup"] = o["dup"] or (len(rows) > 1 and not s.repeats)

    def _mapped_errors(self, fname, sheet, data, s):
        """Столбцы указал пользователь: повтор ИНН, ИНН не из 9 цифр и строки без ИНН — те же правила, что везде."""
        where, details = f"«{sheet}»", self.details
        name = lambda k, r: s.row_names.get((k, r)) or name_at(data, r, s)
        for k, rows in s.all_rows.items():
            if len(rows) > 1 and not s.repeats:          # в выписке контрагент повторяется — это не ошибка
                details["dup_row"].append([fname, where, ", ".join(map(str, rows)), s.keys[k], name(k, rows[0]),
                                           f"{len(rows)} раза"])
            inn = s.keys[k]
            if not s.by_name and not inn_ok(inn):
                fix = inn.zfill(INN_DIGITS) if inn.isdigit() and len(inn) == INN_DIGITS - 1 else ""
                details["bad_list"].append([fname, where, rows[0], inn, name(k, rows[0]), inn_problem(inn), fix])
                self.bad.add(inn)
        for k, rows in s.all_rows.items():
            details["hidden_list"] += [[fname, where, r, s.keys[k], name(k, r)] for r in rows if r in s.hidden_rows]
        money = money_cols(data, s)
        for r in sorted({r for rows in s.all_rows.values() for r in rows}):     # уловки в суммах и в ИНН
            cell = lambda c: [fname, where, r, get_column_letter(c), str(cell_val(data, r, c)).strip(), name_at(data, r, s)]
            details["bad_num_list"] += [cell(c) for c in money if not_a_number(cell_val(data, r, c))]
            details["text_num_list"] += [cell(c) for c in money if number_as_text(cell_val(data, r, c))]
            if has_invisible(raw := cell_val(data, r, s.key_col)):
                details["spaces_list"].append([fname, where, r, get_column_letter(s.key_col), show_invisible(raw),
                                               name_at(data, r, s)])
        if s.by_name or not s.name_col:
            return
        with_inn, end = {r for rows in s.all_rows.values() for r in rows}, s.last_row
        while str(cell_val(data, end + 1, s.name_col) or "").strip():       # таблица — до первой пустой строки
            end += 1
        for r in range(s.header + 1, end + 1):                   # есть название, а ИНН нет (кроме строк итогов)
            text = str(cell_val(data, r, s.name_col) or "").strip()
            if r not in with_inn and len(text) > 3 and not is_total_label(text):
                details["bad_list"].append([fname, where, r, "(пусто)", text, "ИНН не указан", ""])

    def _add_errors(self, fname, sheets):
        res = check_inn(sheets)
        if not res:
            return
        self.details["dup_row"] += [[fname, t.label, ", ".join(map(str, rows)), inn, name, f"{len(rows)} раза"]
                                    for t, inn, name, rows in res.duplicates]
        self.details["bad_list"] += [[fname, t.label, row, raw or "(пусто)", name, problem, fix]
                                     for t, row, raw, name, problem, fix in res.invalid]
        self.bad |= {raw for _, row, raw, *_ in res.invalid if raw}

    def entities(self, offset: int, n_parts: int) -> list:
        """Организации; их суммы — в позициях offset… общей таблицы из n_parts частей."""
        out = []
        for o in self.orgs.values():
            sums, counts, cols = [0.0] * n_parts, [0] * n_parts, [None] * n_parts
            for i, v in o["sums"].items():
                sums[offset + i], counts[offset + i], cols[offset + i] = v, o["counts"][i], o["cols"][i]
            e = Entity(o["inn"], o["names"][0] if o["names"] else "", ORG, sums, counts,
                       names=max(1, len(distinct_names(o["names"]))), cols=cols, variants=distinct_names(o["names"]))
            if o["inn"] in self.bad:
                e.flags.add(BAD_INN)
            if o["dup"]:
                e.flags.add(DUP_ROW)
            e.flags |= o["notes"]
            out.append(e)
        return out
