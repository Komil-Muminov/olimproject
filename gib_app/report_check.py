"""Генерация Excel-отчета о проверке ИНН / РМА."""
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from report_styles import BORDER, CENTER, VCENTER, fill_color, styled_cell


def _simple_sheet(wb, name, title, color, hint, head, widths, rows, shade_by=None):
    """Лист: цветной заголовок, пояснение, таблица с фильтром."""
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = color
    ws.sheet_view.showGridLines = False
    n = len(head)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    total_count = len({r[shade_by - 1] for r in rows}) if shade_by else len(rows)
    styled_cell(ws, 1, 1, f"{title} — {total_count}", font=Font(bold=True, size=14, color="FFFFFF"),
                fill=fill_color(color), alignment=Alignment(vertical="center", indent=1))
    for c in range(2, n + 1):
        ws.cell(1, c).fill = fill_color(color)
    ws.row_dimensions[1].height = 30
    styled_cell(ws, 2, 1, hint, font=Font(italic=True, color="475569"))
    h = 4
    for c, t in enumerate(head, 1):
        styled_cell(ws, h, c, t, font=Font(bold=True, color="0F172A"), fill=fill_color("E2E8F0"), border=BORDER, alignment=CENTER)
    ws.row_dimensions[h].height = 30
    ws.freeze_panes = ws.cell(h + 1, 1)
    if not rows:
        styled_cell(ws, h + 1, 2, "нет — всё в порядке ✓", font=Font(italic=True, color="16A34A"))
        return ws
    shade, prev = False, object()
    for i, row in enumerate(rows):
        r = h + 1 + i
        if shade_by and row[shade_by - 1] != prev:
            shade, prev = not shade, row[shade_by - 1]
        for c, v in enumerate(row, 1):
            x = styled_cell(ws, r, c, v, border=BORDER, alignment=Alignment(vertical="top", wrap_text=c > 2))
            if shade_by and shade:
                x.fill = fill_color("F8FAFC")
    ws.auto_filter.ref = f"A{h}:{get_column_letter(n)}{h + len(rows)}"
    return ws


def write_check_report(path, src, res):
    """Сборка Excel-отчета с результатами проверки таблиц с ИНН."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Итого"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "0F172A"
    for col, w in zip("ABC", (46, 16, 80)):
        ws.column_dimensions[col].width = w
    styled_cell(ws, 1, 1, "Проверка таблиц с ИНН / РМА", font=Font(bold=True, size=16, color="0F172A"))
    ws.row_dimensions[1].height = 28
    styled_cell(ws, 2, 1, f"Файл: {os.path.basename(src)}   ·   создано {datetime.now():%d.%m.%Y %H:%M}", font=Font(color="64748B"))
    styled_cell(ws, 4, 1, "Проверены таблицы:", font=Font(bold=True, color="334155"))
    r = 5
    for t in res.tables:
        styled_cell(ws, r, 1, t.label, font=Font(color="334155"), alignment=Alignment(indent=1))
        styled_cell(ws, r, 2, f"{t.count} орг.", font=Font(color="64748B"))
        r += 1

    cats = [
        ("Дубликаты", "👥 ДУБЛИКАТЫ", "7C3AED", "EDE9FE", len(res.duplicates),
         "один и тот же ИНН записан в таблице несколько раз — возможно, организация задвоена"),
        ("Разные названия", "✏ ОДИН ИНН — РАЗНЫЕ НАЗВАНИЯ", "0891B2", "CFFAFE", len(res.names),
         "у одного ИНН в разных строках или на разных листах название записано по-разному"),
        ("Неправильные ИНН", "⚠ НЕПРАВИЛЬНЫЕ ИНН", "DC2626", "FEE2E2", len(res.invalid),
         "ИНН пустой, не того количества цифр, с буквами или пробелами, потерялся ноль в начале"),
    ]
    r += 1
    for c, t in enumerate(("Что проверяли", "Найдено", "Что это значит"), 1):
        styled_cell(ws, r, c, t, font=Font(bold=True, color="0F172A"), fill=fill_color("E2E8F0"), border=BORDER, alignment=CENTER)
    r += 1
    for sheet_name, title, color, light, count, hint in cats:
        link = styled_cell(ws, r, 1, title, font=Font(bold=True, size=12, color=color, underline="single"),
                           alignment=Alignment(vertical="center", indent=1))
        link.hyperlink = f"#'{sheet_name}'!A1"
        styled_cell(ws, r, 2, count if count else "0 ✓", font=Font(bold=True, size=12, color=color if count else "16A34A"),
                    alignment=Alignment(horizontal="center", vertical="center"))
        styled_cell(ws, r, 3, hint, font=Font(color="475569"), alignment=VCENTER)
        for c in range(1, 4):
            styled_cell(ws, r, c, fill=fill_color(light), border=BORDER)
        ws.row_dimensions[r].height = 26
        r += 1
    styled_cell(ws, r + 1, 1, "Название проверки — ссылка на её лист (вкладки внизу).", font=Font(italic=True, color="64748B"))

    _simple_sheet(wb, "Дубликаты", "👥 ДУБЛИКАТЫ", "7C3AED",
                  "Один и тот же ИНН записан в таблице несколько раз. Проверьте, не задвоена ли организация.",
                  ["№", "ИНН / РМА", "Организация", "Таблица", "Строки", "Сколько раз"], [6, 13, 50, 46, 18, 12],
                  [[i, k, n, t.label, ", ".join(map(str, rows)), len(rows)]
                   for i, (t, k, n, rows) in enumerate(res.duplicates, 1)])
    name_rows = []
    for i, (k, variants) in enumerate(res.names, 1):
        for name, t, row in variants:
            name_rows.append([i, k, name, t.label, row])
    _simple_sheet(wb, "Разные названия", "✏ ОДИН ИНН — РАЗНЫЕ НАЗВАНИЯ", "0891B2",
                  "У одного ИНН название записано по-разному. Каждый вариант — отдельной строкой; строки одного ИНН "
                  "выделены одинаковым фоном.",
                  ["№", "ИНН / РМА", "Как записано", "Где (таблица)", "Строка"], [6, 13, 56, 46, 9], name_rows, shade_by=1)
    _simple_sheet(wb, "Неправильные ИНН", "⚠ НЕПРАВИЛЬНЫЕ ИНН", "DC2626",
                  "ИНН, который выглядит неправильно. В последнем столбце — как, вероятно, должно быть.",
                  ["№", "Таблица", "Строка", "Что стоит", "Организация", "Проблема", "Как должно быть"],
                  [6, 44, 9, 16, 46, 40, 16],
                  [[i, t.label, row, raw, name, problem, fix]
                   for i, (t, row, raw, name, problem, fix) in enumerate(res.invalid, 1)])
    wb.save(path)
    return path
