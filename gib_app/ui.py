"""Интерфейс: открыть выписку → «Заполнить ГИБ» → готово.

Программа сама находит лист, столбец с назначениями и строку шапки, ставит суммы
в столбцы ГИБ той же строки (недостающие создаёт) и сохраняет этот же файл.
"""
import os
import sys

from openpyxl.utils import get_column_letter
from PySide6.QtCore import QSettings, Qt, QThread, Signal
from PySide6.QtWidgets import (QFileDialog, QFrame, QHBoxLayout, QLabel, QMainWindow, QMessageBox,
                               QProgressBar, QPushButton, QStackedWidget, QVBoxLayout, QWidget)

from detect import by_header, detect, sheet_score
from excel_io import Book, WriteError, embed_automation, fill_in_place
from parser import parse_purpose_ex, to_number

EXCEL_EXT = (".xlsx", ".xlsm", ".xls")


def fmt(x):
    return f"{x:,.2f}".replace(",", " ").replace(".", ",")


def rows_list(rows, limit=8):
    """«строка 48» / «строки 48, 52, 60 и ещё 3»"""
    shown = ", ".join(str(r) for r in rows[:limit])
    return ("строка " if len(rows) == 1 else "строки ") + shown + (
        f" и ещё {len(rows) - limit}" if len(rows) > limit else "")


def template_path():
    """Шаблон с макросом — рядом с программой (и внутри exe)."""
    base = getattr(sys, "_MEIPASS", os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
    return os.path.join(base, "assets", "gib_template.xlsm")


class Worker(QThread):
    progress = Signal(int, str)
    done = Signal(dict)
    failed = Signal(str)

    def __init__(self, func, *args):
        super().__init__()
        self.func, self.args = func, args

    def run(self):
        try:
            self.done.emit(self.func(*self.args, progress=lambda p, t: self.progress.emit(p, t)))
        except WriteError as e:
            self.failed.emit(str(e))
        except Exception as e:  # ошибки COM и прочее
            self.failed.emit(f"{type(e).__name__}: {e}")


class Plan:
    """Что нашлось в файле и что будет сделано."""

    def __init__(self, path):
        self.path = path
        book = Book(path)
        # лист, где больше всего назначений с ГИБ
        scores = {n: sheet_score(d) for n, d in book.sheets.items()}
        self.sheet = max(scores, key=scores.get) if scores else None
        data = book.sheets.get(self.sheet, [])
        self.layout = detect(data) if self.sheet else None
        self.items, self.no_gib, self.done, self.mismatch = [], [], [], []
        self.long_codes = []          # (строка, ["22110", …]) — в ГИБ больше 4 цифр, взяты первые 4
        self.codes, self.new_codes, self.total = [], [], 0.0
        # назначений ещё нет (пустой Бучет) — автоматику всё равно можно встроить по заголовку
        self.header_only = None
        if not self.layout:
            for name in [self.sheet] + [n for n in book.sheets if n != self.sheet]:
                lay = by_header(book.sheets.get(name, []))
                if lay:
                    self.sheet, self.header_only = name, lay
                    break
            return
        lay = self.layout
        for r in range(lay.first_row, lay.last_row + 1):
            v = Book.value(data, r, lay.purpose_col)
            if not isinstance(v, str) or not v.strip():
                continue
            gibs, longs = parse_purpose_ex(v)
            if longs:
                self.long_codes.append((r, longs))
            if not gibs:
                self.no_gib.append(r)
                continue
            if lay.gib_cols and all(code in lay.gib_cols and to_number(Book.value(data, r, lay.gib_cols[code])) == a
                                    for code, a in gibs.items()):
                self.done.append(r)
                continue
            if lay.amount_col:
                paid = to_number(Book.value(data, r, lay.amount_col))
                if paid is not None and abs(paid - sum(gibs.values())) > 0.005:
                    self.mismatch.append(r)
            self.items.append((r, gibs))
        self.codes = sorted({c for _, g in self.items for c in g})
        self.new_codes = [c for c in self.codes if c not in lay.gib_cols]
        self.total = round(sum(sum(g.values()) for _, g in self.items), 2)


class MainWindow(QMainWindow):
    def __init__(self, settings=None):
        super().__init__()
        self.setWindowTitle("Распределение по ГИБ")
        self.resize(900, 640)
        self.setAcceptDrops(True)
        self.settings = settings or QSettings("Olim", "GibDistributor")
        self.plan = None
        self.worker = None
        self.backup = None
        self.result_path = None

        root = QWidget()
        root.setObjectName("root")
        self.setCentralWidget(root)
        lay = QVBoxLayout(root)
        lay.setContentsMargins(40, 30, 40, 30)
        lay.setSpacing(16)
        title = QLabel("Распределение по ГИБ")
        title.setObjectName("appTitle")
        title.setAlignment(Qt.AlignCenter)
        lay.addWidget(title)
        sub = QLabel("Суммы из назначений ставятся в столбцы ГИБ этой же строки. Файл сохраняется на месте.")
        sub.setObjectName("muted")
        sub.setAlignment(Qt.AlignCenter)
        lay.addWidget(sub)

        self.pages = QStackedWidget()
        lay.addWidget(self.pages, 1)

        # 1. выбор файла
        self.drop = QPushButton("📂\n\nПеретащите сюда файл выписки\nили нажмите, чтобы выбрать")
        self.drop.setObjectName("dropZone")
        self.drop.clicked.connect(self.choose_file)
        self.pages.addWidget(self.drop)

        # 2. что найдено
        found = QFrame()
        found.setObjectName("card")
        fl = QVBoxLayout(found)
        fl.setContentsMargins(28, 24, 28, 24)
        fl.setSpacing(12)
        self.file_label = QLabel()
        self.file_label.setObjectName("fileName")
        fl.addWidget(self.file_label)
        self.found_text = QLabel()
        self.found_text.setObjectName("bigText")
        self.found_text.setWordWrap(True)
        self.found_text.setTextFormat(Qt.RichText)
        fl.addWidget(self.found_text)
        self.warn_text = QLabel()
        self.warn_text.setObjectName("warn")
        self.warn_text.setWordWrap(True)
        self.warn_text.setTextFormat(Qt.RichText)
        fl.addWidget(self.warn_text)
        fl.addStretch()
        note = QLabel("🛡  Перед изменением программа сохранит копию файла в папку «_резервные копии» рядом с ним.")
        note.setObjectName("muted")
        note.setWordWrap(True)
        fl.addWidget(note)
        self.progress = QProgressBar()
        self.progress.setVisible(False)
        fl.addWidget(self.progress)
        btns = QHBoxLayout()
        other = QPushButton("Другой файл")
        other.clicked.connect(self.choose_file)
        btns.addWidget(other)
        view = QPushButton("Открыть в Excel")
        view.clicked.connect(self.open_file)
        btns.addWidget(view)
        btns.addStretch()
        self.embed_btn = QPushButton("⚙  Встроить автоматику в Excel")
        self.embed_btn.setObjectName("secondary")
        self.embed_btn.setToolTip("Создаёт копию файла с макросом: дальше ГИБ проставляются сами,\n"
                                  "когда вы вставляете назначения прямо в Excel.")
        self.embed_btn.clicked.connect(self.embed)
        btns.addWidget(self.embed_btn)
        self.go_btn = QPushButton("Заполнить ГИБ")
        self.go_btn.setObjectName("primary")
        self.go_btn.clicked.connect(self.run)
        btns.addWidget(self.go_btn)
        fl.addLayout(btns)
        self.pages.addWidget(found)

        # 3. готово
        result = QFrame()
        result.setObjectName("card")
        rl = QVBoxLayout(result)
        rl.setContentsMargins(28, 24, 28, 24)
        rl.setSpacing(12)
        self.result_head = QLabel()
        self.result_head.setObjectName("resultOk")
        rl.addWidget(self.result_head)
        self.result_text = QLabel()
        self.result_text.setObjectName("bigText")
        self.result_text.setWordWrap(True)
        self.result_text.setTextFormat(Qt.RichText)
        rl.addWidget(self.result_text)
        self.result_warn = QLabel()
        self.result_warn.setObjectName("warn")
        self.result_warn.setWordWrap(True)
        self.result_warn.setTextFormat(Qt.RichText)
        rl.addWidget(self.result_warn)
        rl.addStretch()
        self.backup_label = QLabel()
        self.backup_label.setObjectName("muted")
        self.backup_label.setWordWrap(True)
        rl.addWidget(self.backup_label)
        rb = QHBoxLayout()
        again = QPushButton("Заполнить другой файл")
        again.clicked.connect(self.choose_file)
        rb.addWidget(again)
        rb.addStretch()
        self.backup_btn = QPushButton("Открыть копию (до изменений)")
        self.backup_btn.clicked.connect(lambda: self._start_file(self.backup))
        rb.addWidget(self.backup_btn)
        self.open_btn = QPushButton("Открыть в Excel")
        self.open_btn.setObjectName("primary")
        self.open_btn.clicked.connect(lambda: self._start_file(self.result_path))
        rb.addWidget(self.open_btn)
        rl.addLayout(rb)
        self.pages.addWidget(result)

    # ------------------------------------------------------------ файл
    def open_file(self):
        if self.plan:
            self._start_file(self.plan.path)

    def _start_file(self, path):
        try:
            os.startfile(path)
        except OSError as e:
            QMessageBox.warning(self, "Не удалось открыть", f"{path}\n\n{e}")

    def dragEnterEvent(self, e):
        if e.mimeData().hasUrls():
            e.acceptProposedAction()

    def dropEvent(self, e):
        files = [u.toLocalFile() for u in e.mimeData().urls() if u.toLocalFile().lower().endswith(EXCEL_EXT)]
        if files and not self.worker:
            self.load(files[0])

    def choose_file(self):
        if self.worker:
            return
        path, _ = QFileDialog.getOpenFileName(self, "Выберите файл выписки", self.settings.value("last_dir", ""),
                                              "Excel (*.xlsx *.xlsm *.xls)")
        if path:
            self.load(path)

    def load(self, path):
        self.settings.setValue("last_dir", os.path.dirname(path))
        try:
            plan = Plan(path)
        except Exception as ex:
            QMessageBox.critical(self, "Ошибка", f"Не удалось открыть файл:\n{path}\n\n{ex}")
            return
        self.plan = plan
        self.file_label.setText("📄  " + os.path.basename(path))
        self.warn_text.setText("")
        lay = plan.layout
        self.embed_btn.setEnabled(bool(lay or plan.header_only))
        if not lay and plan.header_only:
            h = plan.header_only
            self.found_text.setText(
                f"Назначений с ГИБ в файле пока нет.<br><br>"
                f"Можно <b>встроить автоматику</b>: после этого ГИБ будут проставляться сами, когда вы "
                f"вставите назначения в столбец {get_column_letter(h.purpose_col)} "
                f"(лист «{plan.sheet}», заголовок в строке {h.header_row}).")
            self.go_btn.setEnabled(False)
        elif not lay:
            self.found_text.setText("В этом файле не нашлось назначений с «ГИБ XXXX-сумма».")
            self.go_btn.setEnabled(False)
        elif not plan.items:
            self.found_text.setText(f"Все суммы ГИБ уже проставлены ✓<br><br>"
                                    f"<span style='color:#64748B'>Лист «{plan.sheet}», "
                                    f"назначений с ГИБ: {len(plan.done)}</span>")
            self.go_btn.setEnabled(False)
        else:
            parts = [f"Найдено назначений с ГИБ: <b>{len(plan.items)}</b> на сумму <b>{fmt(plan.total)}</b>",
                     f"ГИБ: <b>{', '.join(plan.codes)}</b>"]
            if plan.new_codes:
                parts.append(f"Будут созданы столбцы: <b>{', '.join(plan.new_codes)}</b>")
            else:
                parts.append("Все нужные столбцы ГИБ уже есть")
            if plan.done:
                parts.append(f"<span style='color:#64748B'>Уже заполнено раньше: {len(plan.done)} — их не трогаем</span>")
            parts.append("Для сверки с выпиской будут столбцы «Расходы» и «Разница»")
            parts.append(f"<span style='color:#64748B'>Лист «{plan.sheet}», назначения в столбце "
                         f"{get_column_letter(lay.purpose_col)}, строки {lay.first_row}–{lay.last_row}</span>")
            self.found_text.setText("<br>".join(parts))
            self.go_btn.setEnabled(True)
        self.warn_text.setText(self._warnings(plan))
        self.pages.setCurrentIndex(1)

    @staticmethod
    def _warnings(plan, after=False):
        w = []
        if plan.no_gib:
            w.append(f"⚠ Нет ГИБ в назначении — {rows_list(plan.no_gib)}. "
                     + ("Пропущено, заполните вручную." if after else "Программа это пропустит."))
        if plan.mismatch:
            w.append(f"⚠ Сумма ГИБ не равна сумме платежа — {rows_list(plan.mismatch)}. Проверьте.")
        if plan.long_codes:
            shown = ", ".join(f"строка {r} ({', '.join(f'{c} → {c[:4]}' for c in codes)})"
                              for r, codes in plan.long_codes[:6])
            more = f" и ещё {len(plan.long_codes) - 6}" if len(plan.long_codes) > 6 else ""
            w.append(f"⚠ В ГИБ больше 4 цифр — взяты первые 4: {shown}{more}. Проверьте"
                     + (" (в Excel эти строки жёлтые)." if after else "."))
        return "<br>".join(w)

    # ------------------------------------------------------------ заполнение
    def run(self):
        p = self.plan
        if not p or not p.items:
            return
        self._busy(True)
        self.worker = Worker(fill_in_place, p.path, p.sheet, p.layout.purpose_col, p.layout.header_row, p.items,
                             p.layout.first_row, p.layout.last_row, [r for r, _ in p.long_codes])
        self.worker.progress.connect(lambda v, t: (self.progress.setValue(v), self.progress.setFormat(f"{t}  %p%")))
        self.worker.done.connect(self._done)
        self.worker.failed.connect(self._failed)
        self.worker.start()

    def embed(self):
        p = self.plan
        lay = p.layout or p.header_only if p else None
        if not lay:
            return
        base = os.path.splitext(p.path)[0]
        dst, _ = QFileDialog.getSaveFileName(self, "Куда сохранить файл с автоматикой",
                                             f"{base} (автоматика).xlsm", "Книга Excel с макросами (*.xlsm)")
        if not dst:
            return
        if not dst.lower().endswith(".xlsm"):
            dst += ".xlsm"
        self._busy(True)
        self.worker = Worker(embed_automation, p.path, dst, template_path(), p.sheet,
                             lay.header_row, lay.purpose_col)
        self.worker.progress.connect(lambda v, t: (self.progress.setValue(v), self.progress.setFormat(f"{t}  %p%")))
        self.worker.done.connect(self._embedded)
        self.worker.failed.connect(self._failed)
        self.worker.start()

    def _busy(self, on):
        self.go_btn.setEnabled(not on and bool(self.plan and self.plan.items))
        self.embed_btn.setEnabled(not on)
        self.progress.setVisible(on)
        self.progress.setValue(0)

    def _embedded(self, res):
        self.worker = None
        self.progress.setVisible(False)
        self.result_path = res["path"]
        self._set_head(True, "✓  Автоматика встроена")
        p = self.plan
        lines = [f"Создан файл: <b>{os.path.basename(res['path'])}</b>"]
        if p.items:
            lines.append(f"Сразу заполнено назначений: <b>{len(p.items)}</b>"
                         + (f", созданы столбцы: <b>{', '.join(p.new_codes)}</b>" if p.new_codes else ""))
        lines.append("Добавлены столбцы «Расходы» и «Разница» — для сверки с выпиской.")
        lines += ["",
                  "<b>Как работать дальше:</b>",
                  "1. Откройте этот файл в Excel и нажмите «Включить содержимое» на жёлтой панели.",
                  "2. Вставляйте назначения в столбец назначений — ГИБ проставятся сами.",
                  "   Суммы из выписки вставляйте в «Расходы» — «Разница» покажет, где не сходится.",
                  "3. Кнопка «Пересчитать ГИБ» вверху листа — заново пройтись по всем назначениям."]
        self.result_text.setText("<br>".join(lines))
        self.result_warn.setText("Исходный файл не изменён." + (
            "<br>" + self._warnings(p, after=True).replace("Пропущено, заполните вручную.",
                                                           "В Excel эти строки подсвечены розовым.")
            if p.no_gib else ""))
        self.backup = None
        self.backup_label.setText("")
        self.backup_btn.setVisible(False)
        self.pages.setCurrentIndex(2)

    def _set_head(self, ok, text):
        self.result_head.setText(text)
        self.result_head.setObjectName("resultOk" if ok else "resultBad")
        self.result_head.style().unpolish(self.result_head)
        self.result_head.style().polish(self.result_head)

    def _done(self, res):
        self.worker = None
        self.progress.setVisible(False)
        check = res["check"]
        ok = check["ok"]
        self.result_path = self.plan.path
        self._set_head(ok, "✓  Готово" if ok else "⚠  Готово, но итоги не сходятся")
        lines = [f"Заполнено строк: <b>{res['filled']}</b> на сумму <b>{fmt(check['written_sum'])}</b>"]
        if res["created"]:
            lines.append(f"Созданы столбцы ГИБ: <b>{', '.join(res['created'])}</b>")
        if res.get("header_inserted"):
            lines.append("Над таблицей добавлена строка с номерами ГИБ")
        lines.append("Всё, что вписала программа, подсвечено <span style='background:#DCFCE7'>&nbsp;зелёным&nbsp;</span>.")
        if res.get("left_open"):
            lines.append("Файл был открыт в Excel — изменения уже в нём и сохранены.")
        recon = check.get("recon")
        if recon and recon["empty"]:
            lines.append("📋 Сверка: вставьте суммы из выписки в столбец «Расходы» — «Разница» посчитается сама.")
        elif recon and not recon["bad_rows"]:
            lines.append("✓ Сверка с выпиской: все суммы совпали, разница 0.")
        self.result_text.setText("<br>".join(lines))
        warn = [self._warnings(self.plan, after=True)] if self._warnings(self.plan) else []
        if recon and recon["bad_rows"] and not recon["empty"]:
            warn.append(f"⚠ Сверка с выпиской: не сходится — {rows_list(recon['bad_rows'])}, "
                        f"общая разница {fmt(recon['total_diff'])}. В столбце «Разница» такие строки подсвечены красным.")
        bad = [g for g in check["gibs"] if not (g["written_ok"] and g["total_ok"])]
        for g in bad:
            warn.append(f"⚠ ГИБ {g['code']}: итог «{fmt(g['total'] or 0)}» не равен сумме столбца {fmt(g['column_sum'])}")
        self.result_warn.setText("<br>".join(warn))
        self.backup = res.get("backup")
        self.backup_label.setText(f"🛡  Копия файла до изменений: {self.backup}" if self.backup else "")
        self.backup_btn.setVisible(bool(self.backup))
        self.pages.setCurrentIndex(2)

    def _failed(self, err):
        self.worker = None
        self._busy(False)
        QMessageBox.critical(self, "Ошибка", err)


STYLE = """
* { font-family: 'Segoe UI'; font-size: 10pt; }
QMainWindow, QWidget#root { background: #F5F7FB; }
QLabel#appTitle { font-size: 22pt; font-weight: 700; color: #0F172A; }
QLabel#muted { color: #64748B; }
QLabel#fileName { font-size: 13pt; font-weight: 600; color: #0F172A; }
QLabel#bigText { font-size: 12pt; color: #0F172A; line-height: 150%; }
QLabel#warn { font-size: 11pt; color: #B45309; }
QLabel#resultOk { font-size: 20pt; font-weight: 700; color: #16A34A; }
QLabel#resultBad { font-size: 20pt; font-weight: 700; color: #DC2626; }
QFrame#card { background: white; border: 1px solid #E2E8F0; border-radius: 16px; }
QPushButton { background: #F1F5F9; border: 1px solid #E2E8F0; border-radius: 8px; padding: 8px 16px; color: #0F172A; }
QPushButton:hover { background: #E2E8F0; }
QPushButton#dropZone { border: 2px dashed #CBD5E1; border-radius: 16px; color: #64748B; font-size: 15pt; background: white; }
QPushButton#dropZone:hover { border-color: #2563EB; color: #2563EB; background: #EFF6FF; }
QPushButton#primary { background: #2563EB; color: white; border: none; padding: 12px 28px; font-weight: 600; font-size: 12pt; }
QPushButton#primary:hover { background: #1D4ED8; }
QPushButton#primary:disabled { background: #94A3B8; }
QPushButton#secondary { background: white; color: #2563EB; border: 2px solid #2563EB; padding: 10px 18px;
                        font-weight: 600; font-size: 11pt; }
QPushButton#secondary:hover { background: #EFF6FF; }
QPushButton#secondary:disabled { color: #94A3B8; border-color: #CBD5E1; }
QProgressBar { border: 1px solid #E2E8F0; border-radius: 6px; text-align: center; height: 20px; background: #F1F5F9; }
QProgressBar::chunk { background: #2563EB; border-radius: 6px; }
"""
