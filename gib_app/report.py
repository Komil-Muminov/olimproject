"""Файл с результатом сравнения по ИНН: листы Итог, Новые, Пропавшие, Изменения, Дубликаты.

Создаётся новый файл рядом с исходным, сам исходный файл не меняется.
"""
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NUM = "#,##0.00"
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DBEAFE")
TOTAL_FILL = PatternFill("solid", fgColor="F1F5F9")
UP = Font(color="15803D", bold=True)      # сумма выросла
DOWN = Font(color="B91C1C", bold=True)    # сумма уменьшилась
SHEET_FILL = {"Новые": "FEF9C3", "Пропавшие": "FEE2E2", "Изменения": "DBEAFE", "Дубликаты": "F3E8FF"}


def report_path(src):
    """«файл — различия.xlsx» рядом с исходным; если уже есть — «… (2).xlsx»."""
    base = os.path.splitext(src)[0] + " — различия"
    path, n = base + ".xlsx", 2
    while os.path.exists(path):
        path, n = f"{base} ({n}).xlsx", n + 1
    return path


def _table(ws, head, rows, widths, num_cols=(), total_cols=(), color_col=None):
    """Таблица с шапкой, рамками, фильтром, закреплённой шапкой и строкой «Итого»."""
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for row in rows:
        ws.append(row)
        r = ws.max_row
        for c in range(1, len(head) + 1):
            cell = ws.cell(r, c)
            cell.border = BORDER
            if c in num_cols:
                cell.number_format = NUM
        if color_col and isinstance(ws.cell(r, color_col).value, (int, float)):
            ws.cell(r, color_col).font = UP if ws.cell(r, color_col).value > 0 else DOWN
    last = ws.max_row
    if rows and total_cols:
        t = last + 1
        ws.cell(t, 2, "Итого").font = Font(bold=True)
        for c in range(1, len(head) + 1):
            ws.cell(t, c).fill = TOTAL_FILL
            ws.cell(t, c).border = BORDER
        for c in total_cols:
            L = get_column_letter(c)
            cell = ws.cell(t, c, f"=SUM({L}2:{L}{last})")
            cell.number_format = NUM
            cell.font = Font(bold=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{last}"
    else:
        ws.append(["— нет —"])


def write_compare_report(path, src, sheet, diff, a, b):
    lab1 = f"Список 1{' (' + a.year + ')' if a.year else ''}"
    lab2 = f"Список 2{' (' + b.year + ')' if b.year else ''}"
    wb = Workbook()

    # Итог
    ws = wb.active
    ws.title = "Итог"
    ws["A1"] = "Сравнение по ИНН / РМА"
    ws["A1"].font = Font(bold=True, size=14)
    info = [
        ("Файл", os.path.basename(src)),
        ("Лист", sheet),
        (lab1, a.label(1)),
        (lab2, b.label(2)),
        ("Создано", datetime.now().strftime("%d.%m.%Y %H:%M")),
        (None, None),
        ("🆕 Новые (есть только в списке 2)", len(diff.only2)),
        ("❌ Пропавшие (были только в списке 1)", len(diff.only1)),
        ("🔄 Изменилась сумма (организаций)", len({c[0] for c in diff.changed})),
        ("✓ Без изменений", diff.same),
        ("👥 Дубликаты ИНН", len(diff.duplicates)),
        (None, None),
        ("Сопоставлено по ИНН, названия могут различаться", f"{diff.names_differ} назв. записаны по-разному"),
    ]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
        if k and k[:1] in "🆕❌🔄✓👥":
            ws.cell(i, 1).font = Font(bold=True)
            ws.cell(i, 2).font = Font(bold=True)
    r = len(info) + 5
    if diff.totals:
        ws.cell(r, 1, "Итоги по спискам").font = Font(bold=True, size=12)
        r += 1
        for c, h in enumerate(["Показатель", lab1, lab2, "Разница"], 1):
            cell = ws.cell(r, c, h)
            cell.font = Font(bold=True)
            cell.fill = HEAD_FILL
            cell.border = BORDER
        for title, s1, s2 in diff.totals:
            r += 1
            for c, v in enumerate([title, s1, s2, round(s2 - s1, 2)], 1):
                cell = ws.cell(r, c, v)
                cell.border = BORDER
                if c > 1:
                    cell.number_format = NUM
            ws.cell(r, 4).font = UP if s2 - s1 > 0 else DOWN
    ws.column_dimensions["A"].width = 48
    for col in "BCD":
        ws.column_dimensions[col].width = 26

    # Новые / Пропавшие
    for title, items in (("Новые", diff.only2), ("Пропавшие", diff.only1)):
        sh = wb.create_sheet(title)
        sh.sheet_properties.tabColor = SHEET_FILL[title]
        metrics = list(items[0][2]) if items else []
        head = ["№", "ИНН / РМА", "Организация"] + metrics
        rows = [[i, k, n] + [v.get(m) for m in metrics] for i, (k, n, v) in enumerate(items, 1)]
        nums = tuple(range(4, 4 + len(metrics)))
        _table(sh, head, rows, [6, 14, 50] + [18] * len(metrics), nums, nums)

    # Изменения
    sh = wb.create_sheet("Изменения")
    sh.sheet_properties.tabColor = SHEET_FILL["Изменения"]
    rows = [[i, k, n, t, v1, v2, d] for i, (k, n, t, v1, v2, d) in enumerate(diff.changed, 1)]
    _table(sh, ["№", "ИНН / РМА", "Организация", "Показатель", f"Было — {lab1}", f"Стало — {lab2}", "Разница"],
           rows, [6, 14, 46, 34, 18, 18, 16], (5, 6, 7), (), color_col=7)

    # Дубликаты
    sh = wb.create_sheet("Дубликаты")
    sh.sheet_properties.tabColor = SHEET_FILL["Дубликаты"]
    rows = [[i, f"Список {n}", k, name, ", ".join(map(str, rr)), len(rr)]
            for i, (n, k, name, rr) in enumerate(diff.duplicates, 1)]
    _table(sh, ["№", "Список", "ИНН / РМА", "Организация", "Строки в исходном листе", "Сколько раз"],
           rows, [6, 12, 14, 46, 24, 12])

    wb.save(path)
    return path
