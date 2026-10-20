"""Чтение Excel файлов через openpyxl и базовые утилиты адресации ячеек."""
import os
import re
import tempfile
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from constants import XL_OPENXML
from file_format import real_path

CELL_RE = re.compile(r"^\s*([A-Za-z]{1,3})\s*(\d+)\s*$")
GIB_HEADER_RE = re.compile(r"^\s*(\d{3,4})(?!\d)")


def parse_cell(addr: str):
    """'B5' -> (5, 2). None, если адрес некорректный."""
    m = CELL_RE.match(addr or "")
    if not m:
        return None
    return int(m.group(2)), column_index_from_string(m.group(1).upper())


def cell_name(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def gib_code(value) -> str | None:
    """Заголовок столбца -> код ГИБ (2111, '2111', '2111 Музди маош')."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    m = GIB_HEADER_RE.match(str(value))
    return m.group(1) if m else None


def _xls_to_xlsx(path: str) -> str:
    """Старый .xls читаем через Excel: сохраняем временную копию в .xlsx."""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    out = os.path.join(tempfile.gettempdir(), f"gib_preview_{os.getpid()}.xlsx")
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(os.path.abspath(path), ReadOnly=True)
        wb.SaveAs(out, FileFormat=XL_OPENXML)
        wb.Close(False)
    finally:
        xl.Quit()
    return out


def _calamine_rows(sheet) -> list:
    """Значения листа calamine с первой ячейки (A1), как у openpyxl: пусто — None."""
    row0, col0 = sheet.start or (0, 0)
    return [[] for _ in range(row0)] + [[None] * col0 + [None if v == "" else v for v in r] for r in sheet.to_python()]


def _calamine_sheets(path: str) -> dict:
    """Старый .xls — быстро, без Excel: значения всех листов."""
    from python_calamine import CalamineWorkbook

    wb = CalamineWorkbook.from_path(path)
    return {name: _calamine_rows(wb.get_sheet_by_name(name)) for name in wb.sheet_names}


def sheet_values(path: str, sheet: str) -> list:
    """Один лист любого формата (.xlsx / .xlsm / .xls) — быстро, через calamine; строки и столбцы — с A1."""
    from python_calamine import CalamineWorkbook

    return _calamine_rows(CalamineWorkbook.from_path(real_path(path)).get_sheet_by_name(sheet))


class Book:
    """Значения всех листов книги для показа в интерфейсе."""

    def __init__(self, path: str):
        self.path = path
        read_path = real_path(path)                 # «.xls», внутри которого .xlsx, — читается как .xlsx
        self.sheets: dict[str, list[list]] = {}
        if read_path.lower().endswith(".xls"):
            try:
                raw = _calamine_sheets(read_path)
            except Exception:                       # не прочиталось напрямую — через Excel (медленно)
                raw = self._openpyxl(_xls_to_xlsx(read_path))
        else:
            raw = self._openpyxl(read_path)
        for name, rows in raw.items():
            width = max((len(r) for r in rows), default=0)
            while rows and all(v is None for v in rows[-1]):
                rows.pop()
            self.sheets[name] = [r + [None] * (width - len(r)) for r in rows]

    @staticmethod
    def _openpyxl(path: str) -> dict:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
        wb.close()
        return out

    @staticmethod
    def value(data, row: int, col: int):
        if 1 <= row <= len(data) and 1 <= col <= len(data[row - 1]):
            return data[row - 1][col - 1]
        return None
