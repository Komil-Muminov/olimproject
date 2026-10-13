"""Настройка формул сверки с выпиской, итогов и валидации данных."""
import re
from openpyxl.utils import column_index_from_string, get_column_letter

from constants import (
    NO_GIB_NOTE, RED_FILL, RED_FONT, TOTAL_HEADERS,
    XL_NO_GIB, XL_NO_GIB_FONT
)
from excel_columns import col_values, gib_columns, insert_column, row_values, sum_row_below, used_bounds
from excel_reader import cell_name
from utils import norm_text, to_num_safe

SIMPLE_SUM_RE = re.compile(r"^=SUM\(\$?([A-Z]+)\$?\d+:\$?([A-Z]+)\$?\d+\)$", re.IGNORECASE)


def find_header(ws, header_row: int, after_col: int, prefixes: tuple) -> int | None:
    _, last_col = used_bounds(ws)
    for c, v in enumerate(row_values(ws, header_row, last_col), 1):
        if c > after_col and norm_text(v).startswith(prefixes):
            return c
    return None


def setup_reconciliation(ws, header_row: int, purpose_col: int, first_row: int, last_row: int | None) -> dict:
    total = find_header(ws, header_row, purpose_col, TOTAL_HEADERS)
    rash = find_header(ws, header_row, purpose_col, ("расход",))
    diff = find_header(ws, header_row, purpose_col, ("разниц",))
    total_created = total is None
    if total is None:
        total = (rash or purpose_col) + 1
        insert_column(ws, header_row, total, "Ҳамагӣ")
    if rash is None:
        insert_column(ws, header_row, total, "Расходы")
        rash, total = total, total + 1
    if diff is None:
        insert_column(ws, header_row, total, "Разница")
        diff, total = total, total + 1

    if last_row is None:
        sr = sum_row_below(ws, first_row, purpose_col)
        last_row = sr - 1 if sr else first_row
    sum_row = sum_row_below(ws, last_row, purpose_col)

    rng = ws.Range(ws.Cells(first_row, diff), ws.Cells(last_row, diff))
    current = col_values(ws, diff, first_row, last_row, "Formula")
    out = [f if f not in ("", None) else
           f"=ROUND({cell_name(first_row + i, total)}-{cell_name(first_row + i, rash)},2)"
           for i, f in enumerate(current)]
    rng.Formula = [(v,) for v in out]
    try:
        rng.NumberFormat = "0.00"
    except Exception:
        pass
    try:
        fc = rng.FormatConditions.Add(Type=1, Operator=4, Formula1="=0")
        fc.Interior.Color = RED_FILL
        fc.Font.Color = RED_FONT
        fc.Font.Bold = True
    except Exception:
        pass

    if sum_row:
        for col in (rash, diff) + ((total,) if total_created else ()):
            cell = ws.Cells(sum_row, col)
            if cell.Formula in ("", None):
                cell.Formula = f"=SUM({cell_name(first_row, col)}:{cell_name(last_row, col)})"
    return {"total": total, "rash": rash, "diff": diff, "total_created": total_created, "sum_row": sum_row}


def fill_total_formulas(ws, header_row: int, total: int, first_row: int, last_row: int):
    cols = sorted(gib_columns(ws, header_row).values())
    if not cols:
        return
    rng = ws.Range(ws.Cells(first_row, total), ws.Cells(last_row, total))
    out = []
    for r, f in zip(range(first_row, last_row + 1), col_values(ws, total, first_row, last_row, "Formula")):
        a, b = cols[0], cols[-1]
        m = SIMPLE_SUM_RE.match(str(f or ""))
        if f not in ("", None) and not m:
            out.append(f)
            continue
        if m:
            a = min(a, column_index_from_string(m.group(1).upper()))
            b = max(b, column_index_from_string(m.group(2).upper()))
        out.append(f"=SUM({get_column_letter(a)}{r}:{get_column_letter(b)}{r})")
    rng.Formula = [(v,) for v in out]


def read_reconciliation(ws, header_row: int, first_row: int, last_row: int):
    rash = find_header(ws, header_row, 0, ("расход",))
    diff = find_header(ws, header_row, 0, ("разниц",))
    if not (rash and diff):
        return None
    paid = col_values(ws, rash, first_row, last_row)
    diffs = col_values(ws, diff, first_row, last_row)
    bad = [first_row + i for i, d in enumerate(diffs) if abs(to_num_safe(d)) > 0.005]
    return {"empty": all(v in (None, "") for v in paid), "bad_rows": bad,
            "total_diff": round(sum(to_num_safe(d) for d in diffs), 2)}


def mark_no_gib(ws, col: int, no_gib_rows: list, gib_rows: list):
    for r in no_gib_rows:
        c = ws.Cells(r, col)
        c.Interior.Color = XL_NO_GIB
        c.Font.Color = XL_NO_GIB_FONT
        try:
            c.ClearComments()
            c.AddComment(NO_GIB_NOTE)
        except Exception:
            pass
    for r in gib_rows:
        c = ws.Cells(r, col)
        if int(c.Interior.Color) == XL_NO_GIB:
            c.Interior.Pattern = -4142
            c.Font.ColorIndex = -4105
            try:
                if c.Comment is not None and c.Comment.Text() == NO_GIB_NOTE:
                    c.ClearComments()
            except Exception:
                pass


def verify_rows(ws, items: list, header_row: int, last_row: int) -> dict:
    used_last, _ = used_bounds(ws)
    expected = {}
    for _, gibs in items:
        for code, amount in gibs.items():
            expected[code] = round(expected.get(code, 0.0) + amount, 2)

    result = []
    for code, col in sorted(gib_columns(ws, header_row).items()):
        values = col_values(ws, col, header_row + 1, last_row)
        written = round(sum(to_num_safe(values[r - header_row - 1]) for r, g in items if code in g), 2)
        column_sum = round(sum(to_num_safe(v) for v in values), 2)
        total = None
        if used_last > last_row:
            formulas = col_values(ws, col, last_row + 1, used_last, "Formula")
            for i, f in enumerate(formulas):
                if isinstance(f, str) and f.startswith("="):
                    total = round(to_num_safe(ws.Cells(last_row + 1 + i, col).Value), 2)
                    break
        exp = expected.get(code, 0.0)
        if not exp and not column_sum:
            continue
        result.append({
            "code": code,
            "expected": exp,
            "written": written,
            "column_sum": column_sum,
            "total": total,
            "written_ok": abs(written - exp) < 0.005,
            "total_ok": total is None or abs(total - column_sum) < 0.005,
        })
    return {
        "gibs": result,
        "written_sum": round(sum(r["written"] for r in result), 2),
        "ok": all(r["written_ok"] and r["total_ok"] for r in result),
    }
