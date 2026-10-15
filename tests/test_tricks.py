"""Уловки в ячейках: белый шрифт, числа текстом, лишние пробелы и невидимые знаки в ИНН; один ИНН — разные названия."""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "gib_app")))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from analysis_scan import build, scan_files  # noqa: E402
from excel_sheet_info import white_cells  # noqa: E402
from utils import has_invisible, number_as_text, show_invisible  # noqa: E402

INNS = [f"{100000000 + i * 7919:09d}" for i in range(3)]


class TricksTest(unittest.TestCase):
    def test_helpers(self):
        self.assertEqual([has_invisible(v) for v in ("040000011", " 040000011", "040 000 011", "04000\u200b0011", 5)],
                         [False, True, True, True, False])
        self.assertEqual(show_invisible(" 0400\u00a0011 "), "␣0400⟨неразрывный пробел⟩011␣")
        self.assertEqual([number_as_text(v) for v in ("1 234,56", "'100"[1:], "нет", 5)], [True, True, False, False])

    def test_list_file(self):
        path = os.path.join(tempfile.mkdtemp(), "т.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Л"
        for row in (["ИНН", "Название", "Сумма"], [INNS[0], "ООО «Альфа»", 10], [f" {INNS[1]}", "Бета", "1 000,50"],
                    [INNS[0], "альфа ооо", 5], [INNS[2], "Тайная", 99]):
            ws.append(row)
        ws["C5"].font = Font(color="FFFFFF")
        ws["A1"].font = Font(color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")      # белое на синем — видно
        wb.save(path)
        self.assertEqual(white_cells(path, "Л"), [("C5", "99")])
        maps = {path: [{"sheet": "Л", "head": 0, "row0": 0, "col0": 0, "kind": "list", "cols": {"inn": 0, "name": 1},
                        "sums": [2], "repeats": True}]}
        data = build(*scan_files([path], maps)["raw"])
        d = data.details
        self.assertEqual([r[4] for r in d["text_num_list"]], ["1 000,50"])
        self.assertEqual([r[4] for r in d["spaces_list"]], [f"␣{INNS[1]}"])
        self.assertEqual([r[2:4] for r in d["white_text"]], [["C5", "99"]])
        self.assertEqual([(r[0], r[2], r[3]) for r in d["inn_names"]], [(INNS[0], 2, "ООО «Альфа» | альфа ооо")])


if __name__ == "__main__":
    unittest.main()
