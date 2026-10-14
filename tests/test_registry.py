import os
import sys
import unittest
from datetime import date

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "gib_app")))

from collections import Counter  # noqa: E402
from analysis_lists import ListCollector  # noqa: E402
from analysis_model import APPEARED, BAD_INN, DUP_PAY, GONE, GREW, ODD_DATE, PAYER, status  # noqa: E402
from report_analysis import write_analysis_report  # noqa: E402
from analysis_registry import inn_in_name  # noqa: E402
from analysis_scan import build  # noqa: E402
import tempfile  # noqa: E402

from openpyxl import Workbook  # noqa: E402

from registry_columns import guess_header, registry_columns  # noqa: E402
from registry_read import Payment, inn_text, read_registry, to_payment  # noqa: E402
from registry_stats import Registry, period_label  # noqa: E402
from app_paths import ANALYSIS_FOLDER  # noqa: E402
from report_paths import analysis_report_path  # noqa: E402

HEAD = ["№ Документа", "Дата документа", "Дата учета", "Плательщик", "ИНН Плат", "Счет Плательщика", "Сумма",
        "БИК плат", "Получатель", "ИНН получ", "Счет Получателя", "БИК Получ", "КоррСчет получ", "ТипДокумента",
        "Основание"]


def pay(row, payer_inn, rec_inn, amount, d=date(2024, 1, 10), doc=None):
    return Payment(row, doc or str(row), d, "Орг " + payer_inn, payer_inn, "Получ " + rec_inn, rec_inn, amount)


class TestRegistryRead(unittest.TestCase):
    def test_columns_by_header(self):
        cols = registry_columns(HEAD)
        self.assertEqual((cols["payer"], cols["payer_inn"], cols["amount"], cols["rec"], cols["rec_inn"]),
                         (3, 4, 6, 8, 9))
        self.assertIsNone(registry_columns(["№", "Основание", "Расходы"]))

    def test_inn_text_restores_zero(self):
        self.assertEqual(inn_text(10002120), "010002120")
        self.assertEqual(inn_text(" 040000011 "), "040000011")

    def test_totals_row_skipped(self):
        cols = registry_columns(HEAD)
        self.assertIsNone(to_payment(5, [None] * 6 + [3219817.0] + [None] * 8, cols))
        p = to_payment(2, ["1", date(2024, 1, 4), None, "А", "010002120", None, 200, None, "Б", "251010603"], cols)
        self.assertEqual((p.row, p.amount, p.rec_inn), (2, 200.0, "251010603"))


class TestOtherDocuments(unittest.TestCase):
    def test_tajik_and_other_headers(self):
        cols = registry_columns(["№", "Сана", "Номи пардохткунанда", "РМА пардохткунанда", "Маблағ",
                                 "Гиранда", "РМА гиранда"])
        self.assertEqual((cols["payer_inn"], cols["amount"], cols["rec_inn"], cols["date"]), (3, 4, 6, 1))
        cols = registry_columns(["Наименование плательщика", "ИНН плательщика", "Сумма платежа"])
        self.assertEqual((cols["payer"], cols["payer_inn"], cols["amount"], cols["rec_inn"]), (0, 1, 2, None))

    def test_one_side_document(self):
        cols = registry_columns(["Получатель", "ИНН получ", "Сумма"])
        p = to_payment(2, ["Орг", "040000011", 100], cols)
        self.assertEqual((p.payer_inn, p.rec_inn), (None, "040000011"))
        r = Registry()
        f = r.start_file("x.xlsx", "Лист1")
        r.add(f, p)
        self.assertEqual((list(r.recs), r.payers, r.bad_inn), (["040000011"], {}, []))

    def test_guess_header(self):
        g = guess_header([["Отчёт"], [None], ["Колонка1", "Сумма", "Колонка3"], [1, 100, "x"]])
        self.assertEqual((g["head"], g["cols"]["amount"], g["heads"][0]), (2, 1, "Колонка1"))

    def _book(self, sheets):
        path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
        wb = Workbook()
        wb.remove(wb.active)
        for name, rows in sheets:
            ws = wb.create_sheet(name)
            for row in rows:
                ws.append(row)
        wb.save(path)
        return path

    def test_all_registry_sheets_are_read(self):
        head = ["Плательщик", "ИНН Плат", "Сумма", "Получатель", "ИНН получ"]
        path = self._book([("Январь", [head, ["А", "111111111", 10, "Б", "222222222"]]),
                           ("Свод", [["Итого", 10]]),
                           ("Февраль", [head, ["А", "111111111", 20, "Б", "222222222"]])])
        sheets, guess = read_registry(path)
        self.assertEqual(([s[0] for s in sheets], guess), (["Январь", "Февраль"], None))
        self.assertEqual([p.amount for _, _, rows in sheets for p in rows], [10.0, 20.0])

    def test_user_columns(self):
        path = self._book([("Лист1", [["К1", "К2", "К3"], ["111111111", "Орг", 55]])])
        sheets, guess = read_registry(path)
        self.assertEqual((sheets, guess["sheet"], guess["head"]), ([], "Лист1", 0))
        mapping = {"sheet": "Лист1", "head": 0, "cols": {"payer_inn": 0, "payer": 1, "amount": 2}}
        sheets, _ = read_registry(path, mapping)
        p = next(sheets[0][2])
        self.assertEqual((p.payer_inn, p.payer, p.amount, p.rec_inn, p.row), ("111111111", "Орг", 55.0, None, 2))


class TestRegistryStats(unittest.TestCase):
    def setUp(self):
        r = self.reg = Registry()
        f1 = r.start_file("a.xlsx", "Лист1")
        for i in range(200):
            r.add(f1, pay(i + 2, "111111111", "999999999", 10.0, date(2024, 1 + i % 4, 5)))
        r.add(f1, pay(300, "222222222", "999999999", 100.0))
        r.add(f1, pay(301, "333333333", "999999999", 50.0, d=date(2004, 3, 27)))     # странная дата
        r.add(f1, pay(302, "0", "999999999", 5.0))                                 # плохой ИНН
        r.add(f1, pay(303, "222222222", "999999999", 100.0, doc="300"))            # двойной платёж
        f2 = r.start_file("b.xlsx", "Лист1")
        for i in range(200):
            r.add(f2, pay(i + 2, "111111111", "999999999", 10.0, date(2024, 5 + i % 4, 5)))
        r.add(f2, pay(300, "333333333", "999999999", 500.0, date(2024, 6, 1)))
        r.add(f2, pay(301, "444444444", "999999999", 7.0, date(2024, 6, 1)))
        r.finish()

    def test_periods(self):
        self.assertEqual([self.reg.files[f].label for f in self.reg.order], ["Янв–Апр 2024", "Май–Авг 2024"])
        self.assertEqual(period_label([(2024, 1), (2025, 2)]), "Янв 2024 – Фев 2025")

    def test_errors(self):
        self.assertEqual([p.row for _, p in self.reg.odd_dates], [301])
        self.assertEqual([(p.row, side) for _, p, side, _ in self.reg.bad_inn], [(302, "плательщик")])
        self.assertEqual([(p.row, first) for _, p, first in self.reg.dups], [(303, (0, 300))])

    def _data(self):
        return build(self.reg, ListCollector(), [])

    def _payers(self, d):
        return {e.inn: e for e in d.entities if e.role == PAYER}

    def test_what_happened(self):
        d = self._data()
        p = self._payers(d)
        self.assertEqual([status(p[i], d.parts) for i in ("222222222", "444444444", "333333333", "0")],
                         [GONE, APPEARED, GREW, GONE])          # «0» платил только в 1-м файле

    def test_error_flags_and_details(self):
        d = self._data()
        p = self._payers(d)
        self.assertEqual((p["222222222"].flags, p["0"].flags, p["333333333"].flags), ({DUP_PAY}, {BAD_INN}, {ODD_DATE}))
        self.assertEqual({g: len(r) for g, r in d.details.items() if r}, {"dup": 2, "bad_pay": 1, "odd": 1})
        first, again = d.details["dup"]                       # пара: где впервые и где повтор
        self.assertEqual((first[2], first[6], again[2], again[6]), (300, "платёж", 303, "повтор"))
        self.assertEqual(d.details["bad_pay"][0][5:9], ["плательщика", "Орг 0", "0", "вместо ИНН стоит 0 (не указан)"])

    def test_excel_report(self):
        from openpyxl import load_workbook
        path = os.path.join(tempfile.mkdtemp(), "a.xlsx")
        res = write_analysis_report(path, self._data())
        # двойной платёж помечен у обеих сторон: плательщик 222… и получатель 999…
        self.assertEqual((res["statuses"][GONE], res["flags"][DUP_PAY]), (2, 2))
        wb = load_workbook(path)
        ws = wb["Анализ"]
        self.assertEqual(wb.sheetnames[:3], ["Анализ", "Кто кому платит", "Двойной платёж"])
        self.assertEqual(ws.cell(5, 2).value, "ИНН")
        self.assertIn("Новый: в «Янв–Апр 2024» этой организации не было", ws.cell(5, ws.max_column - 2).comment.text)
        self.assertIn("а в «Май–Авг 2024» она есть", ws.cell(4, 4).comment.text)   # подсказка в легенде («Новый»)
        err = wb["Двойной платёж"]
        self.assertEqual(err.cell(1, 1).value, "ДВОЙНОЙ ПЛАТЁЖ — 1")                # каждый вид — свой лист
        rules = sum(len(cf.rules) for cf in ws.conditional_formatting)
        self.assertEqual(rules, 6)                          # 5 цветов строк + оранжевые ошибки

    def test_change_and_compare_pair(self):
        from analysis_model import change
        d = self._data()
        p = self._payers(d)
        self.assertEqual(change(p["333333333"], d.parts, d.compare), (450.0, 9.0))     # 50 → 500: +900%
        self.assertEqual(change(p["444444444"], d.parts, d.compare), (None, None))     # нет в первой части
        rev = build(self.reg, ListCollector(), [], pair=(1, 0))                        # «было → стало» наоборот
        self.assertEqual(status(self._payers(rev)["444444444"], rev.parts, rev.compare), GONE)

    def test_months_mode(self):
        from analysis_periods import MONTHS_MODE, available_modes
        self.assertIn(MONTHS_MODE, available_modes(self.reg))
        d = build(self.reg, ListCollector(), [], mode=MONTHS_MODE)
        self.assertEqual([p.label for p in d.parts][:3], ["Янв 2024", "Фев 2024", "Мар 2024"])
        self.assertEqual(len(d.parts), 8)
        self.assertEqual(sum(self._payers(d)["111111111"].sums), 4000.0)              # все месяцы вместе

    def test_pairs(self):
        d = self._data()
        pairs = {(x.payer_inn, x.rec_inn): x for x in d.pairs}
        self.assertEqual(pairs[("333333333", "999999999")].sums, [50.0, 500.0])

    def test_same_name_different_inn(self):
        from analysis_checks import mark_same_names
        from analysis_model import SAME_NAME, Entity
        a, b = Entity("111111111", "ООО «Ромашка»", PAYER, [1.0], [1]), Entity("222222222", "ООО Ромашка", PAYER, [5.0], [1])
        c = Entity("333333333", "Другая", PAYER, [1.0], [1])
        rows = mark_same_names([a, b, c])
        self.assertEqual([(r[0], r[3]) for r in rows], [(1, "222222222"), (1, "111111111")])
        self.assertEqual((SAME_NAME in a.flags, SAME_NAME in c.flags), (True, False))

    def test_inn_in_name(self):
        self.assertEqual(inn_in_name("Сарраесати хазинадори ИНН 010010058"), "010010058")
        self.assertEqual(inn_in_name("счёт 20204972221022300242"), "")


class TestInnPayments(unittest.TestCase):
    def _book(self, rows):
        path = os.path.join(tempfile.mkdtemp(), "r.xlsx")
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        wb.save(path)
        return path

    def test_payments_of_one_inn(self):
        from report_inn import clean_inn, write_inn_payments
        self.assertEqual(clean_inn(" 4000 0011 "), "040000011")
        reg = self._book([["Плательщик", "ИНН Плат", "Сумма", "Получатель", "ИНН получ"],
                          ["А", "040000011", 100, "Б", "222222222"], ["В", "333333333", 50, "А", "040000011"],
                          ["В", "333333333", 70, "Г", "444444444"]])
        out = os.path.join(tempfile.mkdtemp(), "p.xlsx")
        res = write_inn_payments(out, [reg], "40000011")
        self.assertEqual((res["paid"], res["got"], res["paid_sum"], res["got_sum"]), (1, 1, 100.0, 50.0))
        self.assertTrue(os.path.exists(out))
        none = write_inn_payments(os.path.join(tempfile.mkdtemp(), "n.xlsx"), [reg], "999999999")
        self.assertIsNone(none["path"])                     # ничего не нашлось — файл не создаётся

    def test_no_registries(self):
        """Реестров нет — ИНН ищется в таблицах; листов «Платил» / «Получал» нет."""
        from openpyxl import load_workbook
        from report_inn import write_inn_payments
        lst = self._book([["Номи корхона", "РМА"], ["Орг", "040000011"]])
        res = write_inn_payments(os.path.join(tempfile.mkdtemp(), "x.xlsx"), [lst], "040000011")
        self.assertEqual((res["tables"], res["registries"]), (1, 0))
        self.assertEqual(load_workbook(res["path"]).sheetnames, ["Где нашлось", "Из таблиц"])


class TestPaths(unittest.TestCase):
    def test_report_goes_to_analysis_folder(self):
        path = analysis_report_path()
        self.assertEqual(os.path.basename(os.path.dirname(path)), ANALYSIS_FOLDER)
        self.assertTrue(os.path.basename(path).startswith("Анализ "))


class TestLists(unittest.TestCase):
    def test_two_tables_side_by_side(self):
        rows = [["№", "Номи корхона", "РМА", "Маблағ 2024", None, "№", "Номи корхона", "РМА", "Маблағ 2025"]]
        for i in range(12):
            rows.append([i + 1, f"Орг {i}", f"0100000{i:02d}", 100.0, None, i + 1, f"Орг {i}", f"0100000{i:02d}",
                         100.0 if i < 10 else 300.0])
        rows.append([13, "Только слева", "020000001", 50.0, None, 13, "Только справа", "030000001", 70.0])
        c = ListCollector()
        self.assertEqual(c.add_file("t.xlsx", {"Лист1": rows}), 2)
        d = build(Registry(), c, ["t.xlsx"])
        self.assertEqual([p.label for p in d.parts], ["2024 · «Лист1» A–D", "2025 · «Лист1» F–I"])
        st = Counter(status(e, d.parts) for e in d.entities)
        self.assertEqual((st[GONE], st[APPEARED], st[GREW]), (1, 1, 2))

    def test_part_without_sums_compares_presence(self):
        from analysis_model import SAME
        rows = [["Номи корхона", "РМА", "Маблағ 2024", None, "Номи корхона", "РМА"]]
        rows += [[f"Орг {i}", f"0100000{i:02d}", 50.0 + 7 * i, None, f"Орг {i}", f"0100000{i:02d}"] for i in range(12)]
        c = ListCollector()
        c.add_file("t.xlsx", {"Лист1": rows})
        d = build(Registry(), c, ["t.xlsx"])
        self.assertEqual([p.has_sums for p in d.parts], [True, False])
        self.assertEqual(Counter(status(e, d.parts) for e in d.entities), Counter({SAME: 12}))


if __name__ == "__main__":
    unittest.main()
