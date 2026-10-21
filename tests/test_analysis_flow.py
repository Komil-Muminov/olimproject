"""Фильтр по файлам: у каждого файла карточка; не взятый и пустой файл не мешают остальным."""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "gib_app")))

from openpyxl import Workbook  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

import ui_analysis_flow as flow  # noqa: E402
from file_sheets import read_sheets  # noqa: E402

INNS = [f"{100000000 + i * 7919:09d}" for i in range(3)]


def _save(name, amount):
    path = os.path.join(tempfile.mkdtemp(), name)
    wb = Workbook()
    wb.active.title = "Л"
    for r in (["ИНН", "Название", "Сумма"], [INNS[0], "А", amount], [INNS[1], "Б", 10]):
        wb.active.append(r)
    wb.save(path)
    return path


class AnalysisFlowTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication(sys.argv)

    def _win(self, props):
        """Экран анализа без окна: карточки и перехват запуска фоновой задачи."""
        from ui_file_card import FileCard
        win = type("Win", (), {})()
        win.worker, win.an_run_btn = None, type("B", (), {"setEnabled": lambda s, v: setattr(s, "on", v)})()
        win.an_file_cards = [FileCard(p) for p in props]
        from PySide6.QtWidgets import QLineEdit
        win.an_inn, win.an_min, win.an_max = QLineEdit("040058284"), QLineEdit(), QLineEdit()   # фильтр: ИНН
        self.started = None
        flow._run = lambda w, func, *args, done: setattr(self, "started", args)
        return win

    def setUp(self):
        self.saved = flow._run

    def tearDown(self):
        flow._run = self.saved

    def test_only_included_files_go_to_analysis(self):
        a, b, c = (_save(n, 5) for n in ("a.xlsx", "b.xlsx", "c.xlsx"))
        win = self._win([read_sheets(p) for p in (a, b, c)])
        for card in win.an_file_cards:
            card.rows[0].edits["inn"].setText("a")
        win.an_file_cards[1].use.setChecked(False)              # b не берём
        flow.run(win)
        _, paths, mappings, inn, threshold = self.started
        self.assertEqual((paths, set(mappings), inn, threshold), ([a, c], {a, c}, ["040058284"], None))

    def test_not_ready_until_inn_given(self):
        win = self._win([read_sheets(_save("a.xlsx", 5))])
        flow.run(win)
        self.assertIsNone(self.started)                          # колонка ИНН не вписана — поиск не запускается

    def test_sheets_with_checkboxes(self):
        """Отмечены «Янв» и «Мар»: у «Мар» поля пустые — колонки как у «Янв»; «Фев» не берётся."""
        import tempfile
        from ui_file_card import FileCard
        path = os.path.join(tempfile.mkdtemp(), "год.xlsx")
        wb = Workbook()
        for k, name in enumerate(("Янв", "Фев", "Мар")):
            ws = wb.active if k == 0 else wb.create_sheet()
            ws.title = name
            for r in (["ИНН", "Сумма"], [INNS[k], 5]):
                ws.append(r)
        wb.save(path)
        card = FileCard(read_sheets(path))
        self.assertEqual([r.included() for r in card.rows], [True, False, False])     # сначала — первый лист
        card.rows[0].edits["inn"].setText("a")
        card.rows[0].edits["sums"].setText("b")
        card.rows[2].use.setChecked(True)
        self.assertEqual(card.problem(), "")
        self.assertEqual([(m["sheet"], m["sums"]) for m in card.mapping()], [("Янв", [1]), ("Мар", [1])])
        self.assertEqual(card.rows[2].edits["inn"].placeholderText(), "как у «Янв»")
        card.rows[0].use.setChecked(False)                                           # колонки взять неоткуда
        self.assertIn("Мар", card.problem())

    def test_empty_file_card(self):
        from ui_file_card import FileCard
        card = FileCard({"path": "пустой.xlsx", "sheets": {}})
        self.assertEqual((card.included(), card.use.isEnabled(), card.problem()), (False, False, ""))


if __name__ == "__main__":
    unittest.main()
