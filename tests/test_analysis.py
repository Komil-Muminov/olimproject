"""Глубокий анализ по таблицам с ИНН: части (листы), «одно название — разные ИНН», Excel «Анализ», папка отчёта."""
import os
import sys
import tempfile
import unittest
from collections import Counter

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "gib_app")))

from openpyxl import load_workbook  # noqa: E402

from analysis_lists import ListCollector  # noqa: E402
from analysis_model import ORG, SAME_NAME, Entity  # noqa: E402
from analysis_scan import build  # noqa: E402
from app_paths import ANALYSIS_FOLDER  # noqa: E402
from registry_columns import guess_header  # noqa: E402
from report_analysis import write_analysis_report  # noqa: E402
from report_paths import analysis_report_path  # noqa: E402


def _two_tables():
    rows = [["№", "Номи корхона", "РМА", "Маблағ 2024", None, "№", "Номи корхона", "РМА", "Маблағ 2025"]]
    for i in range(12):
        rows.append([i + 1, f"Орг {i}", f"0100000{i:02d}", 100.0, None, i + 1, f"Орг {i}", f"0100000{i:02d}",
                     100.0 if i < 10 else 300.0])
    rows.append([13, "Только слева", "020000001", 50.0, None, 13, "Только справа", "030000001", 70.0])
    c = ListCollector()
    c.add_file("t.xlsx", {"Лист1": rows})
    return c


class TestLists(unittest.TestCase):
    def test_two_tables_side_by_side(self):
        d = build(_two_tables(), ["t.xlsx"])
        self.assertEqual([p.label for p in d.parts], ["2024 · «Лист1» A–D", "2025 · «Лист1» F–I"])
        where = Counter(tuple(c > 0 for c in e.counts) for e in d.entities)
        self.assertEqual(where, Counter({(True, True): 12, (True, False): 1, (False, True): 1}))

    def test_part_without_sums(self):
        rows = [["Номи корхона", "РМА", "Маблағ 2024", None, "Номи корхона", "РМА"]]
        rows += [[f"Орг {i}", f"0100000{i:02d}", 50.0 + 7 * i, None, f"Орг {i}", f"0100000{i:02d}"] for i in range(12)]
        c = ListCollector()
        c.add_file("t.xlsx", {"Лист1": rows})
        d = build(c, ["t.xlsx"])
        self.assertEqual([p.has_sums for p in d.parts], [True, False])
        self.assertTrue(all(e.counts[0] and e.counts[1] for e in d.entities))

    def test_excel_report(self):
        path = os.path.join(tempfile.mkdtemp(), "a.xlsx")
        write_analysis_report(path, build(_two_tables(), ["t.xlsx"]))
        wb = load_workbook(path)
        ws = wb["Анализ"]
        head = [c.value for c in ws[5]]
        self.assertEqual(head[:3], ["№", "ИНН", "Название"])
        self.assertEqual(head[-4:], ["Есть в файлах", "Файлов", "Ошибки", "Названий"])
        self.assertEqual(ws.cell(4, 4).value, "Ошибки")                            # легенда — цвет ошибок
        for gone in ("Кто кому платит", "По месяцам", "Сверка со списком", "Двойной платёж"):
            self.assertNotIn(gone, wb.sheetnames)
        self.assertEqual(sum(len(cf.rules) for cf in ws.conditional_formatting), 3)   # ошибки + разные названия


class TestChecks(unittest.TestCase):
    def test_same_name_different_inn(self):
        from analysis_checks import mark_same_names
        a, b = Entity("111111111", "ООО «Ромашка»", ORG, [1.0], [1]), Entity("222222222", "ООО Ромашка", ORG, [5.0], [1])
        c = Entity("333333333", "Другая", ORG, [1.0], [1])
        rows = mark_same_names([a, b, c])
        self.assertEqual([(r[0], r[3]) for r in rows], [(1, "222222222"), (1, "111111111")])
        self.assertEqual((SAME_NAME in a.flags, SAME_NAME in c.flags), (True, False))

    def test_guess_header(self):
        g = guess_header([["Отчёт"], [None], ["Колонка1", "Сумма", "Колонка3"], [1, 100, "x"]])
        self.assertEqual((g["head"], g["cols"]["amount"], g["heads"][0]), (2, 1, "Колонка1"))


class TestPaths(unittest.TestCase):
    def test_report_goes_to_analysis_folder(self):
        path = analysis_report_path()
        self.assertEqual(os.path.basename(os.path.dirname(path)), ANALYSIS_FOLDER)
        self.assertTrue(os.path.basename(path).startswith("Анализ "))


if __name__ == "__main__":
    unittest.main()
