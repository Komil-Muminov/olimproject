import os
import sys
import tempfile
import unittest
from datetime import date

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "gib_app")))

from openpyxl import Workbook  # noqa: E402

from registry_columns import two_row_columns  # noqa: E402
from registry_read import read_registry  # noqa: E402

TOP = ["№ док", "Дата док", "Плательщик", None, "Получатель", None, "Сумма"]
BOTTOM = [None, None, "Наименование", "ИНН", "Наименование", "ИНН", None]


class TwoRowHeaderTest(unittest.TestCase):
    def test_columns(self):
        cols = two_row_columns(TOP, BOTTOM)
        self.assertEqual({k: cols[k] for k in ("doc", "date", "payer", "payer_inn", "rec", "rec_inn", "amount")},
                         {"doc": 0, "date": 1, "payer": 2, "payer_inn": 3, "rec": 4, "rec_inn": 5, "amount": 6})

    def test_top_below(self):
        """«ИНН» сверху (объединена), «Плательщик | Получатель» снизу."""
        cols = two_row_columns(["Сумма", "ИНН", None], [None, "Плательщика", "Получателя"])
        self.assertEqual((cols["amount"], cols["payer_inn"], cols["rec_inn"]), (0, 1, 2))

    def test_read_file(self):
        path = os.path.join(tempfile.mkdtemp(), "две строки.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.append(["РЕЕСТР ПЛАТЕЖЕЙ"])
        ws.append(TOP)
        ws.append(BOTTOM)
        ws.merge_cells("C2:D2")
        ws.merge_cells("E2:F2")
        ws.append([1, date(2025, 1, 5), "ООО А", "111111111", "ООО Б", "222222222", 100.5])
        wb.save(path)
        sheets, _ = read_registry(path)
        p = next(sheets[0][2])
        self.assertEqual((p.row, p.payer, p.payer_inn, p.rec_inn, p.amount), (4, "ООО А", "111111111", "222222222", 100.5))


if __name__ == "__main__":
    unittest.main()
