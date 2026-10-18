import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "gib_app")))

from openpyxl import Workbook  # noqa: E402

from analysis_list_map import AS_LIST  # noqa: E402
from analysis_propose import propose  # noqa: E402
from analysis_scan import scan_files  # noqa: E402

INNS = [f"{100000000 + i * 7919:09d}" for i in range(12)]


def _save(rows, title="Л"):
    path = os.path.join(tempfile.mkdtemp(), "f.xlsx")
    wb = Workbook()
    wb.active.title = title
    for r in rows:
        wb.active.append(r)
    wb.save(path)
    return path


class ProposeTest(unittest.TestCase):
    def test_inn_found_on_each_sheet(self):
        """На каждом листе колонка ИНН и название — буквами, как вписал бы пользователь; маленький лист — не найдено."""
        path = os.path.join(tempfile.mkdtemp(), "два листа.xlsx")
        wb = Workbook()
        wb.active.title = "Банк"
        for r in [["Контрагент", "Дебет", "ИНН"]] + [[f"Орг {i}", 100, inn] for i, inn in enumerate(INNS)]:
            wb.active.append(r)
        ws = wb.create_sheet("Мало")
        ws.append(["РМА"])
        ws.append([INNS[0]])
        wb.save(path)
        sheets = propose(path)["sheets"]
        self.assertEqual((sheets["Банк"]["inn"], sheets["Банк"]["name"]), ("C", "A"))
        self.assertIsNone(sheets["Мало"]["inn"])

    def test_list_manual_sums(self):
        """Список: пользователь выбрал суммировать Дебет и Кредит."""
        path = _save([["ИНН", "Контрагент", "Дебет", "Кредит"]] + [[inn, f"Орг {i}", 100, 1] for i, inn in enumerate(INNS)])
        mapping = {"sheet": "Л", "head": 0, "row0": 0, "col0": 0, "cols": {"inn": 0, "name": 1}, "kind": AS_LIST,
                   "sums": [2, 3]}
        lists, _ = scan_files([path], {path: [mapping]})["raw"]
        self.assertEqual({e.sums[0] for e in lists.entities(0, 1)}, {101.0})

    def test_table_not_from_a1(self):
        """Таблица с C3: первая колонка листа и колонка ИНН — как в Excel."""
        path = os.path.join(tempfile.mkdtemp(), "c3.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Л"
        for r, row in enumerate([["Название", "ИНН"]] + [[f"Орг {i}", inn] for i, inn in enumerate(INNS)], 3):
            for c, v in enumerate(row, 3):
                ws.cell(r, c, v)
        wb.save(path)
        sh = propose(path)["sheets"]["Л"]
        self.assertEqual((sh["row0"], sh["col0"], sh["inn"], sh["name"]), (0, 2, "D", "C"))

    def test_sheets_dialog(self):
        """Окно «Листы и ИНН»: ИНН нашёлся сам — ✓; не нашёлся — «Анализировать» не нажимается, пока не вписать
        колонку или не отметить «Без ИНН» (лист — в свод как есть)."""
        from PySide6.QtWidgets import QApplication
        from analysis_list_map import RAW
        from ui_sheets_dialog import FOUND, NOT_FOUND, SheetsDialog
        app = QApplication.instance() or QApplication([])  # noqa: F841
        path = os.path.join(tempfile.mkdtemp(), "f.xlsx")
        wb = Workbook()
        wb.active.title = "Банк"
        for r in [["Название", "ИНН"]] + [[f"Орг {i}", inn] for i, inn in enumerate(INNS)]:
            wb.active.append(r)
        wb.create_sheet("Прочее").append(["просто", "текст"])
        wb.save(path)
        d = SheetsDialog(None, [propose(path)])
        ok = d.buttons.button(d.buttons.StandardButton.Ok)
        bank, other = d.widgets
        self.assertEqual((bank["inn"].text(), bank["name"].text(), bank["status"].text()), ("B", "A", FOUND))
        self.assertEqual(other["status"].text(), NOT_FOUND)
        self.assertFalse(ok.isEnabled())
        other["inn"].setText("a")                                   # в колонке A нет ИНН — понятная ошибка
        self.assertIn("нет ИНН", other["status"].text())
        other["raw"].setChecked(True)
        self.assertTrue(ok.isEnabled())
        m = d.mappings()[path]
        self.assertEqual([(x["sheet"], x["kind"]) for x in m], [("Банк", AS_LIST), ("Прочее", RAW)])
        self.assertEqual((m[0]["cols"]["inn"], m[0]["cols"]["name"], m[0]["sums"]), (1, 0, []))
        bank["use"].setChecked(False)                                # лист не берём
        other["use"].setChecked(False)
        self.assertFalse(ok.isEnabled())


if __name__ == "__main__":
    unittest.main()
