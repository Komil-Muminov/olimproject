"""Фильтр по всем файлам: ИНН (в любой колонке) и/или порог суммы — строки как есть, без предела по строкам."""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "gib_app")))

from openpyxl import Workbook, load_workbook  # noqa: E402

from column_entry import to_mapping  # noqa: E402
from file_sheets import read_sheets  # noqa: E402
from report_inn_search import clean_inn, write_filter_report  # noqa: E402

INN = "040058284"


def _save(name, rows):
    path = os.path.join(tempfile.mkdtemp(), name)
    wb = Workbook()
    wb.active.title = "Л"
    for r in rows:
        wb.active.append(r)
    wb.save(path)
    return path


def _map(path, **typed):
    return [to_mapping("Л", read_sheets(path)["sheets"]["Л"], typed)]


class InnSearchTest(unittest.TestCase):
    def test_clean(self):
        self.assertEqual([clean_inn(" 40058284 "), clean_inn("040 058 284"), clean_inn("abc")], [INN, INN, ""])

    def test_rows_from_all_files(self):
        a = _save("30 фоиза.xlsx", [["РМА", "Номи", "30 фоиза"], [INN, "Шабака", 39], ["630000962", "Фурудгоҳ", 27]])
        b = _save("M1-4.xlsx", [["Дата", "ИНН", "Сумма", "Основание"], ["09.01.2024", INN, 100, "налог"],
                                ["10.01.2024", "630000962", 5, "x"], ["11.01.2024", int(INN), 50, "дивиденд"]])
        maps = {a: _map(a, inn="a", sums="c"), b: _map(b, inn="b", sums="c")}
        out = os.path.join(tempfile.mkdtemp(), "r.xlsx")
        res = write_filter_report(out, [a, b], maps, "40058284")              # ноль впереди потерян — не страшно
        self.assertEqual((res["rows"], res["files"]), (3, [("30 фоиза.xlsx", 1, 39.0), ("M1-4.xlsx", 2, 150.0)]))
        rows = list(load_workbook(out).active.iter_rows(min_row=8, values_only=True))
        self.assertEqual(rows[0][:5], ("Файл", "Лист", "Строка", "Сумма строки", "30 фоиза.xlsx: A РМА"))
        self.assertEqual([r[:3] for r in rows[1:]], [("30 фоиза.xlsx", "Л", 2), ("M1-4.xlsx", "Л", 2), ("M1-4.xlsx", "Л", 4)])
        self.assertIn("дивиденд", rows[-1])
        res = write_filter_report(os.path.join(tempfile.mkdtemp(), "p.xlsx"), [a, b], maps, INN, threshold=(60, None))
        self.assertEqual(res["files"], [("M1-4.xlsx", 1, 100.0)])           # порог «от 60»: 39 и 50 — мимо

    def test_sheets_of_one_file(self):
        """Два листа одного файла — две таблицы; колонки подписаны «файл · лист»."""
        path = os.path.join(tempfile.mkdtemp(), "год.xlsx")
        wb = Workbook()
        wb.active.title = "Янв"
        for r in (["ИНН", "Сумма"], [INN, 5], ["630000962", 1]):
            wb.active.append(r)
        ws = wb.create_sheet("Фев")
        for r in (["Сумма", "ИНН"], [7, INN]):
            ws.append(r)
        wb.save(path)
        sheets = read_sheets(path)["sheets"]
        maps = {path: [to_mapping("Янв", sheets["Янв"], {"inn": "a", "sums": "b"}),
                       to_mapping("Фев", sheets["Фев"], {"inn": "b", "sums": "a"})]}
        out = os.path.join(tempfile.mkdtemp(), "r.xlsx")
        res = write_filter_report(out, [path], maps, INN)
        self.assertEqual(res["files"], [("год.xlsx · Янв", 1, 5.0), ("год.xlsx · Фев", 1, 7.0)])
        rows = list(load_workbook(out).active.iter_rows(min_row=8, values_only=True))
        self.assertEqual(rows[0][4:], ("год.xlsx · Янв: A ИНН", "год.xlsx · Янв: B Сумма",
                                       "год.xlsx · Фев: A Сумма", "год.xlsx · Фев: B ИНН"))

    def test_inn_in_any_column(self):
        """ИНН не в колонке ИНН, а в основании — строка тоже находится."""
        a = _save("a.xlsx", [["ИНН", "Сумма", "Основание"], ["630000962", 7, f"за ООО «Шабака» (ИНН {INN})"],
                             ["630000962", 9, "налог 1040058284"]])               # 10 цифр — другой номер
        out = os.path.join(tempfile.mkdtemp(), "r.xlsx")
        res = write_filter_report(out, [a], {a: _map(a, inn="a", sums="b")}, INN)
        self.assertEqual(res["files"], [("a.xlsx", 1, 7.0)])

    def test_threshold_only(self):
        """Без ИНН, только сумма «от 40»: все записи с такой суммой строки из всех файлов."""
        a = _save("a.xlsx", [["ИНН", "Сумма"], [INN, 39], ["630000962", 45]])
        b = _save("b.xlsx", [["ИНН", "Сумма"], [INN, 100]])
        maps = {a: _map(a, inn="a", sums="b"), b: _map(b, inn="a", sums="b")}
        res = write_filter_report(os.path.join(tempfile.mkdtemp(), "r.xlsx"), [a, b], maps, "", (40, None))
        self.assertEqual((res["title"], res["files"]), ("суммы от 40,00", [("a.xlsx", 1, 45.0), ("b.xlsx", 1, 100.0)]))

    def test_not_found(self):
        a = _save("a.xlsx", [["ИНН"], ["630000962"]])
        self.assertIsNone(write_filter_report(os.path.join(tempfile.mkdtemp(), "r.xlsx"), [a], {a: _map(a, inn="a")}, INN)["path"])


if __name__ == "__main__":
    unittest.main()
