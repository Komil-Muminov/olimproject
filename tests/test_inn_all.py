import os
import sys
import tempfile
import unittest
from datetime import date

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "gib_app")))

from openpyxl import Workbook, load_workbook  # noqa: E402

from report_inn import write_inn_payments  # noqa: E402

INN = "040058284"


def _save(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)
    return path


class InnEverywhereTest(unittest.TestCase):
    def test_tables_and_payments(self):
        d = tempfile.mkdtemp()
        lst = _save(os.path.join(d, "30 фоиза.xlsx"), [
            ("Воридот", [["Маълумот"], [], ["№", "Номи", "РМА", "Бакия ба 01.01.2024"],
                         [1, "ҶСК А", INN, -39141.5], [2, "ҶСК Б", "630000962", 0]]),
            ("Пусто", [["нет ИНН"]])])
        reg = _save(os.path.join(d, "M1.xlsx"), [("Л", [
            ["№ док", "Дата док", "Плательщик", "ИНН плательщика", "Получатель", "ИНН получателя", "Сумма"],
            [1, date(2024, 1, 9), "ЧСК А", INN, "Вазорат", "111111111", 100.0],
            [2, date(2024, 1, 9), "Вазорат", "111111111", "ЧСК А", INN, 7.0]])])
        res = write_inn_payments(os.path.join(d, "out.xlsx"), [lst, reg], "40058284")   # ноль впереди потерян
        self.assertEqual((res["tables"], res["paid"], res["got"]), (1, 1, 1))
        wb = load_workbook(res["path"])
        self.assertEqual(wb.sheetnames, ["Где нашлось", "Из таблиц", "Платил", "Получал"])
        cells = {r[4]: r[5] for r in wb["Из таблиц"].iter_rows(min_row=5, values_only=True)}
        self.assertEqual((cells["B: Номи"], cells["D: Бакия ба 01.01.2024"]), ("ҶСК А", -39141.5))
        where = [r[1:5] for r in wb["Где нашлось"].iter_rows(min_row=5, values_only=True)]
        self.assertEqual(where[0], ("30 фоиза.xlsx", "Воридот", 4, "таблица"))

    def test_not_found(self):
        d = tempfile.mkdtemp()
        lst = _save(os.path.join(d, "a.xlsx"), [("Л", [["ИНН"], ["630000962"]])])
        self.assertIsNone(write_inn_payments(os.path.join(d, "out.xlsx"), [lst], INN)["path"])


if __name__ == "__main__":
    unittest.main()
