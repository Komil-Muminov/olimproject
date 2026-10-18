"""Для аудитора: лист «Как считали», «По месяцам», скрытые листы и строки, формулы без значений, .xls без Excel."""
import os
import sys
import tempfile
import unittest
from datetime import date

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "gib_app")))

from openpyxl import Workbook, load_workbook  # noqa: E402

from analysis_jobs import make_report  # noqa: E402
from analysis_scan import build, scan_files  # noqa: E402
from excel_sheet_info import sheet_scan, workbook_sheets  # noqa: E402

INNS = [f"{100000000 + i * 7919:09d}" for i in range(4)]
LIST = {"sheet": "Л", "head": 0, "row0": 0, "col0": 0, "kind": "list", "cols": {"inn": 1, "name": None}, "sums": [],
        "labels": {"ИНН": "B: ИНН плательщика"}}


def _table(hide_row=None, hidden_sheet=False):
    path = os.path.join(tempfile.mkdtemp(), "р.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Л"
    ws.append(["Дата документа", "ИНН плательщика", "ИНН получателя", "Сумма"])
    for k, m in enumerate((1, 1, 2, 3)):
        ws.append([date(2024, m, 5), INNS[k], INNS[3], 100.0 * (k + 1)])
    if hide_row:
        ws.row_dimensions[hide_row].hidden = True
    if hidden_sheet:
        h = wb.create_sheet("Тайный")
        h.sheet_state = "hidden"
        h.append(["x"])
    wb.save(path)
    return path


class AuditSheetsTest(unittest.TestCase):
    def _report(self, path):
        raw = scan_files([path], {path: [LIST]})["raw"]
        out = os.path.join(tempfile.mkdtemp(), "a.xlsx")
        make_report(out, raw, {path: [LIST]})
        return build(*raw), load_workbook(out)

    def test_settings(self):
        _, wb = self._report(_table())
        text = [[v for v in r if v is not None] for r in wb["Как считали"].iter_rows(values_only=True)]
        self.assertIn(["р.xlsx", "Л", "ИНН", "B: ИНН плательщика"], text)
        self.assertIn(["р.xlsx", "Л", "таблица с ИНН", "Строки", "с 2-й до конца листа"], text)
        self.assertNotIn("По месяцам", wb.sheetnames)

    def test_hidden_row_and_sheet(self):
        path = _table(hide_row=3, hidden_sheet=True)
        self.assertEqual((workbook_sheets(path)["Тайный"]["hidden"], sheet_scan(path, "Л")["hidden_rows"]), (True, {3}))
        data, wb = self._report(path)
        self.assertEqual([r[2] for r in data.details["hidden_list"]], [3])
        self.assertIn("Скрытые строки в таблицах", wb.sheetnames)

    def test_formulas_without_values(self):
        path = os.path.join(tempfile.mkdtemp(), "ф.xlsx")
        wb = Workbook()
        wb.active.append(["ИНН", "Цена", "Сумма"])
        wb.active.append([INNS[0], 5, "=B2*2"])
        wb.save(path)
        self.assertEqual(sheet_scan(path, "Sheet", formulas=True)["formulas"], {2: 1})

    def test_old_xls_without_excel(self):
        """Старый .xls читается напрямую (быстро), значения — как у .xlsx."""
        xls = os.path.join(os.path.dirname(__file__), "..", "ОлимФайлы", "Вписка 2025.xls")
        if not os.path.exists(xls):
            self.skipTest("нет файла-примера")
        from excel_reader import Book
        self.assertTrue(Book(xls).sheets)


if __name__ == "__main__":
    unittest.main()
