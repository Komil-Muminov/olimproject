"""Колонки от руки (как в Olimv3): разбор записи, настройка чтения, сверка с «Итого», лист «Все файлы по ИНН»."""
import os
import sys
import tempfile
import unittest
from datetime import date

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "gib_app")))

from openpyxl import Workbook, load_workbook  # noqa: E402

from analysis_jobs import make_report  # noqa: E402
from analysis_propose import propose  # noqa: E402
from analysis_scan import auto_mode, scan_files  # noqa: E402
from column_entry import to_mapping  # noqa: E402
from column_spec import SpecError, cell_address, parse  # noqa: E402

INNS = [f"{100000000 + i * 7919:09d}" for i in range(4)]


def _save(name, rows):
    path = os.path.join(tempfile.mkdtemp(), name)
    wb = Workbook()
    wb.active.title = "Л"
    for r in rows:
        wb.active.append(r)
    wb.save(path)
    return path


def _map(path, **typed):
    p = propose(path)
    return [to_mapping("Л", p["sheets"]["Л"], typed, {})]


class TypedColumnsTest(unittest.TestCase):
    def test_parse(self):
        self.assertEqual(parse("c"), [(2, None, None)])
        self.assertEqual(parse("C5"), [(2, 5, None)])
        self.assertEqual(parse("c5-300"), [(2, 5, 300)])
        self.assertEqual(parse("B5-B20, С5-С20"), [(1, 5, 20), (2, 5, 20)])     # «С» — русская буква
        self.assertEqual(cell_address("e200"), (4, 200))
        for bad in ("5C", "C5-D9", "C9-C5", "C1, X"[:0] or "C5-"):
            with self.assertRaises(SpecError):
                parse(bad)

    def test_rows_from_inn_field(self):
        """«A3-A4» — только строки 3–4; «A» — с первой строки, где ИНН (шапка с итогом над данными — подписи
        берутся из настоящей шапки)."""
        path = _save("с.xlsx", [["Свод"], ["ИНН", "Название", "Сумма"], [None, "ИТОГО", 60],
                                [INNS[0], "А", 10], [INNS[1], "Б", 20], [INNS[2], "В", 30]])
        m = _map(path, inn="a3-a4", name="b", sums="c")[0]
        self.assertEqual((m["head"], m["last"]), (1, 4))
        _, lists, _ = scan_files([path], {path: [m]})["raw"]
        self.assertEqual([e.name for e in lists.entities(0, 1)], ["А"])          # строка 5 — за пределами
        m = _map(path, inn="a", name="b", sums="c")[0]
        self.assertEqual((m["head"], m["labels_row"], m["labels"]["Название"]), (2, 1, "b  →  B: Название"))

    def test_total_check_and_registry(self):
        path = _save("р.xlsx", [["Дата", "ИНН плат", "ИНН получ", "Сумма"], [date(2024, 1, 5), INNS[0], INNS[3], 70],
                                [date(2024, 1, 6), INNS[1], INNS[3], 30], [None, None, "Итого", 101]])
        maps = {path: _map(path, inn="b", rec_inn="c", sums="d", date="a", total="d4")}
        self.assertEqual(maps[path][0]["kind"], "registry")
        raw = scan_files([path], maps)["raw"]
        check = raw[0].checks[0]
        self.assertEqual((check["cell"], check["in_file"], check["program"], check["diff"]), ("D4", 101, 100.0, -1.0))

    def test_total_row_is_not_payment(self):
        """«Итого» — не платёж, а «Ҷамъияти дорои масъулияти маҳдуд» (ООО) — платёж, хоть и начинается с «Ҷамъ»."""
        from registry_read import to_payment
        cols = {"payer": 0, "payer_inn": 1, "amount": 2}
        self.assertIsNotNone(to_payment(5, ["Ҷамъияти дорои масъулияти маҳдуди «А»", INNS[0], 10], cols))
        self.assertIsNone(to_payment(6, ["Итого:", None, 10], cols))
        self.assertIsNone(to_payment(7, ["Ҷамъ", None, 10], cols))

    def test_all_files_by_inn(self):
        a = _save("список.xlsx", [["ИНН", "План"], [INNS[0], 5], [INNS[2], 7]])
        b = _save("реестр.xlsx", [["ИНН", "Сумма", "Основание"], [INNS[0], 1, "раз"], [INNS[0], 2, "два"], [INNS[1], 3, "x"]])
        maps = {a: _map(a, inn="a", sums="b"), b: _map(b, inn="a", sums="b")}
        raw = scan_files([a, b], maps)["raw"]
        out = os.path.join(tempfile.mkdtemp(), "r.xlsx")
        res = make_report(out, raw, auto_mode(raw), None, maps)
        self.assertEqual(res["all_by_inn"], {"orgs": 1, "rows": 3, "cut": False})   # только ИНН из обоих файлов
        rows = list(load_workbook(out)["Все файлы по ИНН"].iter_rows(min_row=4, values_only=True))
        self.assertEqual(rows[0][4:], ("список.xlsx: A ИНН", "список.xlsx: B План", "реестр.xlsx: A ИНН",
                                       "реестр.xlsx: B Сумма", "реестр.xlsx: C Основание"))
        self.assertEqual([r[:4] for r in rows[1:]], [(INNS[0], "список.xlsx", "Л", 2), (INNS[0], "реестр.xlsx", "Л", 2),
                                                     (INNS[0], "реестр.xlsx", "Л", 3)])

    def test_all_by_inn_sheets_of_one_file(self):
        """Три таблицы на трёх листах одного файла — тоже «Все файлы по ИНН»; колонки подписаны «файл · лист»."""
        path = os.path.join(tempfile.mkdtemp(), "год.xlsx")
        wb = Workbook()
        data = {"Янв": [["ИНН", "Сумма"], [INNS[0], 5], [INNS[1], 1]], "Фев": [["Сумма", "ИНН"], [7, INNS[0]]],
                "Мар": [["ИНН", "Название", "Сумма"], [INNS[2], "В", 9], [INNS[0], "А", 3]]}
        for k, (title, rows) in enumerate(data.items()):
            ws = wb.active if k == 0 else wb.create_sheet()
            ws.title = title
            for r in rows:
                ws.append(r)
        wb.save(path)
        p = propose(path)
        typed = {"Янв": {"inn": "a", "sums": "b"}, "Фев": {"inn": "b", "sums": "a"}, "Мар": {"inn": "a", "sums": "c"}}
        maps = {path: [to_mapping(s, p["sheets"][s], t, {}) for s, t in typed.items()]}
        raw = scan_files([path], maps)["raw"]
        out = os.path.join(tempfile.mkdtemp(), "r.xlsx")
        res = make_report(out, raw, auto_mode(raw), None, maps)
        self.assertEqual(res["all_by_inn"], {"orgs": 1, "rows": 3, "cut": False})   # только ИНН из 2+ листов
        rows = list(load_workbook(out)["Все файлы по ИНН"].iter_rows(min_row=4, values_only=True))
        self.assertEqual(rows[0][4:6], ("год.xlsx · Янв: A ИНН", "год.xlsx · Янв: B Сумма"))
        self.assertEqual([r[:4] for r in rows[1:]], [(INNS[0], "год.xlsx", "Янв", 2), (INNS[0], "год.xlsx", "Фев", 2),
                                                     (INNS[0], "год.xlsx", "Мар", 3)])

    def test_all_by_inn_needs_two_tables(self):
        path = _save("один.xlsx", [["ИНН", "Сумма"], [INNS[0], 5]])
        maps = {path: _map(path, inn="a", sums="b")}
        raw = scan_files([path], maps)["raw"]
        res = make_report(os.path.join(tempfile.mkdtemp(), "r.xlsx"), raw, auto_mode(raw), None, maps)
        self.assertIsNone(res["all_by_inn"])


if __name__ == "__main__":
    unittest.main()
